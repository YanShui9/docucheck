"""Build snapshot-bound, formula-free downloads from verified Run checks."""

import io
import json
import re

from openpyxl import Workbook
from sqlalchemy import Engine, select

from ..db.api_queries import ApiQueryError
from ..db.models import document_versions, evidence_blocks

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_MAX_CELL_CHARS = 32_767


def export_payload(engine: Engine, result: dict) -> dict:
    """Resolve only evidence cited by the already verified Run snapshot."""
    if not result["checks"]:
        raise ApiQueryError("checks_unavailable")
    evidence_ids = {
        evidence_id
        for check in result["checks"]
        for key in ("evidence_refs", "reason_evidence_refs")
        for evidence_id in check[key]
    }
    sources = []
    attribution = None
    if evidence_ids:
        ordered_ids = sorted(evidence_ids)
        with engine.connect() as connection:
            for offset in range(0, len(ordered_ids), 500):
                batch = ordered_ids[offset : offset + 500]
                rows = connection.execute(
                    select(
                        evidence_blocks.c.id,
                        evidence_blocks.c.document_version_id,
                        evidence_blocks.c.source_text,
                        evidence_blocks.c.page_idx,
                        evidence_blocks.c.sheet,
                        evidence_blocks.c.cell,
                        document_versions.c.original_name,
                    )
                    .join(
                        document_versions,
                        evidence_blocks.c.document_version_id == document_versions.c.id,
                    )
                    .where(
                        document_versions.c.case_id == result["case_id"],
                        document_versions.c.snapshot_id == result["snapshot_id"],
                        evidence_blocks.c.id.in_(batch),
                    )
                ).all()
                sources.extend(
                    {
                        "evidence_id": row.id,
                        "document_id": row.document_version_id,
                        "document_name": row.original_name,
                        "page": row.page_idx + 1 if row.page_idx is not None else None,
                        "sheet": row.sheet,
                        "cell": row.cell,
                        "source_text": row.source_text,
                    }
                    for row in rows
                )
        if {source["evidence_id"] for source in sources} != evidence_ids:
            raise ApiQueryError("memory_scope")
    if result.get("checklist_version_id"):
        from ..db.dynamic_repository import DynamicRepository

        # Recheck after the separate source read: metadata/Case/corrections may
        # have changed while the download was being built.
        dynamic = DynamicRepository(engine)
        if dynamic.checks(result["run_id"]) != result:
            raise ApiQueryError("run_changed")
        row, _, repo = dynamic.verified(result["run_id"])
        frozen = row.get("item_bindings")
        if frozen is not None:
            from ..db.item_binding_repository import ItemBindingRepository

            version = (ItemBindingRepository(engine).version(result["case_id"], frozen["version_id"])
                       if frozen["version_id"] is not None else {
                           "id": None, "digest": frozen["digest"], "payload": frozen["payload"],
                           "confirmed_by": None, "reason": "提交时没有可用的确认对应", "confirmed_at": None,
                       })
            anchor_ids = {ref for entry in frozen["payload"]["entries"]
                          for anchor in (entry["anchor"], entry.get("sku_anchor") or {})
                          for ref in anchor.get("evidence_refs", [])}
            attribution = {
                "version": version,
                "anchors": [{"evidence_id": b.evidence_id, "document_id": b.document_id,
                             "document_name": repo.source_name(b.document_id), "source_text": b.text,
                             "page": b.page_idx + 1 if b.page_idx is not None else None,
                             "sheet": b.sheet, "cell": b.cell} for b in repo.read(sorted(anchor_ids))],
                "notice": "品目对应锚点仅说明归属；数值结论以核验结果引用的原文为准。",
            }
        verified_sources = {
            b.evidence_id: {
                "source_text": b.text,
                "document_id": b.document_id,
                "document_name": repo.source_name(b.document_id),
                "page": b.page_idx + 1 if b.page_idx is not None else None,
                "sheet": b.sheet,
                "cell": b.cell,
            }
            for b in repo.read(list(evidence_ids))
        }
        if any(
            any(
                source[k] != value
                for k, value in verified_sources[source["evidence_id"]].items()
            )
            for source in sources
        ):
            raise ApiQueryError("dynamic_source_changed")
    if result.get("checklist_version_id") and dynamic.checks(result["run_id"]) != result:
        raise ApiQueryError("run_changed")
    payload = {
        "schema_version": 1,
        "result": result,
        "sources": sorted(sources, key=lambda item: item["evidence_id"]),
    }
    if attribution is not None:
        payload["item_attribution"] = attribution
    return payload


def render_json(payload: dict) -> bytes:
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode(
        "utf-8"
    )


def _text(sheet, row: int, column: int, value: object) -> None:
    """Force every cell that may contain external text to an XLSX string."""
    content = "" if value is None else str(value)
    content = content.replace("\r\n", "\n").replace("\r", "\n")
    content = _ILLEGAL_XML.sub("\ufffd", content)
    if len(content) > _MAX_CELL_CHARS:
        raise ApiQueryError("xlsx_cell_limit")
    cell = sheet.cell(row=row, column=column, value=content)
    cell.data_type = "s"
    cell.number_format = "@"


def render_xlsx(payload: dict) -> bytes:
    workbook = Workbook()
    info = workbook.active
    info.title = "运行信息"
    checks = workbook.create_sheet("核验结果")
    sources = workbook.create_sheet("原文来源")
    result = payload["result"]

    _text(info, 1, 1, "字段")
    _text(info, 1, 2, "值")
    for row, (label, value) in enumerate(
        (
            ("Run ID", result["run_id"]),
            ("Case ID", result["case_id"]),
            ("快照 ID", result["snapshot_id"]),
            ("资料版本", result["revision"]),
            ("运行状态", result["run_status"]),
            ("导出时为当前版本", "是" if result["is_current"] else "否"),
            (
                "说明",
                "每项是否正式以核验结果的正式列为准；历史和部分结果不作为当前正式结论。",
            ),
            ("原文换行", "XLSX 将回车换行统一显示为换行；JSON 保留原始字符。"),
        ),
        start=2,
    ):
        _text(info, row, 1, label)
        _text(info, row, 2, value)
    if result.get("checklist_version_id"):
        _text(info, 10, 1, "清单版本")
        _text(info, 10, 2, result["checklist_version_id"])
    if payload.get("item_attribution") is not None:
        attribution = payload["item_attribution"]
        version = attribution["version"]
        sheet = workbook.create_sheet("品目对应依据")
        metadata = (("对应版本", version["id"]), ("摘要", version["digest"]),
                    ("确认人", version["confirmed_by"]), ("确认时间", version["confirmed_at"]),
                    ("原因", version["reason"]), ("说明", attribution["notice"]))
        for row, (label, value) in enumerate(metadata, 1):
            _text(sheet, row, 1, label)
            _text(sheet, row, 2, value)
        for row, entry in enumerate(version["payload"]["entries"], 8):
            values = (entry["document_id"], entry["item_id"], entry["identifier_class"],
                      entry.get("supplier_sku"), json.dumps(entry["selector"], ensure_ascii=False),
                      "\n".join(entry["anchor"]["evidence_refs"]), entry["anchor"]["quote"],
                      json.dumps(entry.get("sku_anchor"), ensure_ascii=False))
            for column, value in enumerate(values, 1):
                _text(sheet, row, column, value)
        for column, label in enumerate(("文档ID", "品目ID", "标识类型", "供应商SKU", "确认范围", "锚点ID", "锚点摘录", "SKU锚点"), 1):
            _text(sheet, 7, column, label)
        anchor_sheet = workbook.create_sheet("对应锚点原文")
        for column, label in enumerate(("来源ID", "文档ID", "文件名", "页码", "工作表", "单元格", "完整原文"), 1):
            _text(anchor_sheet, 1, column, label)
        for row, anchor in enumerate(attribution["anchors"], 2):
            for column, key in enumerate(("evidence_id", "document_id", "document_name", "page", "sheet", "cell", "source_text"), 1):
                _text(anchor_sheet, row, column, anchor[key])
        for target in (sheet, anchor_sheet):
            target.column_dimensions["A"].width = 31
            target.column_dimensions["G"].width = 50

    headers = (
        "供应商",
        "指标",
        "处理状态",
        "业务结果",
        "原因代码",
        "规则版本",
        "正式",
        "来源 ID",
        "原因来源 ID",
        "事实 ID",
        "核验结果 ID",
    )
    if result.get("checklist_version_id"):
        headers += ("品目 ID", "指标 ID", "清单版本", "品目")
    for column, label in enumerate(headers, start=1):
        _text(checks, 1, column, label)
    for row, check in enumerate(result["checks"], start=2):
        values = (
            check["supplier"],
            check.get("criterion_name") or check["field_key"],
            check["processing_status"],
            check["business_status"],
            check["reason_code"],
            check["rule_version"],
            "是" if check["is_formal"] else "否",
            "\n".join(check["evidence_refs"]),
            "\n".join(check["reason_evidence_refs"]),
            "\n".join(check["fact_ids"]),
            check["check_result_id"],
        )
        if result.get("checklist_version_id"):
            values += (
                check.get("item_id"),
                check.get("criterion_id"),
                result["checklist_version_id"],
                check.get("item_name"),
            )
        for column, value in enumerate(values, start=1):
            _text(checks, row, column, value)

    material_groups = [
        (check, group)
        for check in result["checks"]
        for group in check.get("material_source_groups", [])
    ]
    if material_groups:
        material = workbook.create_sheet("材料引用分组")
        for column, label in enumerate(("核验结果 ID", "指标 ID", "品目 ID", "事实 ID", "材料来源 ID"), 1):
            _text(material, 1, column, label)
        for row, (check, group) in enumerate(material_groups, 2):
            values = (check["check_result_id"], check["criterion_id"], check["item_id"],
                      group["fact_id"], "\n".join(group["evidence_refs"]))
            for column, value in enumerate(values, 1):
                _text(material, row, column, value)
        material.freeze_panes = "A2"
        material.auto_filter.ref = material.dimensions
        for column in ("A", "B", "C", "D", "E"):
            material.column_dimensions[column].width = 31

    tax_groups = [
        (check, group)
        for check in result["checks"]
        for group in check.get("tax_source_groups", [])
    ]
    if tax_groups:
        tax = workbook.create_sheet("税口径引用分组")
        for column, label in enumerate(("核验结果 ID", "指标 ID", "品目 ID", "事实 ID", "税说明来源 ID"), 1):
            _text(tax, 1, column, label)
        for row, (check, group) in enumerate(tax_groups, 2):
            values = (check["check_result_id"], check["criterion_id"], check["item_id"],
                      group["fact_id"], "\n".join(group["evidence_refs"]))
            for column, value in enumerate(values, 1):
                _text(tax, row, column, value)
        tax.freeze_panes = "A2"
        tax.auto_filter.ref = tax.dimensions
        for column in ("A", "B", "C", "D", "E"):
            tax.column_dimensions[column].width = 31

    price_groups = [
        (check, group)
        for check in result["checks"]
        for group in check.get("price_source_groups", [])
    ]
    if price_groups:
        price = workbook.create_sheet("单价引用分组")
        for column, label in enumerate(("核验结果 ID", "指标 ID", "品目 ID", "事实 ID", "单价来源 ID"), 1):
            _text(price, 1, column, label)
        for row, (check, group) in enumerate(price_groups, 2):
            values = (check["check_result_id"], check["criterion_id"], check["item_id"],
                      group["fact_id"], "\n".join(group["evidence_refs"]))
            for column, value in enumerate(values, 1):
                _text(price, row, column, value)
        price.freeze_panes = "A2"
        price.auto_filter.ref = price.dimensions
        for column in ("A", "B", "C", "D", "E"):
            price.column_dimensions[column].width = 31

    source_headers = (
        "来源 ID",
        "文档 ID",
        "原件文件名",
        "PDF 页码",
        "工作表",
        "单元格",
        "原文",
    )
    for column, label in enumerate(source_headers, start=1):
        _text(sources, 1, column, label)
    for row, source in enumerate(payload["sources"], start=2):
        values = (
            source["evidence_id"],
            source["document_id"],
            source["document_name"],
            source["page"],
            source["sheet"],
            source["cell"],
            source["source_text"],
        )
        for column, value in enumerate(values, start=1):
            _text(sources, row, column, value)

    info.column_dimensions["A"].width = 23
    info.column_dimensions["B"].width = 72
    checks.freeze_panes = "A2"
    checks.auto_filter.ref = checks.dimensions
    sources.freeze_panes = "A2"
    sources.auto_filter.ref = sources.dimensions
    for sheet in (checks, sources):
        sheet.column_dimensions["A"].width = 31
        sheet.column_dimensions["B"].width = 31
        sheet.column_dimensions["C"].width = 26
        sheet.column_dimensions["D"].width = 20
        sheet.column_dimensions["G"].width = 45
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
