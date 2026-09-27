"""Turn supported document parser outputs into source-located evidence blocks."""

from __future__ import annotations

import json
import subprocess
import tempfile
import unicodedata
from collections.abc import Sequence
from dataclasses import dataclass, replace
from hashlib import sha256
from pathlib import Path, PurePosixPath
from typing import Literal
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from pypdf import PageObject, PdfReader
from pypdf.errors import PdfReadError
from pypdf.generic import ArrayObject, ContentStream, DictionaryObject, StreamObject

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_SHEETS = 5
MAX_ROWS = 1000
MAX_COLUMNS = 30
MAX_NONEMPTY_CELLS = 20_000
MAX_CELL_CHARACTERS = 2000
XLSX_RELATIONSHIP_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
XLSX_COMMENTS_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
)
PDF_VECTOR_PAINT_OPERATORS = frozenset({
    b"S", b"s", b"f", b"F", b"f*", b"B", b"B*", b"b", b"b*", b"sh",
})


class UnsupportedContentError(ValueError):
    """The document is outside the supported input contract."""


class DocumentParsingError(RuntimeError):
    """A supported document could not be fully parsed."""


@dataclass(frozen=True)
class EvidenceBlock:
    evidence_id: str
    document_id: str
    text: str
    page_idx: int | None = None
    bbox: tuple[float, float, float, float] | None = None
    page_width_pt: float | None = None
    page_height_pt: float | None = None
    sheet: str | None = None
    cell: str | None = None


@dataclass(frozen=True)
class ParseResult:
    document_id: str
    parser_version: str
    status: Literal["complete", "partial"]
    blocks: tuple[EvidenceBlock, ...]
    total_pages: int | None = None
    processed_pages: tuple[int, ...] = ()
    total_sheets: int | None = None
    processed_sheets: tuple[str, ...] = ()
    coverage_issues: tuple[str, ...] = ()


def _evidence_id(document_id: str, locator: str, text: str) -> str:
    return sha256(f"{document_id}\0{locator}\0{text}".encode()).hexdigest()[:24]


def parse_mineru_middle(path: Path, *, document_id: str) -> ParseResult:
    """Read MinerU 4 middle JSON; source PDF coverage is checked separately."""
    source = json.loads(Path(path).read_text(encoding="utf-8"))
    if source.get("schema") != "docvortex.middle" or source.get("schema_version") != "2.0":
        raise UnsupportedContentError("unsupported MinerU middle JSON schema")
    pages = source.get("pages")
    if not isinstance(pages, list):
        raise UnsupportedContentError("MinerU output has no pages")

    metadata = source.get("metadata", {})
    total_pages = metadata.get("document", {}).get("page_count")
    if not isinstance(total_pages, int) or total_pages < 1 or total_pages > 30:
        raise UnsupportedContentError("PDF page count is missing or exceeds limit")
    layout_pages = {
        item["page_idx"]: item
        for item in source.get("extensions", {}).get("docvortex_layout", {}).get("pages", [])
    }
    blocks: list[EvidenceBlock] = []
    processed_pages: list[int] = []
    coverage_issues: list[str] = []
    for page in pages:
        page_idx = page.get("page_idx")
        if not isinstance(page_idx, int) or page_idx < 0 or page_idx >= total_pages:
            raise UnsupportedContentError("invalid PDF page index")
        if page_idx in processed_pages:
            raise UnsupportedContentError("duplicate PDF page index")
        layout = layout_pages.get(page_idx)
        if layout is None or layout.get("width_pt", 0) <= 0 or layout.get("height_pt", 0) <= 0:
            raise UnsupportedContentError("PDF page dimensions are missing")
        processed_pages.append(page_idx)
        page_blocks = page.get("blocks")
        if not isinstance(page_blocks, list):
            raise UnsupportedContentError("PDF page blocks are missing or malformed")
        body_text_found = False
        for block_index, block in enumerate(page_blocks):
            if not isinstance(block, dict):
                raise UnsupportedContentError("PDF block is malformed")
            content = block.get("content", [])
            if not isinstance(content, list) or any(not isinstance(item, dict) for item in content):
                raise UnsupportedContentError("PDF block content is malformed")
            if (block.get("type") not in {"text", "doc_title", "page_number"}
                    or not content or any(item.get("type") != "text" for item in content)):
                coverage_issues.append(f"page:{page_idx}:block:{block_index}:unsupported_content")
            text = " ".join(
                str(item["content"]).strip()
                for item in block.get("content", [])
                if item.get("type") == "text" and item.get("content")
            ).strip()
            if not text:
                continue
            if block.get("type") != "page_number":
                body_text_found = True
            raw_bbox = block.get("bbox")
            if (
                not isinstance(raw_bbox, list)
                or len(raw_bbox) != 4
                or not all(isinstance(value, (int, float)) and 0 <= value <= 1 for value in raw_bbox)
                or raw_bbox[0] > raw_bbox[2]
                or raw_bbox[1] > raw_bbox[3]
            ):
                raise UnsupportedContentError("invalid PDF evidence bbox")
            locator = f"page:{page_idx}:block:{block_index}"
            blocks.append(
                EvidenceBlock(
                    evidence_id=_evidence_id(document_id, locator, text),
                    document_id=document_id,
                    text=text,
                    page_idx=page_idx,
                    bbox=tuple(float(value) for value in raw_bbox),
                    page_width_pt=float(layout["width_pt"]),
                    page_height_pt=float(layout["height_pt"]),
                )
            )
        if not body_text_found:
            coverage_issues.append(f"page:{page_idx}:no_body_text")

    complete = not coverage_issues and source.get("is_full_document") is True and sorted(processed_pages) == list(
        range(total_pages)
    )
    producer = metadata.get("producer", {})
    parser_config = source.get("extensions", {}).get("mineru", {})
    config_hash = sha256(json.dumps(parser_config, sort_keys=True).encode()).hexdigest()[:16]
    parser_version = f"{producer.get('name', 'mineru')}:{producer.get('version', 'unknown')}/middle:2.0/config:{config_hash}"
    return ParseResult(
        document_id=document_id,
        parser_version=parser_version,
        status="complete" if complete else "partial",
        blocks=tuple(blocks),
        total_pages=total_pages,
        processed_pages=tuple(sorted(processed_pages)),
        coverage_issues=tuple(coverage_issues),
    )


def verify_pdf_source_coverage(parsed: ParseResult, source_pdf: Path) -> ParseResult:
    """Mark omitted source text and unverified page content as incomplete."""
    try:
        pdf = PdfReader(source_pdf)
        if len(pdf.pages) != parsed.total_pages:
            raise DocumentParsingError("MinerU page count differs from the source PDF")
        issues = list(parsed.coverage_issues)
        for page_idx, page in enumerate(pdf.pages):
            if len(page.images):
                issues.append(f"page:{page_idx}:embedded_image_unverified")
            annotation_ref = page.get("/Annots")
            if annotation_ref is not None:
                annotations = annotation_ref.get_object()
                if not isinstance(annotations, ArrayObject):
                    issues.append(f"page:{page_idx}:annotations_unverified")
                else:
                    issues.extend(
                        f"page:{page_idx}:annotation:{index}:unverified"
                        for index in range(len(annotations))
                    )
            issues.extend(
                f"page:{page_idx}:{issue}" for issue in sorted(_pdf_graphics_issues(page, pdf))
            )
            try:
                source_text = page.extract_text()
            except (PdfReadError, UnicodeError, ValueError):
                source_text = None
            if not source_text or not source_text.strip():
                issues.append(f"page:{page_idx}:source_text_unavailable")
                continue
            parsed_text = "".join(block.text for block in parsed.blocks if block.page_idx == page_idx)
            normalized_parsed = _normalize_pdf_text(parsed_text)
            used_positions = bytearray(len(normalized_parsed))
            for line_idx, line in enumerate(source_text.splitlines(), start=1):
                normalized_line = _normalize_pdf_text(line)
                if not normalized_line:
                    continue
                match_at = normalized_parsed.find(normalized_line)
                while match_at >= 0 and any(used_positions[match_at:match_at + len(normalized_line)]):
                    match_at = normalized_parsed.find(normalized_line, match_at + 1)
                if match_at < 0:
                    issues.append(f"page:{page_idx}:source_line:{line_idx}:not_extracted")
                else:
                    used_positions[match_at:match_at + len(normalized_line)] = b"\x01" * len(normalized_line)
    except DocumentParsingError:
        raise
    except Exception as exc:
        raise UnsupportedContentError("PDF source coverage could not be checked") from exc
    version = parsed.parser_version
    if "/source-text:1" not in version:
        version += "/source-text:1"
    if "/embedded-image:1" not in version:
        version += "/embedded-image:1"
    if "/vector-graphics:1" not in version:
        version += "/vector-graphics:1"
    if "/annotations:1" not in version:
        version += "/annotations:1"
    issues = list(dict.fromkeys(issues))
    return replace(parsed, parser_version=version,
                   status="partial" if issues else parsed.status,
                   coverage_issues=tuple(issues))


def _normalize_pdf_text(text: str) -> str:
    return "".join(character for character in unicodedata.normalize("NFKC", text)
                   if not character.isspace())


def _pdf_graphics_issues(page: PageObject, pdf: PdfReader) -> set[str]:
    """Inspect painted vectors in the page and forms it actually invokes."""
    content = page.get_contents()
    if content is None:
        return set()
    root_resource_ref = page.get("/Resources")
    root_resources = root_resource_ref.get_object() if root_resource_ref is not None else DictionaryObject()
    if not isinstance(root_resources, DictionaryObject):
        root_resources = DictionaryObject()
    issues: set[str] = set()
    active_forms: set[int] = set()

    def scan(stream: StreamObject | ArrayObject, resources: DictionaryObject, depth: int) -> None:
        if depth > 20:
            issues.add("xobject_unverified")
            return
        for operands, operator in ContentStream(stream, pdf).operations:
            if operator in PDF_VECTOR_PAINT_OPERATORS:
                issues.add("vector_graphics_unverified")
            if operator != b"Do":
                continue
            xobject_ref = resources.get("/XObject")
            xobjects = xobject_ref.get_object() if xobject_ref else None
            if not operands or not isinstance(xobjects, DictionaryObject):
                issues.add("xobject_unverified")
                continue
            target_ref = xobjects.get(operands[0])
            target = target_ref.get_object() if target_ref else None
            if not isinstance(target, StreamObject):
                issues.add("xobject_unverified")
                continue
            subtype = target.get("/Subtype")
            if subtype == "/Image":
                continue  # page.images checks images, including nested and inline images.
            if subtype != "/Form":
                issues.add("xobject_unverified")
                continue
            form_id = id(target)
            if form_id in active_forms:
                issues.add("xobject_unverified")
                continue
            form_resource_ref = target.get("/Resources")
            form_resources = form_resource_ref.get_object() if form_resource_ref is not None else resources
            if not isinstance(form_resources, DictionaryObject):
                issues.add("xobject_unverified")
                continue
            active_forms.add(form_id)
            try:
                scan(target, form_resources, depth + 1)
            finally:
                active_forms.remove(form_id)

    scan(content, root_resources, 0)
    return issues


def parse_pdf(
    path: Path,
    *,
    document_id: str,
    mineru_command: Sequence[str],
    timeout_seconds: int = 180,
) -> ParseResult:
    """Run local MinerU Basic on a text PDF and convert its middle JSON output."""
    source = Path(path)
    if source.suffix.lower() != ".pdf" or source.stat().st_size > MAX_FILE_BYTES:
        raise UnsupportedContentError("expected a PDF file of at most 10 MiB")
    if source.read_bytes()[:5] != b"%PDF-":
        raise UnsupportedContentError("file content is not PDF")
    try:
        pdf = PdfReader(source)
        page_count = len(pdf.pages)
        if page_count < 1 or page_count > 30:
            raise UnsupportedContentError("PDF page count exceeds 30")
        if any(not page.extract_text().strip() for page in pdf.pages):
            raise UnsupportedContentError("scanned or empty PDF pages are unsupported")
    except UnsupportedContentError:
        raise
    except Exception as exc:
        raise UnsupportedContentError("PDF structure could not be read") from exc

    with tempfile.TemporaryDirectory(prefix="docucheck-mineru-") as temp_dir:
        output_path = Path(temp_dir) / "document.middle.json"
        command = [
            *mineru_command,
            "parse",
            str(source),
            "-o",
            str(output_path),
            "--format",
            "middle_json",
            "--tier",
            "basic",
            "--ocr-mode",
            "txt",
        ]
        try:
            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=timeout_seconds,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise DocumentParsingError("MinerU timed out") from exc
        except OSError as exc:
            raise DocumentParsingError("MinerU could not be started") from exc
        if completed.returncode != 0 or not output_path.is_file():
            raise DocumentParsingError(f"MinerU failed with exit code {completed.returncode}")
        result = parse_mineru_middle(output_path, document_id=document_id)
    return verify_pdf_source_coverage(result, source)


def parse_xlsx(path: Path, *, document_id: str) -> ParseResult:
    """Read literal XLSX cells, rejecting formulas and inputs beyond phase-one limits."""
    source = Path(path)
    if source.suffix.lower() != ".xlsx" or source.stat().st_size > MAX_FILE_BYTES:
        raise UnsupportedContentError("expected an XLSX file of at most 10 MiB")
    coverage_issues: list[str] = []
    try:
        with ZipFile(source) as archive:
            uncompressed_bytes = sum(item.file_size for item in archive.infolist())
            if uncompressed_bytes > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise UnsupportedContentError("XLSX uncompressed size exceeds 50 MiB")
    except BadZipFile as exc:
        raise UnsupportedContentError("XLSX is not a valid ZIP container") from exc

    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise UnsupportedContentError("XLSX sheet count exceeds 5")
        blocks: list[EvidenceBlock] = []
        nonempty = 0
        for sheet in workbook:
            if (sheet.max_row is not None and sheet.max_row > MAX_ROWS) or (
                sheet.max_column is not None and sheet.max_column > MAX_COLUMNS
            ):
                raise UnsupportedContentError("XLSX sheet dimensions exceed 1000 rows or 30 columns")
            declared_bounds = (sheet.min_row, sheet.min_column, sheet.max_row, sheet.max_column)
            dimension_mismatch = False
            row_index = 0
            column_index = 0
            try:
                # Read the worksheet part resolved by openpyxl's workbook relationships.
                worksheet_part = PurePosixPath(sheet._worksheet_path.lstrip("/"))
                with workbook._archive.open(str(worksheet_part)) as stream:
                    for event, element in iterparse(stream, events=("start", "end")):
                        if event == "end":
                            element.clear()
                            continue
                        tag = element.tag.rsplit("}", 1)[-1]
                        if tag in {"drawing", "legacyDrawing", "legacyDrawingHF"}:
                            drawing_ref = element.attrib.get(XLSX_RELATIONSHIP_ID, "unknown")
                            if len(drawing_ref) > 80:
                                drawing_ref = "unknown"
                            kind = "drawing" if tag == "drawing" else "legacy_drawing"
                            coverage_issues.append(
                                f"sheet:{sheet.title}:{kind}_unverified:{drawing_ref}"
                            )
                        elif tag == "mergeCell":
                            coverage_issues.append(
                                f"sheet:{sheet.title}:merged_cell:{element.attrib.get('ref', 'unknown')}"
                            )
                        elif tag == "row":
                            row_ref = element.get("r")
                            try:
                                row_index = int(row_ref) if row_ref is not None else row_index + 1
                            except ValueError as exc:
                                raise UnsupportedContentError("XLSX row coordinate is invalid") from exc
                            if row_index < 1 or row_index > MAX_ROWS:
                                raise UnsupportedContentError(
                                    "XLSX sheet dimensions exceed 1000 rows or 30 columns"
                                )
                            column_index = 0
                        elif tag == "c":
                            cell_ref = element.get("r")
                            if cell_ref:
                                try:
                                    cell_row, cell_column = coordinate_to_tuple(cell_ref)
                                except ValueError as exc:
                                    raise UnsupportedContentError(
                                        "XLSX cell coordinate is invalid"
                                    ) from exc
                                column_index = cell_column
                            else:
                                column_index += 1
                                cell_row, cell_column = row_index, column_index
                                cell_ref = f"{get_column_letter(cell_column)}{cell_row}"
                            if cell_row != row_index or cell_column < 1:
                                raise UnsupportedContentError("XLSX cell coordinate differs from row")
                            if cell_row > MAX_ROWS or cell_column > MAX_COLUMNS:
                                raise UnsupportedContentError(
                                    "XLSX sheet dimensions exceed 1000 rows or 30 columns"
                                )
                            min_row, min_col, max_row, max_col = declared_bounds
                            if not dimension_mismatch and (
                                (min_row is not None and cell_row < min_row)
                                or (min_col is not None and cell_column < min_col)
                                or (max_row is not None and cell_row > max_row)
                                or (max_col is not None and cell_column > max_col)
                            ):
                                coverage_issues.append(
                                    f"sheet:{sheet.title}:dimension_mismatch:{cell_ref}"
                                )
                                dimension_mismatch = True
                relationship_part = str(
                    worksheet_part.parent / "_rels" / f"{worksheet_part.name}.rels"
                )
                try:
                    relationship_info = workbook._archive.getinfo(relationship_part)
                except KeyError:
                    pass
                else:
                    with workbook._archive.open(relationship_info) as stream:
                        for event, element in iterparse(stream, events=("start", "end")):
                            if event == "end":
                                element.clear()
                                continue
                            if (element.tag.rsplit("}", 1)[-1] == "Relationship"
                                    and element.attrib.get("Type") == XLSX_COMMENTS_RELATIONSHIP):
                                comment_ref = element.attrib.get("Id", "unknown")
                                if len(comment_ref) > 80:
                                    comment_ref = "unknown"
                                coverage_issues.append(
                                    f"sheet:{sheet.title}:comment_unverified:{comment_ref}"
                                )
            except (DefusedXmlException, ParseError) as exc:
                raise UnsupportedContentError("XLSX worksheet or relationship XML is malformed") from exc
            if dimension_mismatch:
                sheet.reset_dimensions()
            for row in sheet.iter_rows(min_row=1, min_col=1):
                for cell in row:
                    if cell.value is None:
                        continue
                    if cell.row > MAX_ROWS or cell.column > MAX_COLUMNS:
                        raise UnsupportedContentError(
                            "XLSX sheet dimensions exceed 1000 rows or 30 columns"
                        )
                    nonempty += 1
                    if nonempty > MAX_NONEMPTY_CELLS:
                        raise UnsupportedContentError("XLSX nonempty cell count exceeds 20000")
                    if cell.data_type == "f":
                        raise UnsupportedContentError("formula cells are unsupported")
                    text = str(cell.value)
                    if len(text) > MAX_CELL_CHARACTERS:
                        raise UnsupportedContentError("XLSX cell text exceeds 2000 characters")
                    locator = f"sheet:{sheet.title}:cell:{cell.coordinate}"
                    blocks.append(
                        EvidenceBlock(
                            evidence_id=_evidence_id(document_id, locator, text),
                            document_id=document_id,
                            text=text,
                            sheet=sheet.title,
                            cell=cell.coordinate,
                        )
                    )
        return ParseResult(
            document_id=document_id,
            parser_version=(
                "openpyxl:3.1.5/merged-coverage:1/dimension:1/drawing:1"
                "/legacy-drawing:1/comments-rel:1"
            ),
            status="partial" if coverage_issues else "complete",
            blocks=tuple(blocks),
            total_sheets=len(workbook.sheetnames),
            processed_sheets=tuple(workbook.sheetnames),
            coverage_issues=tuple(coverage_issues),
        )
    finally:
        workbook.close()
