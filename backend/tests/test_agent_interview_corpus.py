"""The public synthetic evaluation corpus is isolated, bounded and tamper evident."""

import hashlib
import importlib.util
import json
from pathlib import Path

import pytest
from pydantic import ValidationError


def api():
    assert importlib.util.find_spec("evals.agent_interview"), "agent interview corpus missing"
    from evals.agent_interview import corpus
    return corpus


def test_dataset_has_frozen_stratification_and_disjoint_groups():
    dataset = api().load_dataset()
    assert len(dataset.tasks) == 40
    assert len(dataset.repeat_task_ids) == 6
    development = [t for t in dataset.tasks if t.split == "development"]
    measurement = [t for t in dataset.tasks if t.split == "measurement"]
    assert len(development) == 10 and len(measurement) == 30
    assert not {t.group_id for t in development} & {t.group_id for t in measurement}
    assert len({t.prompt for t in dataset.tasks}) == 40
    for tasks, counts in [(development, [2, 3, 3, 1, 1]), (measurement, [6, 9, 7, 3, 5])]:
        assert [sum(t.category == c for t in tasks) for c in
                ("intent", "rag", "review", "memory", "safety")] == counts
    rag = [t for t in dataset.tasks if t.category == "rag"]
    assert sum(t.expected.answerable is True for t in rag) == 8
    assert sum(t.expected.answerable is False for t in rag) == 4
    assert {t.category for t in measurement if t.task_id in dataset.repeat_task_ids} == {
        "intent", "rag", "review", "memory", "safety"}


def test_contract_rejects_bad_counts_scope_locations_and_extra_fields():
    from evals.agent_interview.contracts import Dataset
    raw = api().load_dataset().model_dump(mode="json")
    changes = [lambda d: d["tasks"].pop(),
               lambda d: d["sources"][0].update(path="../escape.pdf"),
               lambda d: d["anchors"][0].update(page=0),
               lambda d: d["tasks"][0].update(unexpected="label"),
               lambda d: d["tasks"][0].update(group_id=d["tasks"][-1]["group_id"])]
    for change in changes:
        value = json.loads(json.dumps(raw))
        change(value)
        with pytest.raises(ValidationError):
            Dataset.model_validate(value)


def test_build_roundtrip_originals_gold_separation_and_no_overwrite(tmp_path):
    corpus = api()
    root = tmp_path / "corpus"
    manifest = corpus.build_corpus(root)
    assert corpus.validate_corpus(root) == manifest
    assert {r.role for r in manifest.files} == {"input", "gold"}
    assert any(r.path.endswith(".pdf") for r in manifest.files)
    assert any(r.path.endswith(".xlsx") for r in manifest.files)
    agent_inputs = json.loads((root / "inputs/scenarios.json").read_text(encoding="utf-8"))
    encoded = json.dumps(agent_inputs, ensure_ascii=False)
    assert "necessary_anchor_groups" not in encoded and "fact_claims" not in encoded
    assert not any(t.task_id in encoded for t in corpus.load_dataset().tasks)
    with pytest.raises(FileExistsError):
        corpus.build_corpus(root)
    record = next(r for r in manifest.files if r.role == "input" and r.path.endswith(".pdf"))
    (root / record.path).write_bytes(b"modified")
    with pytest.raises(ValueError, match="hash"):
        corpus.validate_corpus(root)


def test_validate_checks_anchor_text_and_manifest_completeness(tmp_path):
    corpus = api()
    root = tmp_path / "corpus"
    corpus.build_corpus(root)
    manifest_path = root / "manifest.json"
    raw = json.loads(manifest_path.read_text(encoding="utf-8"))
    raw["files"].pop()
    manifest_path.write_text(json.dumps(raw), encoding="utf-8")
    with pytest.raises(ValueError, match="manifest"):
        corpus.validate_corpus(root)


def test_no_presence_true_or_incomplete_coverage_formal_pass():
    dataset = api().load_dataset()
    reviews = [t for t in dataset.tasks if t.category == "review"]
    assert any(t.expected.outcome == "manual" for t in reviews)
    assert any(t.expected.outcome == "partial" for t in reviews)
    for task in reviews:
        if any(c.kind == "presence" and c.value is True for c in task.expected.fact_claims):
            assert task.expected.outcome == "manual"
        if "coverage_incomplete" in task.expected.required_behaviors:
            assert task.expected.outcome == "partial"


def test_setup_sources_are_supplier_owned_and_product_checklists_are_valid():
    from backend.app.criteria.models import ChecklistDefinition
    dataset = api().load_dataset()
    sources = {s.source_id: s for s in dataset.sources}
    for setup in dataset.setups:
        for binding in setup.bindings:
            assert sources[binding.source_id].supplier_id == binding.supplier_id
        if setup.confirmed_rules:
            ChecklistDefinition.model_validate({
                "items": [{"item_id": i.item_id, "name": i.name, "specification": i.sku}
                          for i in setup.items], "criteria": setup.confirmed_rules})
        else:
            assert setup.state["confirmed"] is False
    for task in dataset.tasks:
        if task.category == "review":
            assert "review_count" not in {a.key for a in task.expected.immutable_state_assertions}


def test_contract_blocks_unowned_sources_bad_bindings_and_untyped_assertions():
    from evals.agent_interview.contracts import Dataset
    raw = api().load_dataset().model_dump(mode="json")
    mutations = [lambda d: d["setups"][0]["bindings"][0].update(pages=[999]),
                 lambda d: d["setups"][0]["bindings"][0].update(supplier_id="b"),
                 lambda d: d["tasks"][0]["expected"]["immutable_state_assertions"][0].update(
                     key="free_text_assertion"),
                 lambda d: d["tasks"][0]["expected"]["immutable_state_assertions"][0].update(
                     expected={"looks": "unchanged"}),
                 lambda d: d["sources"][0].update(path="inputs/dphone/C:evil.pdf")]
    for mutate in mutations:
        value = json.loads(json.dumps(raw))
        mutate(value)
        with pytest.raises(ValidationError):
            Dataset.model_validate(value)


def test_validator_rejects_rehashed_original_with_bad_anchor(tmp_path):
    corpus = api()
    root = tmp_path / "corpus"
    corpus.build_corpus(root)
    manifest_path = root / "manifest.json"
    raw = json.loads(manifest_path.read_text(encoding="utf-8"))
    record = next(r for r in raw["files"] if r["path"].endswith("comparison.xlsx"))
    from openpyxl import load_workbook
    workbook = load_workbook(root / record["path"])
    workbook.active["C2"] = "9999"
    workbook.save(root / record["path"])
    workbook.close()
    payload = (root / record["path"]).read_bytes()
    record.update(sha256=hashlib.sha256(payload).hexdigest(), size=len(payload))
    manifest_path.write_text(json.dumps(raw), encoding="utf-8")
    with pytest.raises(ValueError, match="anchor"):
        corpus.validate_corpus(root)


def test_review_gold_is_full_supplier_criterion_matrix_with_six_rule_types():
    dataset = api().load_dataset()
    setups = {s.setup_id: s for s in dataset.setups}
    rule_types = set()
    for task in dataset.tasks:
        if task.category != "review":
            continue
        setup = setups[task.setup_id]
        rule_types.update(c["rule"]["type"] for c in setup.confirmed_rules)
        required = {(s.supplier_id, c["item_id"], c["criterion_id"])
                    for s in setup.suppliers for c in setup.confirmed_rules}
        checks = task.expected.business_checks
        assert {(c.supplier_id, c.item_id, c.criterion_id) for c in checks} == required
        assert task.expected.extra_result_policy == "reject_unexpected_checks"
        assert all(c.necessary_anchor_groups and c.reference_answer for c in checks)
    assert rule_types == {"numeric", "money", "enum", "presence", "duration", "manual"}


def test_replacement_rag_has_distinct_reasoning_and_real_long_cell_tail():
    dataset = api().load_dataset()
    terminal = next(t for t in dataset.tasks if t.task_id == "measure_rag_terminal")
    assert "低温" in terminal.prompt
    assert len(terminal.expected.necessary_anchor_groups) >= 2
    task = next(t for t in dataset.tasks if t.task_id == "measure_rag_table_scope")
    assert "尾部" in task.prompt
    anchors = {a.anchor_id: a for a in dataset.anchors}
    used = [anchors[a] for group in task.expected.necessary_anchor_groups for a in group]
    assert any(a.cell == "H3" and 1500 < len(a.text) <= 2000
               and a.text.endswith("随板附赠磁扣12枚。") for a in used)
    assert all("安装顺序" not in a.text[:1500] for a in used if a.cell == "H3")


@pytest.mark.parametrize("change", ["extra_cell", "extra_sheet", "extra_pdf_line"])
def test_validator_rejects_rehashed_additional_original_content(tmp_path, change):
    corpus = api()
    root = tmp_path / "corpus"
    corpus.build_corpus(root)
    manifest_path = root / "manifest.json"
    raw = json.loads(manifest_path.read_text(encoding="utf-8"))
    record = next(r for r in raw["files"]
                  if r["path"].endswith("supplier_a.pdf" if change == "extra_pdf_line"
                                       else "comparison.xlsx"))
    path = root / record["path"]
    if change == "extra_pdf_line":
        dataset = corpus.load_dataset()
        source = next(s for s in dataset.sources if s.path == record["path"])
        pages = [list(page) for page in source.pages]
        pages[0].append("未经协议的新业务说明：单价全部改为1元。")
        corpus._pdf(path, pages)
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(path)
        if change == "extra_cell":
            workbook.active["Z9"] = "未经协议的额外价格"
        else:
            workbook.create_sheet("额外业务说明")["A1"] = "隐藏报价"
        workbook.save(path)
        workbook.close()
    payload = path.read_bytes()
    record.update(sha256=hashlib.sha256(payload).hexdigest(), size=len(payload))
    manifest_path.write_text(json.dumps(raw), encoding="utf-8")
    with pytest.raises(ValueError, match="content|sheet|matrix"):
        corpus.validate_corpus(root)


def test_review_gold_preserves_exact_conflict_insufficient_and_manual_states():
    dataset = api().load_dataset()
    tasks = {t.task_id: t for t in dataset.tasks}
    tax = tasks["measure_review_tax_conflict"].expected.business_checks
    coverage = tasks["measure_review_coverage"].expected.business_checks
    assert next(c for c in tax if c.supplier_id == "a").business_status == "conflict"
    assert next(c for c in tax if c.supplier_id == "b").business_status == "manual"
    assert next(c for c in coverage if c.supplier_id == "b").business_status == "insufficient_evidence"
    assert next(c for c in coverage if c.supplier_id == "a").business_status is None


def test_public_scenarios_filter_adapter_only_answer_labels_by_allowlist():
    corpus = api()
    dataset = corpus.load_dataset().model_copy(deep=True)
    sentinels = {
        "memory_target_message_id": "GOLD_SENTINEL_MEMORY_TARGET",
        "automatic_resend_allowed": "GOLD_SENTINEL_RESEND_ANSWER",
        "evaluation_target": "GOLD_SENTINEL_EVALUATION_TARGET",
        "history_clock": "GOLD_SENTINEL_HISTORY_CLOCK",
        "fixture_confirmation": "GOLD_SENTINEL_FIXTURE_CONFIRMATION",
        "unrecognized_future_label": "GOLD_SENTINEL_FUTURE_LABEL",
    }
    dataset.setups[0].context.update(sentinels)
    scenarios, _ = corpus._agent_scenarios(dataset)
    encoded = json.dumps(scenarios, ensure_ascii=False)
    for key, value in sentinels.items():
        assert key not in encoded
        assert value not in encoded
    assert "on_unavailable" not in encoded
    assert "preconditions" not in encoded
    assert "draft_rules" in encoded and "actor_role" in encoded


def test_unconfirmed_draft_has_no_saved_checklist_or_binding_version():
    dataset = api().load_dataset()
    task = next(t for t in dataset.tasks if t.task_id == "measure_intent_draft")
    setup = next(s for s in dataset.setups if s.setup_id == task.setup_id)
    assert setup.state["confirmed"] is False
    for key in ("checklist_version", "bindings_version"):
        assert setup.state[key] == 0
        assert next(a.expected for a in task.expected.immutable_state_assertions if a.key == key) == 0


def test_published_originals_match_frozen_corpus():
    root = Path(__file__).resolve().parents[2] / 'evals/agent_interview/fixtures'
    manifest = api().validate_corpus(root)
    assert manifest.dataset_sha256 == hashlib.sha256(api().DATASET_PATH.read_bytes()).hexdigest()
    assert len([f for f in manifest.files if Path(f.path).suffix in {'.pdf', '.xlsx'}]) == 40
    assert any(f.path == 'inputs/scenarios.json' for f in manifest.files)


def test_published_xlsx_inputs_are_within_actual_product_parser_contract():
    from backend.app.parsing import parse_xlsx

    root = Path(__file__).resolve().parents[2] / 'evals/agent_interview/fixtures'
    dataset = api().load_dataset()
    for source in dataset.sources:
        if source.format == 'xlsx':
            result = parse_xlsx(root / source.path, document_id=source.source_id)
            assert result.status == 'complete', source.source_id


def test_source_contract_rejects_xlsx_cell_beyond_supported_input_limit():
    source = next(s for s in api().load_dataset().sources if s.format == 'xlsx')
    value = source.model_dump(mode='json')
    sheet = next(iter(value['sheets']))
    value['sheets'][sheet][0][0] = 'x' * 2001
    with pytest.raises(ValueError, match='xlsx_cell_exceeds_supported_contract'):
        type(source).model_validate(value)


@pytest.mark.parametrize("role", ["quote", "spec", "requirement"])
def test_contract_rejects_sources_that_replace_the_same_product_slot(role):
    from evals.agent_interview.contracts import Dataset

    raw = api().load_dataset().model_dump(mode="json")
    # Start from a slot-safe protocol, including when checking the old broken fixture.
    for source in raw["sources"]:
        source["role"] = "spec" if source["format"] == "pdf" else "quote"
    setup = raw["setups"][0]
    sources = {source["source_id"]: source for source in raw["sources"]}
    first = sources[setup["source_ids"][0]]
    other = next(sources[key] for key in setup["source_ids"]
                 if key != first["source_id"]
                 and sources[key]["supplier_id"] == first["supplier_id"])
    if role == "requirement":
        # Requirements occupy the one case-wide (None, requirement) slot.
        for index, original in enumerate((first, other)):
            copy = dict(original, source_id=f"extra_requirement_{index}",
                        path=f"inputs/{setup['group_id']}/extra_{index}.{original['format']}",
                        supplier_id=None, role=role)
            raw["sources"].append(copy)
            setup["source_ids"].append(copy["source_id"])
    else:
        first["role"] = other["role"] = role
    with pytest.raises(ValidationError, match="duplicate_setup_document_slot"):
        Dataset.model_validate(raw)


@pytest.mark.parametrize("role", ["quote", "spec", "requirement"])
def test_actual_document_repository_replaces_same_slot_across_file_formats(tmp_path, role):
    from datetime import UTC, datetime

    from sqlalchemy import create_engine, insert

    from backend.app.db.document_repository import DocumentRepository, UploadCommand
    from backend.app.db.models import cases, metadata

    engine = create_engine(f"sqlite:///{tmp_path / 'slots.db'}")
    try:
        metadata.create_all(engine)
        with engine.begin() as connection:
            connection.execute(insert(cases).values(
                id="slot-case", revision=0, created_at=datetime.now(UTC),
                updated_at=datetime.now(UTC)))
        repository = DocumentRepository(engine)
        for revision, extension in enumerate(("pdf", "xlsx")):
            repository.upload(UploadCommand(
                case_id="slot-case", expected_revision=revision, role=role,
                supplier_handle=None if role == "requirement" else "supplier-a",
                supplier_name=None if role == "requirement" else "供应商A",
                idempotency_key=f"upload-{revision}", document_id=f"document-{revision}",
                original_name=f"original-{revision}.{extension}",
                storage_key=f"test/{revision}.{extension}", source_sha256=str(revision) * 64,
                media_type="application/pdf" if extension == "pdf" else "application/octet-stream",
                pdf_pages=1 if extension == "pdf" else 0))
        current = repository.list_current("slot-case")
        assert current["revision"] == 2
        assert [document["id"] for document in current["documents"]] == ["document-1"]
        assert current["documents"][0]["model_name"] is None
    finally:
        engine.dispose()


def test_published_sources_explicitly_keep_pdf_specs_and_xlsx_quotes_in_distinct_slots():
    dataset = api().load_dataset()
    assert dataset.revision == 6
    assert sum(s.format == "pdf" and s.role == "spec" for s in dataset.sources) == 20
    assert sum(s.format == "xlsx" and s.role == "quote" for s in dataset.sources) == 20
    sources = {source.source_id: source for source in dataset.sources}
    for setup in dataset.setups:
        slots = [(sources[key].supplier_id, sources[key].role) for key in setup.source_ids]
        assert len(slots) == len(set(slots)) == 4


def test_source_requirement_without_supplier_is_one_supported_document_slot():
    from evals.agent_interview.contracts import Source

    source = next(s for s in api().load_dataset().sources if s.format == "pdf")
    raw = source.model_dump(mode="json")
    raw.update(role="requirement", supplier_id=None)
    assert Source.model_validate(raw).supplier_id is None
    for role, supplier in (("quote", None), ("spec", None), ("requirement", "a")):
        raw.update(role=role, supplier_id=supplier)
        with pytest.raises(ValidationError, match="source_role_supplier_mismatch"):
            Source.model_validate(raw)


@pytest.mark.parametrize("remove_all", [False, True])
def test_confirmed_setup_rejects_items_without_an_explicit_rule(remove_all):
    from evals.agent_interview.contracts import Setup

    raw = api().load_dataset().setups[0].model_dump(mode="json")
    if remove_all:
        raw["confirmed_rules"] = []
    else:
        # A saved first-item rule cannot retain another unreferenced confirmed item.
        raw["confirmed_rules"] = [c for c in raw["confirmed_rules"]
                                  if c["item_id"] == raw["items"][0]["item_id"]]
    with pytest.raises(ValidationError, match="confirmed_setup_has_unreferenced_items"):
        Setup.model_validate(raw)


def test_v6_confirms_every_declared_item_and_exposes_all_secondary_sku_cells():
    dataset = api().load_dataset()
    assert dataset.revision == 6
    setups = {s.setup_id: s for s in dataset.setups}
    for setup in dataset.setups:
        if setup.state["confirmed"] is True:
            assert {i.item_id for i in setup.items} == {
                c["item_id"] for c in setup.confirmed_rules}
            secondary = setup.items[1]
            criterion = next(c for c in setup.confirmed_rules
                             if c["criterion_id"] == "secondary_sku")
            assert criterion["item_id"] == secondary.item_id
            assert criterion["rule"] == {"type": "enum", "operator": "eq",
                                          "values": [secondary.sku]}
            assert criterion["source"]["kind"] == "user_defined"
        else:
            assert setup.confirmed_rules == []
    reviews = [t for t in dataset.tasks if t.category == "review"]
    assert sum(len(t.expected.business_checks) for t in reviews) == 44
    anchors = {a.anchor_id: a for a in dataset.anchors}
    sources = {s.source_id: s for s in dataset.sources}
    for task in reviews:
        setup = setups[task.setup_id]
        secondary = setup.items[1]
        checks = [c for c in task.expected.business_checks if c.criterion_id == "secondary_sku"]
        assert len(checks) == len(setup.suppliers) == 2
        for check in checks:
            assert check.item_id == secondary.item_id
            assert (check.business_status, check.processing_status, check.is_formal) == (
                "satisfied", "complete", task.expected.outcome != "partial")
            for group in check.necessary_anchor_groups:
                for key in group:
                    anchor = anchors[key]
                    source = sources[anchor.source_id]
                    assert source.supplier_id == check.supplier_id
                    assert source.group_id == task.group_id
                    assert secondary.sku in anchor.text


def test_partial_review_rejects_formal_flags_on_other_complete_cells():
    from evals.agent_interview.contracts import ExpectedOutcome

    raw = next(t for t in api().load_dataset().tasks
               if t.task_id == "measure_review_coverage").expected.model_dump(mode="json")
    complete = next(c for c in raw["business_checks"] if c["processing_status"] == "complete")
    complete["is_formal"] = True
    with pytest.raises(ValidationError, match="partial_run_not_formal"):
        ExpectedOutcome.model_validate(raw)


@pytest.mark.parametrize("formal", [False, True])
def test_partial_run_with_all_complete_cells_still_cannot_be_formal(formal):
    from evals.agent_interview.contracts import ExpectedOutcome

    raw = next(t for t in api().load_dataset().tasks
               if t.task_id == "measure_review_coverage").expected.model_dump(mode="json")
    for check in raw["business_checks"]:
        check["processing_status"] = "complete"
        check["is_formal"] = formal
    if formal:
        with pytest.raises(ValidationError, match="partial_run_not_formal"):
            ExpectedOutcome.model_validate(raw)
    else:
        assert ExpectedOutcome.model_validate(raw).outcome == "partial"
