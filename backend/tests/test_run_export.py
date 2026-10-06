"""Run exports retain the checked snapshot contract and literal source text."""

import io
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from backend.app.agent.session_memory import _repository_fingerprint
from backend.app.db.agent_snapshot import load_agent_snapshot
from backend.app.db.models import agent_runs, evidence_blocks
from backend.app.rules.memory import RULE_VERSION
from backend.tests.test_m4_workbench_api import _fixture


def _run_with_result(tmp_path, monkeypatch):
    engine, client, alice, bob, case_id, receipt, _ = _fixture(tmp_path, monkeypatch)
    conversation = client.post(f"/api/v1/cases/{case_id}/conversations",
                               headers=alice).json()
    submitted = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages",
        headers={**alice, "Idempotency-Key": "export-question"},
        json={"content": "审查内存", "mode": "standard", "expected_revision": 1},
    ).json()
    run_id = submitted["run_id"]
    with engine.connect() as connection:
        evidence_id = connection.scalar(select(evidence_blocks.c.id).where(
            evidence_blocks.c.document_version_id == receipt["document_id"],
            evidence_blocks.c.cell == "B6"))
    repo = load_agent_snapshot(engine, case_id, receipt["snapshot_id"])
    memory = {
        "version": 2, "snapshot_id": receipt["snapshot_id"],
        "repository_fingerprint": _repository_fingerprint(repo),
        "fact_sets_available": False, "fact_sets": [],
        "check_results": [{
            "check_result_id": "checked-memory", "snapshot_id": receipt["snapshot_id"],
            "supplier": "vendor-a", "field_key": "memory",
            "processing_status": "complete", "business_status": "满足",
            "reason_code": "threshold_met", "evidence_refs": [evidence_id],
            "reason_evidence_refs": [evidence_id], "fact_ids": [],
            "rule_version": RULE_VERSION,
        }],
    }
    with engine.begin() as connection:
        connection.execute(update(agent_runs).where(agent_runs.c.id == run_id).values(
            status="partial", memory_json=memory))
    return engine, client, alice, bob, case_id, receipt, run_id, evidence_id


def test_export_json_reuses_checked_result_and_scoped_source(tmp_path, monkeypatch):
    engine, client, alice, bob, case_id, receipt, run_id, evidence_id = (
        _run_with_result(tmp_path, monkeypatch))
    url = f"/api/v1/agent-runs/{run_id}/export?format=json"
    checks = client.get(f"/api/v1/agent-runs/{run_id}/checks", headers=alice).json()
    assert client.get(url, headers=bob).status_code == 404
    exported = client.get(url, headers=alice)
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-disposition"].startswith("attachment;")
    assert exported.headers["x-content-type-options"] == "nosniff"
    payload = exported.json()
    assert payload["schema_version"] == 1
    assert payload["result"] == checks
    assert payload["result"]["checks"][0]["is_formal"] is False
    assert payload["sources"] == [{
        "evidence_id": evidence_id, "document_id": receipt["document_id"],
        "document_name": "vendor-a-quote.xlsx", "sheet": "报价", "cell": "B6",
        "page": None, "source_text": "已装内存 16 GB",
    }]
    assert payload["result"]["case_id"] == case_id
    assert client.get(url.replace("json", "xml"), headers=alice).status_code == 422
    engine.dispose()


def test_xlsx_export_matches_json_and_contains_no_formula(tmp_path, monkeypatch):
    engine, client, alice, _, _, _, run_id, evidence_id = (
        _run_with_result(tmp_path, monkeypatch))
    base = f"/api/v1/agent-runs/{run_id}/export"
    payload = client.get(f"{base}?format=json", headers=alice).json()
    response = client.get(f"{base}?format=xlsx", headers=alice)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["运行信息", "核验结果", "原文来源"]
    assert workbook["运行信息"]["B2"].value == run_id
    assert workbook["核验结果"]["A2"].value == "vendor-a"
    assert workbook["核验结果"]["D2"].value == "满足"
    assert workbook["核验结果"]["G2"].value == "否"
    assert workbook["核验结果"]["H2"].value == evidence_id
    assert workbook["原文来源"]["A2"].value == payload["sources"][0]["evidence_id"]
    assert workbook["原文来源"]["G2"].value == "已装内存 16 GB"
    with ZipFile(io.BytesIO(response.content)) as archive:
        sheets = [archive.read(name) for name in archive.namelist()
                  if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")]
    assert all(b"<f" not in sheet for sheet in sheets)
    engine.dispose()


def test_export_rejects_run_without_structured_checks(tmp_path, monkeypatch):
    engine, client, alice, _, case_id, receipt, _, _ = (
        _run_with_result(tmp_path, monkeypatch))
    conversation = client.post(f"/api/v1/cases/{case_id}/conversations",
                               headers=alice).json()
    submitted = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages",
        headers={**alice, "Idempotency-Key": "empty-export"},
        json={"content": "尚未处理的问题", "mode": "standard",
              "expected_revision": receipt["revision"]},
    ).json()
    response = client.get(
        f"/api/v1/agent-runs/{submitted['run_id']}/export?format=json",
        headers=alice)
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "checks_unavailable"
    engine.dispose()


@pytest.mark.parametrize("attack", (
    '=HYPERLINK("https://example.invalid/?x=1", "click")',
    "+SUM(1,2)", "-SUM(1,2)", "@SUM(1,2)",
    "\t=SUM(1,2)", "\r=SUM(1,2)", "\n=SUM(1,2)", "＝SUM(1,2)",
))
def test_xlsx_renderer_keeps_attacker_strings_as_text(attack):
    from backend.app.api.run_export import render_xlsx

    payload = {
        "schema_version": 1,
        "result": {"run_id": "run", "case_id": "case", "snapshot_id": "snap",
                   "revision": 2, "run_status": "partial", "is_current": False,
                   "checks": [{"supplier": attack, "field_key": "memory",
                               "processing_status": "complete", "business_status": attack,
                               "reason_code": attack, "is_formal": False,
                               "evidence_refs": ["e-1"], "reason_evidence_refs": ["e-1"],
                               "fact_ids": [], "check_result_id": "check",
                               "rule_version": "memory-v1"}]},
        "sources": [{"evidence_id": "e-1", "document_id": "doc",
                     "document_name": attack, "sheet": "Sheet1", "cell": "B6",
                     "page": None, "source_text": attack}],
    }
    content = render_xlsx(payload)
    book = load_workbook(io.BytesIO(content))
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "f"
    display = attack.replace("\r\n", "\n").replace("\r", "\n")
    assert book["核验结果"]["A2"].value == display
    assert book["核验结果"]["A2"].data_type == "s"
    assert book["原文来源"]["G2"].value == display
    with ZipFile(io.BytesIO(content)) as archive:
        sheets = [archive.read(name) for name in archive.namelist()
                  if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")]
    assert all(b"<f" not in sheet for sheet in sheets)
