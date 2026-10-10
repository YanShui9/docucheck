"""Category examples exercise uploaded documents, confirmation and formal APIs."""

import json
import re
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from backend.app.db.ingestion import DocumentIngestion
from backend.app.db.models import document_versions
from backend.app.db.run_executor import execute_run
from backend.tests import test_document_ingestion
from backend.tests.item_binding_helpers import (
    confirm_single_item_bindings,
    confirm_xlsx_item_ranges,
)
from backend.tests.test_dynamic_agent import ScriptProvider, call
from backend.tests.test_dynamic_workflow import workbook

ingestion_case = test_document_ingestion.ingestion_case


@pytest.mark.parametrize("item,metric,unit,minimum,above,below", [
    ("苹果", "净重", "g", "100", "120", "80"),
    ("衬衫", "棉含量", "%", "80", "90", "70"),
    ("显示器", "屏幕亮度", "nit", "300", "400", "250"),
    ("大米", "净含量", "kg", "5", "6", "4"),
    ("陶瓷杯", "耐热温度", "摄氏度", "100", "120", "90"),
])
def test_uploaded_arbitrary_metric_proposal_confirmation_and_formal_export(
    ingestion_case, item, metric, unit, minimum, above, below,
):
    client, engine, storage, case_id = ingestion_case
    rows = [("requirement", None, f"{item} {metric}不低于{minimum}{unit}"),
            ("quote", "vendor-a", f"{item} {metric}{above}{unit}"),
            ("quote", "vendor-b", f"{item} {metric}{below}{unit}")]
    for revision, (role, supplier, text) in enumerate(rows):
        data = {"role": role, "expected_revision": str(revision)}
        if supplier:
            data.update(supplier_handle=supplier, supplier_name=supplier)
        response = client.post(f"/api/v1/cases/{case_id}/documents", data=data,
                               headers={"Idempotency-Key": f"category-upload-{revision}"},
                               files={"file": (f"{supplier or '买方需求'}.xlsx", workbook(text),
                                               "application/octet-stream")})
        assert response.status_code == 201, response.text
        snapshot = response.json()["snapshot_id"]
    with engine.connect() as connection:
        documents = list(connection.scalars(select(document_versions.c.id).where(
            document_versions.c.case_id == case_id,
            document_versions.c.snapshot_id == snapshot)))
    ingestion = DocumentIngestion(engine, storage)
    for document_id in documents:
        assert ingestion.process_document(case_id, snapshot, document_id)["parse_status"] == "complete"
    conversation = client.post(f"/api/v1/cases/{case_id}/conversations").json()["id"]

    def submit(operation, key, question):
        response = client.post(f"/api/v1/conversations/{conversation}/messages",
                               headers={"Idempotency-Key": key},
                               json={"content": question, "operation": operation, "expected_revision": 3})
        assert response.status_code == 202, response.text
        receipt = response.json()
        provider = ScriptProvider(decision="propose_checklist" if operation == "propose_checklist"
                                  else "review_checklist")
        report = execute_run(engine, receipt["run_id"], receipt["attempt_id"], provider,
                             heartbeat_interval_seconds=60)
        assert report.status == "completed", report
        return receipt["run_id"]

    proposal_id = submit("propose_checklist", "category-proposal", "根据需求提出审查指标")
    response = client.get(f"/api/v1/agent-runs/{proposal_id}/proposal")
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["requires_user_confirmation"] is True
    assert preview["definition"]["items"][0]["name"] == item
    assert preview["definition"]["criteria"][0]["name"] == metric
    assert preview["definition"]["criteria"][0]["rule"]["value"] == minimum
    checklist_url = f"/api/v1/cases/{case_id}/snapshots/{snapshot}/checklist"
    assert client.get(checklist_url).json()["confirmed_version"] is None
    draft = client.put(checklist_url, headers={"Idempotency-Key": "category-draft"},
                       json={**preview["definition"], "expected_revision": 3,
                             "expected_checklist_revision": 0, "reason": "买方核对要求原文"})
    assert draft.status_code == 200, draft.text
    confirmed = client.post(checklist_url + "/confirm", headers={"Idempotency-Key": "category-confirm"},
                            json={"expected_revision": 3,
                                  "expected_checklist_revision": draft.json()["checklist_revision"],
                                  "criterion_ids": [preview["definition"]["criteria"][0]["criterion_id"]],
                                  "confirmed": True, "reason": "确认本次采购要求"})
    assert confirmed.status_code == 200, confirmed.text
    confirm_single_item_bindings(engine, case_id, snapshot, confirmed.json())
    review_id = submit("review", "category-review", "完整审查两家供应商")
    checks = client.get(f"/api/v1/agent-runs/{review_id}/checks").json()
    assert checks["checklist_version_id"] == confirmed.json()["confirmed_version"]["id"]
    assert {(c["supplier"], c["business_status"]) for c in checks["checks"]} == {
        ("vendor-a", "满足"), ("vendor-b", "不满足")}
    assert all(c["is_formal"] and c["item_name"] == item and c["criterion_name"] == metric
               for c in checks["checks"])
    result = client.get(f"/api/v1/agent-runs/{review_id}").json()
    assert f"{item} · {metric}" in result["answer_text"]
    assert "采购资料!A1" in result["answer_text"]
    exported = client.get(f"/api/v1/agent-runs/{review_id}/export").json()
    assert exported["result"]["checks"] == checks["checks"]
    excel = client.get(f"/api/v1/agent-runs/{review_id}/export?format=xlsx")
    assert excel.status_code == 200, excel.text
    book = load_workbook(BytesIO(excel.content))
    assert book["核验结果"]["B2"].value == metric


def test_multi_item_partial_confirmation_reports_only_confirmed_scope(ingestion_case):
    client, engine, storage, case_id = ingestion_case
    book = Workbook()
    book.active.title = "报价"
    book.active["A1"] = "苹果 数量12箱"
    book.active["A2"] = "梨 数量19箱"
    output = BytesIO()
    book.save(output)
    response = client.post(f"/api/v1/cases/{case_id}/documents",
                           headers={"Idempotency-Key": "multi-item-upload"},
                           data={"role": "quote", "supplier_handle": "vendor-a",
                                 "supplier_name": "甲供应商", "expected_revision": "0"},
                           files={"file": ("水果报价.xlsx", output.getvalue(), "application/octet-stream")})
    assert response.status_code == 201, response.text
    uploaded = response.json()
    snapshot = uploaded["snapshot_id"]
    assert DocumentIngestion(engine, storage).process_document(
        case_id, snapshot, uploaded["document_id"])["parse_status"] == "complete"
    definition = {
        "items": [{"item_id": "apple", "name": "苹果"}, {"item_id": "pear", "name": "梨"}],
        "criteria": [{"criterion_id": item + "-quantity", "item_id": item, "name": "数量",
                      "rule": {"type": "numeric", "operator": "gte", "value": count, "unit": "箱"},
                      "source": {"kind": "user_defined", "reason": "采购人员明确要求"}}
                     for item, count in (("apple", "10"), ("pear", "20"))] + [
            {"criterion_id": "apple-authenticity", "item_id": "apple", "name": "检测报告真实性",
             "rule": {"type": "manual", "instruction": "通过检测机构核实报告真伪"},
             "source": {"kind": "user_defined", "reason": "人工复核要求"}}],
    }
    url = f"/api/v1/cases/{case_id}/snapshots/{snapshot}/checklist"
    saved = client.put(url, headers={"Idempotency-Key": "multi-item-draft"},
                       json={**definition, "expected_revision": 1, "expected_checklist_revision": 0,
                             "reason": "按两种水果分别设置指标"})
    assert saved.status_code == 200, saved.text
    first = client.post(url + "/confirm", headers={"Idempotency-Key": "multi-item-confirm"},
                        json={"expected_revision": 1, "expected_checklist_revision": 1, "confirmed": True,
                              "criterion_ids": ["apple-quantity", "pear-quantity"],
                              "reason": "先核验已确认数量，报告真伪另外复核"})
    assert first.status_code == 200, first.text
    assert first.json()["excluded_criterion_ids"] == ["apple-authenticity"]
    confirm_xlsx_item_ranges(engine, case_id, snapshot, first.json())
    conversation = client.post(f"/api/v1/cases/{case_id}/conversations").json()["id"]

    class ItemProvider(ScriptProvider):
        def extract(self, messages, schema, **options):
            data = json.loads(messages[-1].content)
            names = {i["item_id"]: i["name"] for i in definition["items"]}
            facts = []
            for criterion in data["criteria"]:
                if criterion["rule"]["type"] != "numeric":
                    continue
                for block in data["evidence"]:
                    match = re.fullmatch(re.escape(names[criterion["item_id"]]) + r" 数量(\d+)箱", block["text"])
                    if block["role"] != "requirement" and match:
                        facts.append({"criterion_id": criterion["criterion_id"],
                                      "item_id": criterion["item_id"], "supplier": data["supplier"],
                                      "value": match[1], "unit": "箱", "quote": block["text"],
                                      "evidence_refs": [block["evidence_id"]]})
            return call(schema["function"]["name"], {"facts": facts})

    def review(key):
        submitted = client.post(f"/api/v1/conversations/{conversation}/messages",
                                headers={"Idempotency-Key": key},
                                json={"content": "审查已确认清单", "expected_revision": 1})
        assert submitted.status_code == 202, submitted.text
        receipt = submitted.json()
        report = execute_run(engine, receipt["run_id"], receipt["attempt_id"], ItemProvider(),
                             heartbeat_interval_seconds=60)
        assert report.status == "completed", report
        return report, client.get(f"/api/v1/agent-runs/{receipt['run_id']}/checks").json()

    partial, checks = review("multi-item-review")
    assert {(c["item_name"], c["criterion_name"], c["business_status"]) for c in checks["checks"]} == {
        ("苹果", "数量", "满足"), ("梨", "数量", "不满足")}
    assert all(c["is_formal"] for c in checks["checks"])
    assert "已确认" in partial.answer_text
    second = client.post(url + "/confirm", headers={"Idempotency-Key": "multi-item-confirm-rest"},
                         json={"expected_revision": 1, "expected_checklist_revision": 2,
                               "confirmed": True, "criterion_ids": ["apple-authenticity"],
                               "reason": "把人工真实性复核项纳入清单"})
    assert second.status_code == 200, second.text
    confirm_xlsx_item_ranges(engine, case_id, snapshot, second.json())
    _, complete = review("multi-item-review-rest")
    assert len(complete["checks"]) == 3
    assert next(c for c in complete["checks"] if c["criterion_id"] == "apple-authenticity")[
        "business_status"] == "待人工确认"
    old = client.get(f"/api/v1/agent-runs/{partial.run_id}/checks").json()
    assert not old["is_current"] and not any(c["is_formal"] for c in old["checks"])
