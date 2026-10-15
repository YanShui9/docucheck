"""Conservative presence decisions survive SQLite history and public exports."""

from copy import deepcopy
from hashlib import sha256
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from backend.app.criteria.facts import BINDING_RULE_VERSION
from backend.app.db.dynamic_repository import DynamicRepository
from backend.app.db.models import dynamic_batches, dynamic_runs
from backend.app.db.run_repository import RepositoryError
from backend.app.evidence.repository import digest
from backend.tests import test_material_repository as material_helpers

ingestion_case = material_helpers.ingestion_case
material_case = material_helpers.material_case


@pytest.mark.parametrize("material", [False, True], ids=["declaration", "attached-looking-body"])
def test_current_presence_report_requires_body_verification(material_case, material):
    """A complete extraction with attached-looking material is still unverified."""
    client, engine, *_ = material_case
    run_id, payload, _ = material_helpers.persist_material_report(
        material_case, rule_version=BINDING_RULE_VERSION, material=material
    )
    check = payload["checks"][0]
    assert check["business_status"] == "待人工确认"
    assert check["reason_code"] == "material_body_unverified"
    assert check["processing_status"] == "complete" and check["is_formal"] is True
    assert payload.get("needs_review", False) is False
    row, verified, _ = DynamicRepository(engine).verified(run_id)
    assert row["status"] == "completed" and verified == payload
    checks = client.get(f"/api/v1/agent-runs/{run_id}/checks")
    assert checks.status_code == 200, checks.text
    public = checks.json()
    assert public["run_status"] == "completed"
    assert public["checks"][0] == check | {"checklist_version_id": public["checklist_version_id"]}
    facts = client.get(f"/api/v1/agent-runs/{run_id}/facts")
    assert facts.status_code == 200, facts.text
    assert facts.json()["fact_sets"][0]["facts"] == [
        fact | {"raw_value": fact["value"], "raw_unit": fact["unit"]}
        for fact in payload["facts"]
    ]


def test_v11_history_and_cache_stay_frozen_while_v12_exports_all_sources(material_case):
    """New manual decisions cannot rewrite or reuse old positive report identities."""
    client, engine, case_id, snapshot, *_ = material_case
    old_id, old_payload, _ = material_helpers.persist_material_report(
        material_case, rule_version="dynamic-rules-v11"
    )
    before = deepcopy(old_payload)
    old_export = client.get(f"/api/v1/agent-runs/{old_id}/export")
    assert old_export.status_code == 200, old_export.text
    old_export = old_export.json()
    new_id, new_payload, repo = material_helpers.persist_material_report(
        material_case, rule_version="dynamic-rules-v12"
    )
    store = DynamicRepository(engine)
    assert store.verified(old_id)[1] == before
    assert store.verified(new_id)[1] == new_payload
    assert before["checks"][0]["business_status"] == "满足"
    assert new_payload["checks"][0]["business_status"] == "待人工确认"
    assert new_payload["checks"][0]["reason_code"] == "material_body_unverified"
    assert new_payload["facts"] == before["facts"]
    assert {e["input_digest"] for e in before["batch_manifest"]}.isdisjoint(
        {e["input_digest"] for e in new_payload["batch_manifest"]}
    )
    with engine.connect() as connection:
        cached = {run_id: list(connection.execute(select(dynamic_batches).where(
            dynamic_batches.c.run_id == run_id)).mappings()) for run_id in (old_id, new_id)}
    for run_id, payload in ((old_id, before), (new_id, new_payload)):
        assert {row["input_digest"] for row in cached[run_id]} == {
            entry["input_digest"] for entry in payload["batch_manifest"]}
        candidates = [candidate for row in cached[run_id] for candidate in row["output_json"]["facts"]]
        assert candidates[0]["material_sources"] == payload["facts"][0]["material_sources"]
        assert all(digest(row["output_json"]) == row["output_digest"] for row in cached[run_id])
    assert client.get(f"/api/v1/agent-runs/{old_id}/export").json() == old_export

    exported = client.get(f"/api/v1/agent-runs/{new_id}/export")
    assert exported.status_code == 200, exported.text
    exported = exported.json()
    fact = new_payload["facts"][0]
    body_refs = fact["material_sources"][0]["evidence_refs"]
    refs = set(fact["evidence_refs"] + body_refs)
    check = exported["result"]["checks"][0]
    assert check["material_source_groups"] == [{"fact_id": fact["fact_id"], "evidence_refs": body_refs}]
    assert set(check["evidence_refs"]) == refs
    assert {source["evidence_id"] for source in exported["sources"]} == refs
    by_id = {block.evidence_id: block for block in repo.blocks}
    for source in exported["sources"]:
        block = by_id[source["evidence_id"]]
        evidence = client.get(
            f"/api/v1/cases/{case_id}/snapshots/{snapshot}/evidence/{block.evidence_id}")
        assert evidence.status_code == 200, evidence.text
        assert evidence.json()["source_text"] == source["source_text"] == block.text
        original = client.get(
            f"/api/v1/cases/{case_id}/snapshots/{snapshot}/documents/{block.document_id}/original")
        assert original.status_code == 200, original.text
        assert sha256(original.content).hexdigest() == block.source_hash
        original_book = load_workbook(BytesIO(original.content))
        assert original_book[source["sheet"]][source["cell"]].value == block.text
    excel = client.get(f"/api/v1/agent-runs/{new_id}/export?format=xlsx")
    assert excel.status_code == 200, excel.text
    book = load_workbook(BytesIO(excel.content))
    assert [book["核验结果"].cell(2, col).value for col in (3, 4, 5, 6)] == [
        "complete", "待人工确认", "material_body_unverified", "dynamic-rules-v12"]
    assert book["材料引用分组"]["D2"].value == fact["fact_id"]
    assert book["材料引用分组"]["E2"].value == "\n".join(body_refs)
    assert {book["原文来源"].cell(row, 1).value for row in range(2, book["原文来源"].max_row + 1)} == refs
    assert {book["原文来源"].cell(row, 7).value for row in range(2, book["原文来源"].max_row + 1)} == {
        by_id[ref].text for ref in refs}


@pytest.mark.parametrize("tamper", ["promote_result", "drop_cached_body"])
def test_v12_history_rejects_rehashed_result_or_cache_tampering(material_case, tamper):
    """A recomputed JSON digest does not authorize promotion or loss of provenance."""
    _, engine, *_ = material_case
    run_id, payload, _ = material_helpers.persist_material_report(
        material_case, rule_version="dynamic-rules-v12"
    )
    store = DynamicRepository(engine)
    assert store.verified(run_id)[1] == payload
    with engine.begin() as connection:
        if tamper == "promote_result":
            changed = deepcopy(payload)
            changed["checks"][0].update(business_status="满足", reason_code="threshold_met")
            connection.execute(update(dynamic_runs).where(dynamic_runs.c.run_id == run_id).values(
                payload_json=changed, payload_digest=digest(changed)))
        else:
            rows = connection.execute(select(dynamic_batches).where(
                dynamic_batches.c.run_id == run_id)).mappings()
            row = next(row for row in rows if row["output_json"]["facts"])
            changed = deepcopy(row["output_json"])
            changed["facts"][0].pop("material_sources")
            connection.execute(update(dynamic_batches).where(
                dynamic_batches.c.run_id == run_id, dynamic_batches.c.batch_id == row["batch_id"]
            ).values(output_json=changed, output_digest=digest(changed)))
    with pytest.raises(RepositoryError, match="dynamic_source_changed|batch_source_changed"):
        store.verified(run_id)
