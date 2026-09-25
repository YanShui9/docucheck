"""Verify that the synthetic development pack and its labels match the originals."""

import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

PACK = Path(__file__).resolve().parents[2] / "examples" / "dev-pack-001"
FIELDS = {
    "memory",
    "storage",
    "screen_size",
    "weight",
    "warranty",
    "quantity",
    "unit_price",
    "delivery_days",
}


def test_source_documents_are_extractable() -> None:
    requirement = PdfReader(PACK / "requirements.pdf")
    specification = PdfReader(PACK / "vendor-a-spec.pdf")
    assert len(requirement.pages) == len(specification.pages) == 1
    assert "已装内存至少 16 GB" in requirement.pages[0].extract_text()
    assert "主机重量不超过 1.6 kg" in requirement.pages[0].extract_text()
    assert "已装内存 32 GB" in specification.pages[0].extract_text()
    assert "标称 SSD 1 TB" in specification.pages[0].extract_text()

    workbook = load_workbook(PACK / "vendor-a-quote.xlsx", read_only=True, data_only=False)
    assert workbook.sheetnames == ["报价"]
    sheet = workbook["报价"]
    assert sheet["B4"].value == "DC-14-A"
    assert sheet["B6"].value == "已装内存 16 GB"
    assert sheet["B7"].value == "10 台"
    assert sheet["B8"].value == "含税单价 5800 元/台"
    assert sheet["B9"].value == "合同签订后 10 个工作日"
    workbook.close()


def test_eight_labels_point_to_original_evidence() -> None:
    labels = json.loads((PACK / "labels.json").read_text(encoding="utf-8"))
    assert labels["dataset_kind"] == "synthetic_development"
    assert {item["field_key"] for item in labels["results"]} == FIELDS
    assert len(labels["results"]) == 8
    assert Counter(item["status"] for item in labels["results"]) == {
        "满足": 4,
        "不满足": 1,
        "资料冲突": 1,
        "证据不足": 1,
        "待人工确认": 1,
    }

    pdf_text = {
        name: PdfReader(PACK / name).pages[0].extract_text()
        for name in ("requirements.pdf", "vendor-a-spec.pdf")
    }
    workbook = load_workbook(PACK / "vendor-a-quote.xlsx", read_only=True)
    for result in labels["results"]:
        assert result["requirement_evidence"]
        assert result["reason"]
        for evidence in result["requirement_evidence"] + result["supplier_evidence"]:
            if evidence["document"].endswith(".pdf"):
                assert evidence["page"] == 1
                assert evidence["text"] in pdf_text[evidence["document"]]
            else:
                assert evidence["document"] == "vendor-a-quote.xlsx"
                assert evidence["sheet"] == "报价"
                assert workbook[evidence["sheet"]][evidence["cell"]].value == evidence["text"]
    workbook.close()

    warranty = next(item for item in labels["results"] if item["field_key"] == "warranty")
    assert warranty["supplier_evidence"] == []
    assert set(warranty["processed_supplier_documents"]) == {
        "vendor-a-quote.xlsx",
        "vendor-a-spec.pdf",
    }
    memory = next(item for item in labels["results"] if item["field_key"] == "memory")
    assert {item["document"] for item in memory["supplier_evidence"]} == {
        "vendor-a-quote.xlsx",
        "vendor-a-spec.pdf",
    }
