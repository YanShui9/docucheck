"""Offline wire contracts for dynamic evaluation; no live service or model."""

import importlib.util
import io
import json
from copy import deepcopy
from hashlib import sha256
from types import SimpleNamespace

import httpx
import pytest
from openpyxl import Workbook, load_workbook


def runner():
    assert importlib.util.find_spec("evals.dynamic.run") is not None, "dynamic HTTP runner missing"
    from evals.dynamic import run

    return run


@pytest.mark.parametrize("url", [
    "http://localhost:8000", "https://example.org", "http://127.0.0.1@evil.test",
    "http://127.0.0.1:8000/private", "http://127.0.0.1:8000?token=x",
    "http://127.0.0.1:8000#token", "http://127.0.0.1:0",
])
def test_target_requires_literal_loopback_without_credentials_or_extra_parts(url):
    with pytest.raises(ValueError, match="literal_loopback"):
        runner().validate_api_url(url)


@pytest.mark.parametrize("url", ["http://127.0.0.1:8000", "http://[::1]:8000/"])
def test_literal_loopback_is_accepted(url):
    assert runner().validate_api_url(url) == url.rstrip("/")


def test_preview_never_reads_credentials_or_opens_http(tmp_path, monkeypatch, capsys):
    run = runner()
    monkeypatch.setattr(run, "prepare_protocol", lambda *a, **k: {
        "planned_tasks": 1, "planned_files": 2, "planned_positions": 1,
        "maximum_agent_requests": 24, "tasks": [],
    })
    monkeypatch.setattr(run.httpx, "Client", lambda *a, **k: pytest.fail("preview_network"))
    monkeypatch.setattr(run, "os", SimpleNamespace(environ=SimpleNamespace(
        get=lambda *a, **k: pytest.fail("preview_credentials"))))
    destination = tmp_path / "preview"
    monkeypatch.setattr("sys.argv", ["dynamic_run", "--output", str(destination)])
    assert run.main() == 0
    assert json.loads(capsys.readouterr().out)["mode"] == "preview"
    assert (destination / "protocol.json").is_file()


def test_output_directory_is_never_overwritten(tmp_path, monkeypatch):
    run = runner()
    monkeypatch.setattr(run, "prepare_protocol", lambda *a, **k: {})
    destination = tmp_path / "existing"
    destination.mkdir()
    marker = destination / "protocol.json"
    marker.write_text("original", encoding="utf-8")
    monkeypatch.setattr("sys.argv", ["dynamic_run", "--output", str(destination)])
    assert run.main() == 2
    assert marker.read_text(encoding="utf-8") == "original"


class Clock:
    value = 0.0

    def now(self):
        return self.value

    def sleep(self, seconds):
        self.value += seconds


def input_pack(tmp_path):
    directory = tmp_path / "inputs" / "dyn-wire"
    directory.mkdir(parents=True)
    workbook = Workbook()
    workbook.active.title = "报价"
    workbook.active["A3"] = "品目：苹果手机 型号：DEMO-P256"
    workbook.active["A4"] = "数量"
    workbook.active["B4"] = "100台"
    workbook.save(directory / "quote.xlsx")
    content = (directory / "quote.xlsx").read_bytes()
    workbook.active.title = "需求"
    workbook.active["A3"] = "品目：苹果手机"
    workbook.active["B4"] = "数量至少100台"
    workbook.save(directory / "requirements.xlsx")
    buyer_content = (directory / "requirements.xlsx").read_bytes()
    anchor = {"document": "quote.xlsx", "role": "quote", "supplier_handle": "vendor-a",
              "item_id": "phone", "text": "品目：苹果手机 型号：DEMO-P256",
              "sheet": "报价", "cell": "A3"}
    manifest = {
        "schema_version": 1, "protocol_id": "dynamic-dev-v1", "pack_id": "dyn-wire",
        "data_class": "ai_authored_seen_development", "category": "phone",
        "source_family_ids": ["wire-test"], "supplier_handles": ["vendor-a"],
        "authors": ["AI fixture"], "reviewers": ["AI fixture"],
        "buyer_anchors": {"quantity": {"document": "requirements.xlsx", "role": "requirement", "supplier_handle": None,
                                      "item_id": "phone", "text": "数量至少100台", "sheet": "需求", "cell": "B4"}},
        "documents": [{"filename": "quote.xlsx", "role": "quote", "supplier_handle": "vendor-a",
                       "sha256": sha256(content).hexdigest(), "bytes": len(content), "pages": 0},
                      {"filename": "requirements.xlsx", "role": "requirement", "supplier_handle": None,
                       "sha256": sha256(buyer_content).hexdigest(), "bytes": len(buyer_content), "pages": 0}],
        "checklist_definition": {
            "items": [{"item_id": "phone", "name": "苹果手机", "specification": "DEMO-P256"}],
            "criteria": [{"criterion_id": "quantity", "item_id": "phone", "name": "数量",
                          "rule": {"type": "numeric", "operator": "gte", "value": "100", "unit": "台"},
                          "source": {"kind": "user_defined", "reason": "合成需求数量100台"}}],
        },
        "bindings": [{"document": "quote.xlsx", "item_id": "phone", "sku": "DEMO-P256",
                      "scope": {"sheet": "报价", "cell_range": "A3:B4"},
                      "anchor": anchor, "sku_anchor": anchor},
                     {"document": "requirements.xlsx", "item_id": "phone", "sku": None,
                      "scope": {"sheet": "需求", "cell_range": "A3:B4"},
                      "anchor": {"document": "requirements.xlsx", "role": "requirement", "supplier_handle": None,
                                 "item_id": "phone", "text": "品目：苹果手机", "sheet": "需求", "cell": "A3"},
                      "sku_anchor": None}],
    }
    return directory, manifest


class WireServer:
    """Deterministic HTTP response fixture; production request schemas validate edits."""

    case_id = "a" * 32
    document_id = "b" * 32
    snapshot_id = "c" * 32
    version_id = "d" * 32
    binding_id = "e" * 32
    conversation_id = "f" * 32
    run_id = "1" * 32
    attempt_id = "2" * 32
    message_id = "3" * 32
    evidence_id = "4" * 32

    def __init__(self, manifest, *, parsing="complete", status="completed",
                 unknown_submit=False, wrong_sha=False, wrong_attempt=False,
                 wrong_anchor=False, xlsx_status=None):
        self.manifest = manifest
        self.parsing, self.status = parsing, status
        self.unknown_submit, self.wrong_sha = unknown_submit, wrong_sha
        self.wrong_attempt, self.wrong_anchor = wrong_attempt, wrong_anchor
        self.xlsx_status = xlsx_status
        self.requests = []
        self.uploaded = 0
        self.confirmed_definition = self.manifest["checklist_definition"]
        self.checks = {
            "run_id": self.run_id, "case_id": self.case_id, "snapshot_id": self.snapshot_id,
            "revision": 2, "run_status": status, "is_current": True,
            "checklist_version_id": self.version_id, "item_binding_version_id": self.binding_id,
            "item_binding_digest": "a" * 64, "payload_digest": "b" * 64,
            "checks": [{"check_result_id": "5" * 64, "snapshot_id": self.snapshot_id,
                        "supplier": "vendor-a", "field_key": "quantity", "item_id": "phone",
                        "criterion_id": "quantity", "item_name": "苹果手机", "criterion_name": "数量",
                        "processing_status": "complete", "business_status": "满足",
                        "reason_code": "rule_match", "evidence_refs": [self.evidence_id],
                        "reason_evidence_refs": ["8" * 32], "fact_ids": [], "rule_version": "wire",
                        "is_formal": True, "checklist_version_id": self.version_id}],
        }

    def request(self, method, path, **kwargs):
        self.requests.append((method, path, deepcopy(kwargs)))
        if method == "POST" and path == "/api/v1/cases":
            assert kwargs["json"] == {"project_id": "9" * 32}
            return httpx.Response(201, json={"id": self.case_id, "revision": 0})
        if method == "POST" and path.endswith("/documents"):
            assert kwargs["data"]["expected_revision"] == str(self.uploaded)
            if self.uploaded == 0:
                assert kwargs["data"]["supplier_handle"] == "vendor-a"
            else:
                assert "supplier_handle" not in kwargs["data"]
            assert kwargs["headers"]["Idempotency-Key"]
            doc = self.manifest["documents"][self.uploaded]
            assert sha256(kwargs["files"]["file"][1]).hexdigest() == doc["sha256"]
            self.uploaded += 1
            return httpx.Response(201, json={"document_id": self.document_id if self.uploaded == 1 else "6" * 32, "case_id": self.case_id,
                "snapshot_id": self.snapshot_id, "revision": self.uploaded,
                "source_sha256": "0" * 64 if self.wrong_sha else doc["sha256"]})
        if method == "GET" and path.endswith("/documents"):
            return httpx.Response(200, json={"case_id": self.case_id, "snapshot_id": self.snapshot_id,
                "revision": 2, "documents": [{"id": self.document_id if i == 0 else "6" * 32, "original_name": doc["filename"],
                 "source_sha256": doc["sha256"], "role": doc["role"],
                 "supplier_handle": doc["supplier_handle"], "model_name": "DEMO-P256" if i == 0 else None,
                 "parse_status": self.parsing} for i, doc in enumerate(self.manifest["documents"])]})
        if method == "GET" and "/documents/" in path and path.endswith("/evidence"):
            offset = kwargs.get("params", {}).get("after_ordinal", -1)
            i = 0 if self.document_id in path else 1
            document_id = self.document_id if i == 0 else "6" * 32
            anchor = self.manifest["bindings"][i]["anchor"]
            evidence = [] if offset >= 0 else [{"id": self.evidence_id if i == 0 else "7" * 32, "document_id": document_id,
                "ordinal": 0, "source_text": "wrong" if self.wrong_anchor else anchor["text"],
                "page_idx": None, "sheet": anchor["sheet"], "cell": "A3", "bbox": None,
                "parser_version": "xlsx-wire-parser"}]
            if i == 1 and offset < 0:
                evidence.append({"id": "8" * 32, "document_id": document_id, "ordinal": 1,
                    "source_text": "数量至少100台", "page_idx": None, "sheet": "需求", "cell": "B4", "bbox": None,
                    "parser_version": "xlsx-wire-parser"})
            return httpx.Response(200, json={"case_id": self.case_id, "snapshot_id": self.snapshot_id,
                "document_id": document_id, "parser_version": "xlsx-wire-parser",
                "parse_status": "complete", "coverage_issues": [], "evidence": evidence,
                "next_ordinal": 0 if i == 0 else 1})
        if method == "PUT" and path.endswith("/checklist"):
            from backend.app.api.checklists import DraftIn
            body = DraftIn.model_validate(kwargs["json"])
            assert body.expected_revision == 2 and body.expected_checklist_revision == 0
            assert body.items[0].item_id == "phone"
            assert body.criteria[0].source.context is not None, "runtime buyer context missing"
            assert body.criteria[0].source.context.evidence_refs == ["8" * 32]
            self.confirmed_definition = body.model_dump(mode="json", include={"items", "criteria"})
            return httpx.Response(200, json={"case_id": self.case_id, "snapshot_id": self.snapshot_id,
                "revision": 2, "checklist_revision": 1, "draft": self.manifest["checklist_definition"]})
        if method == "POST" and path.endswith("/checklist/confirm"):
            from backend.app.api.checklists import ConfirmIn
            body = ConfirmIn.model_validate(kwargs["json"])
            assert body.criterion_ids == ["quantity"] and body.confirmed
            assert body.expected_revision == 2 and body.expected_checklist_revision == 1
            return httpx.Response(200, json={"case_id": self.case_id, "snapshot_id": self.snapshot_id,
                "revision": 2, "checklist_revision": 2, "execution_available": True,
                "confirmed_version": {"id": self.version_id, "definition": self.confirmed_definition}})
        if method == "PUT" and path.endswith("/item-bindings"):
            from backend.app.api.item_bindings import ConfirmBindingsIn
            body = ConfirmBindingsIn.model_validate(kwargs["json"])
            assert body.checklist_version_id == self.version_id
            assert (body.expected_revision, body.expected_checklist_revision, body.expected_binding_revision) == (2, 2, 0)
            assert body.entries[0].document_id == self.document_id
            assert body.entries[0].source_sha256 == self.manifest["documents"][0]["sha256"]
            assert body.entries[0].parser_version == "xlsx-wire-parser"
            assert body.entries[0].anchor.evidence_refs == [self.evidence_id]
            assert body.entries[0].sku_anchor.evidence_refs == [self.evidence_id]
            assert body.entries[0].selector.kind == "xlsx"
            assert len(body.entries) == 1  # Requirement provenance is not a supplier binding.
            return httpx.Response(200, json={"case_id": self.case_id, "snapshot_id": self.snapshot_id,
                "state": "confirmed", "binding_revision": 1,
                "current_version": {"id": self.binding_id, "digest": "a" * 64}})
        if method == "POST" and path.endswith("/conversations"):
            return httpx.Response(201, json={"id": self.conversation_id, "case_id": self.case_id})
        if method == "POST" and path.endswith("/messages"):
            from backend.app.api.schemas import MessageIn
            body = MessageIn.model_validate(kwargs["json"])
            assert body.operation == "review" and body.expected_revision == 2
            assert body.mode == "thorough"
            if self.unknown_submit:
                raise httpx.ReadTimeout("secret remote debug text must not enter attempt")
            return httpx.Response(202, json={"run_id": self.run_id, "attempt_id": self.attempt_id,
                                            "message_id": self.message_id})
        if method == "GET" and path == f"/api/v1/agent-runs/{self.run_id}":
            return httpx.Response(200, json={"id": self.run_id, "case_id": self.case_id,
                "conversation_id": self.conversation_id, "snapshot_id": self.snapshot_id, "revision": 2,
                "active_attempt_id": "8" * 32 if self.wrong_attempt else self.attempt_id,
                "checklist_version_id": self.version_id, "item_binding_version_id": self.binding_id,
                "item_binding_digest": "a" * 64, "status": self.status, "outcome_unknown": False,
                "answer_text": "vendor-a · 苹果手机 · 数量：满足。", "error_code": None,
                "model_id": "offline-wire-fixture", "usage": None})
        if method == "GET" and path.endswith("/checks"):
            return httpx.Response(200, json=self.checks)
        if method == "GET" and path.endswith("/progress"):
            return httpx.Response(200, json={"run_id": self.run_id, "workflow": "review", "run_status": self.status,
                "checklist_version_id": self.version_id, "item_binding_version_id": self.binding_id,
                "item_binding_digest": "a" * 64, "coverage": {"vendor-a": True},
                "item_coverage": {"vendor-a": {"phone": True}}, "validated_batches": 1, "planned_batches": 1})
        if method == "GET" and "/export" in path:
            payload = {"schema_version": 1, "result": self.checks, "sources": []}
            if kwargs.get("params", {}).get("format") == "xlsx" or "format=xlsx" in path:
                from backend.app.api.run_export import render_xlsx
                if self.xlsx_status:
                    payload = deepcopy(payload)
                    payload["result"]["checks"][0]["business_status"] = self.xlsx_status
                return httpx.Response(200, content=render_xlsx(payload))
            return httpx.Response(200, json=payload)
        if method == "POST" and path.endswith("/cancel"):
            return httpx.Response(200, json={"status": "cancelled"})
        pytest.fail(f"Unexpected wire request: {method} {path}")


def execute(tmp_path, **server_options):
    directory, manifest = input_pack(tmp_path)
    server = WireServer(manifest, **server_options)
    clock = Clock()
    result = runner().run_pack(server, directory, manifest=manifest, project_id="9" * 32,
                               deadline_seconds=2, clock=clock.now, sleep=clock.sleep)
    return result, server


def test_formal_review_saves_confirmed_versions_and_real_export_surfaces(tmp_path):
    result, server = execute(tmp_path)
    assert result["stage"] == "finished", result
    assert result["error_code"] is None
    assert result["observation"]["api_cells"][0]["business_status"] == "满足"
    visible = result["observation"]["visible_cells"][0]
    assert visible["surface"] == "status_only" and visible["is_formal"] is None
    assert visible["evidence_refs"] == []
    assert result["observation"]["coverage"][0]["source"].endswith("/progress")
    assert result["usage"]["state"] == "unknown" and result["usage"]["cost"] is None
    assert load_workbook(io.BytesIO(result["xlsx_bytes"]), read_only=True)["核验结果"]["D2"].value == "满足"
    assert result["attempt_id"] == server.attempt_id
    assert any(source["role"] == "requirement" and source["evidence_id"] == "8" * 32
               for source in result["observation"]["sources"])
    assert all("gold" not in json.dumps(kwargs, default=str) for _, _, kwargs in server.requests)
    assert not any(path.endswith(("/retry", "/resume", "/proposal")) for _, path, _ in server.requests)


def test_changed_input_is_rejected_before_creating_resources(tmp_path):
    directory, manifest = input_pack(tmp_path)
    (directory / "quote.xlsx").write_bytes(b"changed")
    server = WireServer(manifest)
    result = runner().run_pack(server, directory, manifest=manifest, project_id="9" * 32)
    assert result["error_code"] == "evaluation_input_changed"
    assert server.requests == []


@pytest.mark.parametrize("options,code", [
    ({"wrong_sha": True}, "upload_receipt_mismatch"),
    ({"wrong_anchor": True}, "binding_anchor_unresolved"),
    ({"wrong_attempt": True}, "run_receipt_mismatch"),
    ({"parsing": "pending"}, "parse_timeout"),
])
def test_failures_are_preserved_without_automatic_retry(tmp_path, options, code):
    result, server = execute(tmp_path, **options)
    assert result["error_code"] == code
    assert result["raw"]
    assert not any(path.endswith(("/retry", "/resume")) for _, path, _ in server.requests)
    assert sum(path.endswith("/messages") for _, path, _ in server.requests) <= 1


def test_unknown_submission_is_retained_once_without_secret_exception_text(tmp_path):
    result, server = execute(tmp_path, unknown_submit=True)
    assert result["error_code"] == "http_transport_error"
    assert result["outcome_unknown"] is True
    assert sum(path.endswith("/messages") for _, path, _ in server.requests) == 1
    assert "secret remote debug" not in json.dumps(result, default=str)


def test_xlsx_is_independently_read_and_mismatch_remains_visible(tmp_path):
    result, _ = execute(tmp_path, xlsx_status="不满足")
    assert result["observation"]["export_cells"][0]["business_status"] == "满足"
    assert result["observation"]["xlsx_export_cells"][0]["business_status"] == "不满足"
    assert "json_xlsx_export_mismatch" in result["observation"]["surface_errors"]


def test_server_unknown_stops_polling_and_retains_unknown_model_usage(tmp_path):
    directory, manifest = input_pack(tmp_path)
    server = WireServer(manifest, status="running")
    original_request = server.request

    def unknown_run(method, path, **kwargs):
        response = original_request(method, path, **kwargs)
        if method == "GET" and path == f"/api/v1/agent-runs/{server.run_id}":
            payload = response.json()
            payload["outcome_unknown"] = True
            return httpx.Response(200, json=payload)
        return response

    server.request = unknown_run
    clock = Clock()
    result = runner().run_pack(server, directory, manifest=manifest, project_id="9" * 32,
                               deadline_seconds=2, clock=clock.now, sleep=clock.sleep)
    assert result["error_code"] == "server_outcome_unknown"
    assert result["observation"]["run_status"] == "outcome_unknown"
    assert sum(path == f"/api/v1/agent-runs/{server.run_id}" for _, path, _ in server.requests) == 1


def test_server_error_after_mutation_is_retained_as_unknown_without_leaking_detail(tmp_path):
    directory, manifest = input_pack(tmp_path)
    server = WireServer(manifest)
    original_request = server.request

    def failed_upload(method, path, **kwargs):
        if method == "POST" and path.endswith("/documents"):
            server.requests.append((method, path, kwargs))
            return httpx.Response(502, json={"detail": "secret internal debug", "error_code": "upstream_timeout"})
        return original_request(method, path, **kwargs)

    server.request = failed_upload
    result = runner().run_pack(server, directory, manifest=manifest, project_id="9" * 32)
    assert result["error_code"] == "http_status_502"
    assert result["outcome_unknown"] is True
    assert "secret internal debug" not in json.dumps(result, default=str)


def test_visible_answer_ambiguous_labels_do_not_copy_api_cells(tmp_path):
    _, manifest = input_pack(tmp_path)
    second = deepcopy(manifest["checklist_definition"]["criteria"][0])
    second["criterion_id"] = "quantity2"
    manifest["checklist_definition"]["criteria"].append(second)
    cells, errors = runner()._visible_cells(manifest, "vendor-a · 苹果手机 · 数量：满足。")
    assert cells == []
    assert errors == ["visible_answer_identity_or_status_unresolved"]


def test_visible_statuses_do_not_parse_rendered_sources_or_source_quote_as_results(tmp_path):
    _, manifest = input_pack(tmp_path)
    answer = ("vendor-a · 苹果手机 · 数量：满足。\n核验来源：\n"
              + f"来源 [{'a' * 64}]：vendor-a-quote.xlsx · 采购资料!B4\n"
              + "原文：vendor-a · 苹果手机 · 数量：不满足。")
    cells, errors = runner()._visible_cells(manifest, answer)
    assert errors == []
    assert len(cells) == 1
    assert cells[0]["business_status"] == "满足"


def test_preflight_rejects_labels_or_gold_in_business_manifest(tmp_path):
    directory, manifest = input_pack(tmp_path)
    manifest["gold"] = {"quantity": "满足"}
    server = WireServer(manifest)
    result = runner().run_pack(server, directory, manifest=manifest, project_id="9" * 32)
    assert result["error_code"] == "evaluation_manifest_invalid"
    assert server.requests == []


def cli_fixture(tmp_path, monkeypatch, *, change_gold=False, project_fails=False, change_index=None):
    run = runner()
    root = tmp_path / "corpus"
    directory, manifest = input_pack(root)
    config_dir, gold_dir = root / "config", root / "gold"
    config_dir.mkdir()
    gold_dir.mkdir()
    config, gold = config_dir / "dyn-wire.json", gold_dir / "dyn-wire.json"
    config.write_text(json.dumps(manifest), encoding="utf-8")
    gold.write_text('{"local_gold_only": true}', encoding="utf-8")
    task = {"task_id": "dyn-wire-01", "pack_id": "dyn-wire", "manifest": manifest,
            "manifest_sha256": sha256(config.read_bytes()).hexdigest(),
            "gold_sha256": sha256(gold.read_bytes()).hexdigest()}
    index = root / "index.json"
    index.write_text(json.dumps({"packs": [{"pack_id": "dyn-wire"}]}), encoding="utf-8")
    protocol = {"tasks": [task], "index_sha256": sha256(index.read_bytes()).hexdigest()}
    monkeypatch.setattr(run, "prepare_protocol", lambda *a, **k: protocol)
    server = WireServer(manifest)
    options, projects = [], []

    class Client:
        def __init__(self, **kwargs):
            options.append(kwargs)
            if change_gold:
                gold.write_text("changed", encoding="utf-8")
            if change_index == "delete":
                index.unlink()
            elif change_index == "tamper":
                index.write_text('{"packs": []}', encoding="utf-8")

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def post(self, path, **kwargs):
            projects.append((path, kwargs))
            return httpx.Response(503 if project_fails else 201, json={"id": "9" * 32})

        def request(self, method, path, **kwargs):
            return server.request(method, path, **kwargs)

    monkeypatch.setattr(run.httpx, "Client", Client)
    monkeypatch.setattr(run, "os", SimpleNamespace(environ={"DOCUCHECK_EVAL_TOKEN": "fixture-secret-token"}))
    output = tmp_path / "output"
    monkeypatch.setattr("sys.argv", ["dynamic_run", "--send", "--root", str(root), "--output", str(output)])
    return run, output, server, options, projects, directory


def test_send_creates_new_synthetic_project_and_preserves_binary_and_json(tmp_path, monkeypatch):
    run, output, server, options, projects, _ = cli_fixture(tmp_path, monkeypatch)
    assert run.main() == 0
    assert len(projects) == 1 and projects[0][0] == "/api/v1/projects"
    assert options[0]["trust_env"] is False and options[0]["follow_redirects"] is False
    record = json.loads((output / "dyn-wire-01.json").read_bytes())
    assert "xlsx_bytes" not in record
    assert sha256((output / "dyn-wire-01.xlsx").read_bytes()).hexdigest() == record["xlsx_export"]["sha256"]
    assert json.loads((output / "observations.json").read_bytes())[0] == record["observation"]
    assert "fixture-secret-token" not in "".join(p.read_text(encoding="utf-8") for p in output.glob("*.json"))
    assert "local_gold_only" not in json.dumps(server.requests, default=str)


def test_hard_exit_during_submission_leaves_durable_unknown_request(tmp_path, monkeypatch):
    run, output, server, _, _, _ = cli_fixture(tmp_path, monkeypatch)
    original = server.request

    def stop_after_send(method, path, **kwargs):
        if method == "POST" and path.endswith("/messages"):
            journal = output / "dyn-wire-01.inflight.json"
            assert journal.is_file(), "request must be durable before it can reach the server"
            assert json.loads(journal.read_bytes())["raw"][-1]["state"] == "sending"
            # SystemExit deliberately bypasses the ordinary transport and
            # KeyboardInterrupt handlers, as a terminated process would.
            raise SystemExit("simulated_process_exit")
        return original(method, path, **kwargs)

    server.request = stop_after_send
    with pytest.raises(SystemExit, match="simulated_process_exit"):
        run.main()
    journal = output / "dyn-wire-01.inflight.json"
    assert journal.is_file(), "sent task must survive before its final report exists"
    record = json.loads(journal.read_bytes())
    assert record["stage"] == "submit"
    assert record["case_id"] == server.case_id
    assert record["outcome_unknown"] is True
    assert record["needs_reconciliation"] is True
    assert record["raw"][-1]["state"] == "sending"
    assert record["raw"][-1]["idempotency_key"]
    assert record["raw"][-1]["path"].endswith("/messages")
    assert not (output / "dyn-wire-01.json").exists()
    assert "fixture-secret-token" not in journal.read_text(encoding="utf-8")
    assert "local_gold_only" not in journal.read_text(encoding="utf-8")


def test_review_receipt_is_durable_before_the_first_poll(tmp_path, monkeypatch):
    run, output, server, _, _, _ = cli_fixture(tmp_path, monkeypatch)
    original = server.request

    def stop_before_poll(method, path, **kwargs):
        if method == "GET" and path.startswith("/api/v1/agent-runs/"):
            journal = output / "dyn-wire-01.inflight.json"
            assert journal.is_file(), "receipt must be durable before polling"
            saved = json.loads(journal.read_bytes())
            submission = next(r for r in saved["raw"] if r["path"].endswith("/messages"))
            assert submission["state"] == "received"
            assert submission["response"]["run_id"] == server.run_id
            raise SystemExit("simulated_process_exit")
        return original(method, path, **kwargs)

    server.request = stop_before_poll
    with pytest.raises(SystemExit, match="simulated_process_exit"):
        run.main()
    journal = output / "dyn-wire-01.inflight.json"
    assert journal.is_file()
    record = json.loads(journal.read_bytes())
    assert record["run_id"] == server.run_id
    assert record["attempt_id"] == server.attempt_id
    assert record["needs_reconciliation"] is True
    submission = next(r for r in record["raw"] if r["path"].endswith("/messages"))
    assert submission["state"] == "received"
    assert submission["response"]["run_id"] == record["run_id"]


def test_project_creation_is_reserved_before_any_http_request(tmp_path, monkeypatch):
    run, output, _, _, _, _ = cli_fixture(tmp_path, monkeypatch)
    original = run.httpx.Client.post

    def stop_before_project(self, path, **kwargs):
        assert (output / "project.inflight.json").is_file()
        assert (output / "dyn-wire-01.inflight.json").is_file()
        record = json.loads((output / "project.inflight.json").read_bytes())
        assert record["outcome_unknown"] is True
        assert record["raw"][-1]["state"] == "sending"
        assert record["raw"][-1]["json"] == kwargs["json"]
        assert record["raw"][-1]["idempotency_key"] == kwargs["headers"]["Idempotency-Key"]
        raise SystemExit("simulated_process_exit")

    assert original
    monkeypatch.setattr(run.httpx.Client, "post", stop_before_project)
    with pytest.raises(SystemExit, match="simulated_process_exit"):
        run.main()


def test_real_process_exit_cannot_erase_the_sent_request(tmp_path):
    import os
    import subprocess
    import sys
    from pathlib import Path

    from scripts.verify_offline import sanitized_env

    program = '''
import json, os, sys
from pathlib import Path
from pytest import MonkeyPatch
from scripts.offline_guard import NetworkGuard
from backend.tests.test_dynamic_eval_run import cli_fixture
guard = NetworkGuard()
guard.install()
patch = MonkeyPatch()
run, output, server, _, _, _ = cli_fixture(Path(sys.argv[1]), patch)
original = server.request
def exit_during_submission(method, path, **kwargs):
    if method == "POST" and path.endswith("/messages"):
        print(f"child_network_attempts={guard.violations}", flush=True)
        os._exit(73)
    return original(method, path, **kwargs)
server.request = exit_during_submission
run.main()
'''
    child = subprocess.run([sys.executable, "-c", program, str(tmp_path)],
                           cwd=Path(__file__).resolve().parents[2], env=sanitized_env(os.environ),
                           capture_output=True, text=True, encoding="utf-8", timeout=30, check=False)
    assert child.returncode == 73, child.stdout + child.stderr
    assert "child_network_attempts=0" in child.stdout
    record = json.loads((tmp_path / "output/dyn-wire-01.inflight.json").read_bytes())
    assert record["stage"] == "submit"
    assert record["outcome_unknown"] is True
    assert record["raw"][-1]["state"] == "sending"
    assert record["raw"][-1]["idempotency_key"]
    assert record["case_id"] == WireServer.case_id
    assert not (tmp_path / "output/dyn-wire-01.json").exists()


def test_journal_write_failure_blocks_external_mutation_and_preserves_reserved_json(tmp_path, monkeypatch):
    run, output, server, _, projects, _ = cli_fixture(tmp_path, monkeypatch)

    def disk_failure(*args):
        raise OSError("simulated_local_write_failure")

    monkeypatch.setattr(run, "atomic_replace", disk_failure)
    assert run.main() == 2
    assert not projects and not server.requests
    record = json.loads((output / "project.inflight.json").read_bytes())
    assert record["journal_state"] == "reserved"
    assert record["raw"] == []


def test_client_deadline_leaves_running_server_outcome_unknown(tmp_path):
    record, server = execute(tmp_path, status="running")
    assert record["error_code"] == "run_timeout"
    assert record["outcome_unknown"] is True
    assert record["observation"]["run_status"] == "running"
    assert record["usage"]["state"] == "unknown"
    assert sum(path.endswith("/messages") for _, path, _ in server.requests) == 1
    assert not any(path.endswith(("/retry", "/resume")) for _, path, _ in server.requests)


@pytest.mark.parametrize("stage", ["case", "submit"])
def test_successful_mutation_with_unusable_receipt_is_unknown(tmp_path, stage):
    directory, manifest = input_pack(tmp_path)
    server = WireServer(manifest)
    original = server.request

    def malformed_receipt(method, path, **kwargs):
        if method == "POST" and (
            stage == "case" and path == "/api/v1/cases"
            or stage == "submit" and path.endswith("/messages")
        ):
            return httpx.Response(201 if stage == "case" else 202, json={})
        return original(method, path, **kwargs)

    server.request = malformed_receipt
    record = runner().run_pack(server, directory, manifest=manifest, project_id="9" * 32)
    assert record["error_code"] == "invalid_input_or_api_shape"
    assert record["outcome_unknown"] is True
    assert record["usage"]["state"] == "unknown"
    assert record["automatic_retries"] is False


@pytest.mark.parametrize("damage", ["non_zip", "empty_sheet", "numeric_source_id"])
def test_corrupt_xlsx_is_preserved_as_artifact_failure_in_final_attempt(tmp_path, monkeypatch, damage):
    run, output, server, _, _, _ = cli_fixture(tmp_path, monkeypatch)
    original = server.request

    def broken_export(method, path, **kwargs):
        if method == "GET" and path.endswith("/export?format=xlsx"):
            if damage == "non_zip":
                return httpx.Response(200, content=b"not a ZIP workbook")
            if damage == "empty_sheet":
                book = Workbook()
                book.active.title = "核验结果"
            else:
                book = load_workbook(io.BytesIO(original(method, path, **kwargs).content))
                headers = [cell.value for cell in book["核验结果"][1]]
                book["核验结果"].cell(2, headers.index("来源 ID") + 1, 42)
            content = io.BytesIO()
            book.save(content)
            book.close()
            return httpx.Response(200, content=content.getvalue())
        return original(method, path, **kwargs)

    server.request = broken_export
    assert run.main() == 1
    record = json.loads((output / "dyn-wire-01.json").read_bytes())
    assert record["observation"]["run_status"] == "completed"
    assert record["observation"]["surface_errors"] == ["xlsx_artifact_unavailable_or_invalid", "json_xlsx_export_mismatch"]
    assert record["observation"]["xlsx_export_cells"] == []
    assert record["outcome_unknown"] is False
    assert json.loads((output / "dyn-wire-01.inflight.json").read_bytes())["journal_state"] == "finished"


def test_frozen_gold_change_preserves_not_started_attempt_and_denominator(tmp_path, monkeypatch):
    run, output, server, _, _, _ = cli_fixture(tmp_path, monkeypatch, change_gold=True)
    assert run.main() == 1
    rows = json.loads((output / "observations.json").read_bytes())
    assert len(rows) == 1 and rows[0]["run_status"] == "not_started"
    assert rows[0]["error_code"] == "evaluation_frozen_configuration_changed"
    assert server.requests == []


@pytest.mark.parametrize("change_index", ["delete", "tamper"])
def test_frozen_index_change_preserves_not_started_attempt(tmp_path, monkeypatch, change_index):
    run, output, server, _, _, _ = cli_fixture(tmp_path, monkeypatch, change_index=change_index)
    assert run.main() == 1
    rows = json.loads((output / "observations.json").read_bytes())
    assert len(rows) == 1 and rows[0]["run_status"] == "not_started"
    assert rows[0]["error_code"] == "evaluation_frozen_index_changed"
    assert server.requests == []


def test_index_is_rechecked_before_each_pack_after_a_completed_attempt(tmp_path, monkeypatch):
    run, output, server, _, _, directory = cli_fixture(tmp_path, monkeypatch)
    protocol = run.prepare_protocol(directory.parent.parent)
    protocol["tasks"].append({**protocol["tasks"][0], "task_id": "dyn-wire-02"})
    original_request = server.request

    def change_after_export(method, path, **kwargs):
        response = original_request(method, path, **kwargs)
        if method == "GET" and path.endswith("/export?format=xlsx"):
            (directory.parent.parent / "index.json").write_text('{"packs": []}', encoding="utf-8")
        return response

    server.request = change_after_export
    assert run.main() == 1
    rows = json.loads((output / "observations.json").read_bytes())
    assert rows[0]["run_status"] == "completed"
    assert rows[1]["run_status"] == "not_started"
    assert rows[1]["error_code"] == "evaluation_frozen_index_changed"
    assert sum(method == "POST" and path == "/api/v1/cases" for method, path, _ in server.requests) == 1


def test_wrapped_chinese_anchor_preserves_one_real_block_as_context_quote():
    anchor = {"document": "requirements.pdf", "role": "requirement", "supplier_handle": None,
              "item_id": "fruit", "text": "小时配送：从供应商接到订单开始计算，需人工核对起算事件和小时口径，不能无依据转换成自然日。", "page": 1}
    source_text = "合成开发评估资料：内容均为虚构。\n" + anchor["text"].replace("转换成", "转\n换成")
    blocks = [{"id": "7" * 32, "document_id": "6" * 32, "ordinal": 0,
               "source_text": source_text, "page_idx": 0, "sheet": None, "cell": None, "bbox": None}]
    context = runner()._anchor_context(anchor, blocks)
    assert context == {"evidence_refs": ["7" * 32], "quote": source_text}
    assert anchor["text"] not in context["quote"]  # Actual PDF line breaks stay intact.


def test_split_blocks_cannot_be_joined_to_resolve_wrapped_anchor():
    anchor = {"document": "requirements.pdf", "role": "requirement", "supplier_handle": None,
              "item_id": "fruit", "text": "人工核对小时口径，不能无依据转换成自然日。", "page": 1}
    texts = ["人工核对小时口径，不能无依据转", "换成自然日。"]
    blocks = [{"id": character * 32, "source_text": text, "page_idx": 0, "sheet": None, "cell": None}
              for character, text in zip(("7", "8"), texts, strict=True)]
    with pytest.raises(ValueError, match="binding_anchor_unresolved"):
        runner()._anchor_context(anchor, blocks)


def test_duplicate_wrapped_anchors_on_same_page_are_unresolved():
    anchor = {"document": "requirements.pdf", "role": "requirement", "supplier_handle": None,
              "item_id": "fruit", "text": "小时口径不能转换成自然日。", "page": 1}
    blocks = [{"id": character * 32, "source_text": "小时口径不能转\n换成自然日。", "page_idx": 0,
               "sheet": None, "cell": None} for character in ("7", "8")]
    with pytest.raises(ValueError, match="binding_anchor_unresolved"):
        runner()._anchor_context(anchor, blocks)


@pytest.mark.parametrize("location", [{"page": 1}, {"sheet": "采购资料", "cell": "A3"}])
def test_parser_normalized_punctuation_uses_real_pdf_or_xlsx_context(location):
    filename = "requirements.pdf" if "page" in location else "requirements.xlsx"
    anchor = {"document": filename, "role": "requirement", "supplier_handle": None,
              "item_id": "phone", "text": "品目：苹果手机；型号：DEMO-P256", **location}
    actual = "合成开发评估资料：内容均为虚构。\n品目:苹果手机；型号:DEMO-P256"
    block = {"id": "7" * 32, "source_text": actual,
             "page_idx": 0 if "page" in location else None,
             "sheet": location.get("sheet"), "cell": location.get("cell"), "bbox": None}
    context = runner()._anchor_context(anchor, [block])
    assert context == {"evidence_refs": ["7" * 32], "quote": actual}
    assert "型号:DEMO-P256" in context["quote"]
    assert anchor["text"] not in context["quote"]


def test_normalized_punctuation_does_not_choose_between_equivalent_blocks():
    anchor = {"document": "requirements.pdf", "role": "requirement", "supplier_handle": None,
              "item_id": "phone", "text": "品目：苹果手机；型号：DEMO-P256", "page": 1}
    blocks = [{"id": character * 32, "source_text": text, "page_idx": 0, "sheet": None, "cell": None}
              for character, text in zip(("7", "8"), (anchor["text"], "品目:苹果手机；型号:DEMO-P256"), strict=True)]
    with pytest.raises(ValueError, match="binding_anchor_unresolved"):
        runner()._anchor_context(anchor, blocks)


def test_unknown_project_creation_preserves_unstarted_tasks(tmp_path, monkeypatch):
    run, output, server, _, projects, _ = cli_fixture(tmp_path, monkeypatch, project_fails=True)
    assert run.main() == 1
    record = json.loads((output / "dyn-wire-01.json").read_bytes())
    assert record["stage"] == "not_started"
    assert record["error_code"] == "project_creation_failed_or_unknown"
    assert len(projects) == 1 and server.requests == []


def test_prepare_freezes_12_packs_104_positions_without_credentials(tmp_path, monkeypatch):
    from backend.tests.test_candidate_corpus import FONT
    from evals.dynamic.corpus import build_corpus

    font = FONT
    if not font.is_file():
        pytest.skip("Chinese font required to generate local originals")
    root = tmp_path / "corpus"
    build_corpus(root, font_path=font)
    run = runner()
    monkeypatch.setattr(run, "os", SimpleNamespace(environ=SimpleNamespace(
        get=lambda *a, **k: pytest.fail("freeze_credentials"))))
    protocol = run.prepare_protocol(root)
    assert (protocol["planned_tasks"], protocol["planned_positions"], protocol["planned_files"]) == (12, 104, 58)
    assert protocol["automatic_retries"] is False and protocol["track"] == "formal_review"
    assert protocol["candidate_generation_measured"] is False
    for task in protocol["tasks"]:
        assert task["gold_sha256"] == sha256((root / "gold" / f"{task['pack_id']}.json").read_bytes()).hexdigest()
        assert task["manifest_sha256"] == sha256((root / "config" / f"{task['pack_id']}.json").read_bytes()).hexdigest()
        assert "gold" not in task["manifest"] and "cells" not in task["manifest"]
    with pytest.raises(ValueError, match="selection_invalid"):
        run.prepare_protocol(root, packs=["dyn-d01", "dyn-d01"])


def test_real_fastapi_routes_accept_uploads_buyer_context_binding_and_review(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine

    from backend.app.api.main import create_app
    from backend.app.db.ingestion import DocumentIngestion
    from backend.app.db.models import metadata
    from backend.tests.test_api_workspace_authorization import DemoAuthenticator

    directory, manifest = input_pack(tmp_path)
    storage = tmp_path / "storage"
    monkeypatch.setenv("DOCUCHECK_DATA_DIR", str(storage))
    engine = create_engine(f"sqlite:///{tmp_path / 'wire.db'}")
    metadata.create_all(engine)
    client = TestClient(create_app(engine, authenticator=DemoAuthenticator()),
                        headers={"Authorization": "Bearer alice"})
    project_id = client.post("/api/v1/projects", json={"name": "Offline dynamic wire fixture"}).json()["id"]
    ingestion = DocumentIngestion(engine, storage)
    requests = []

    class ParsedClient:
        def request(self, method, path, **kwargs):
            if method == "GET" and path.endswith("/documents"):
                scope = client.get(path).json()
                ingestion.process_snapshot(scope["case_id"], scope["snapshot_id"])
            kwargs.pop("timeout", None)  # In-process TestClient has no transport deadline.
            response = client.request(method, path, **kwargs)
            requests.append((method, path, response.status_code))
            return response

    try:
        clock = Clock()
        result = runner().run_pack(ParsedClient(), directory, manifest=manifest, project_id=project_id,
                                   deadline_seconds=2, clock=clock.now, sleep=clock.sleep)
        assert result["error_code"] == "run_timeout", result["raw"][-1]
        assert result["run_id"] and result["attempt_id"]
        assert any(method == "PUT" and path.endswith("/item-bindings") and status == 200
                   for method, path, status in requests)
        assert sum(path.endswith("/messages") for _, path, _ in requests) == 1
        assert not any(path.endswith(("/retry", "/resume")) for _, path, _ in requests)
    finally:
        client.close()
        engine.dispose()
