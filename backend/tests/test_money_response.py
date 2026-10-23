"""Money evidence keeps its separate identity in public responses and downloads."""

import json
from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from backend.app.api.run_export import render_json, render_xlsx
from backend.app.api.schemas import CheckOut


def money_check():
    return {
        "check_result_id": "check-money", "snapshot_id": "snapshot-money",
        "supplier": "vendor-a", "field_key": "price", "processing_status": "complete",
        "business_status": "conflict", "reason_code": "tax_basis_conflict",
        "evidence_refs": ["price-row", "tax-row"], "reason_evidence_refs": ["tax-row"],
        "fact_ids": ["a" * 64], "rule_version": "dynamic-rules-v13", "is_formal": True,
        "criterion_id": "price", "criterion_name": "单价", "item_id": "apple",
        "item_name": "苹果", "checklist_version_id": "checklist-money",
        "tax_source_groups": [{"fact_id": "a" * 64, "evidence_refs": ["tax-row"]}],
    }


def money_export():
    return {
        "schema_version": 1,
        "result": {
            "run_id": "run-money", "case_id": "case-money", "snapshot_id": "snapshot-money",
            "revision": 2, "run_status": "completed", "is_current": True,
            "checklist_version_id": "checklist-money", "checks": [money_check()],
        },
        "sources": [
            {"evidence_id": "price-row", "document_id": "quote-money",
             "document_name": "报价.xlsx", "page": None, "sheet": "报价", "cell": "A2",
             "source_text": "苹果 不含税单价25元/箱"},
            {"evidence_id": "tax-row", "document_id": "quote-money",
             "document_name": "报价.xlsx", "page": None, "sheet": "报价", "cell": "A4",
             "source_text": "税口径说明：本表单价均为含税价。\n由供应商确认。"},
        ],
    }


@pytest.mark.parametrize("field", ["tax_source_groups", "price_source_groups"])
def test_money_groups_survive_response_filtering_without_unrelated_fields(field):
    check = money_check()
    check[field] = [{"fact_id": "a" * 64, "evidence_refs": ["source-row"]}]
    check["private_context"] = {"credential": "not-public"}
    check[field][0]["private_context"] = "not-public"
    encoded = CheckOut.model_validate(check).model_dump(mode="json")
    assert encoded[field] == [
        {"fact_id": "a" * 64, "evidence_refs": ["source-row"]},
    ]
    assert "private_context" not in encoded


@pytest.mark.parametrize("field", ["tax_source_groups", "price_source_groups"])
@pytest.mark.parametrize("group", [
    {"fact_id": "invalid", "evidence_refs": ["tax-row"]},
    {"fact_id": "a" * 64, "evidence_refs": []},
    {"fact_id": "a" * 64, "evidence_refs": [f"tax-{i}" for i in range(5)]},
])
def test_money_group_response_keeps_bounded_source_contract(field, group):
    with pytest.raises(ValidationError):
        CheckOut.model_validate(money_check() | {field: [group]})


def test_legacy_response_has_no_empty_tax_group_field():
    check = money_check()
    del check["tax_source_groups"]
    check["rule_version"] = "dynamic-rules-v12"
    encoded = CheckOut.model_validate(check).model_dump(mode="json")
    assert "tax_source_groups" not in encoded
    assert "price_source_groups" not in encoded


def test_json_export_keeps_tax_groups_and_separate_complete_sources():
    payload = money_export()
    encoded = json.loads(render_json(payload))
    assert encoded == payload
    assert encoded["result"]["checks"][0]["tax_source_groups"][0]["evidence_refs"] == ["tax-row"]
    assert len(encoded["sources"]) == 2
    assert encoded["sources"][1]["source_text"] == payload["sources"][1]["source_text"]


def test_xlsx_export_has_distinct_tax_groups_and_complete_original_sources():
    payload = money_export()
    # A report may contain material checks as well as money checks. Each group
    # keeps its own semantic label and never becomes a combined source span.
    material = deepcopy(payload["result"]["checks"][0])
    material.update({
        "check_result_id": "check-material", "criterion_id": "report",
        "evidence_refs": ["body-row"], "reason_evidence_refs": [], "fact_ids": ["b" * 64],
        "material_source_groups": [{"fact_id": "b" * 64, "evidence_refs": ["body-row"]}],
    })
    del material["tax_source_groups"]
    payload["result"]["checks"].append(material)
    payload["sources"].append({
        "evidence_id": "body-row", "document_id": "report-document",
        "document_name": "检测报告.xlsx", "page": None, "sheet": "报告", "cell": "A2",
        "source_text": "苹果检测报告：检验结果为合格。",
    })
    book = load_workbook(BytesIO(render_xlsx(payload)))
    tax = book["税口径引用分组"]
    assert list(tax.values) == [
        ("核验结果 ID", "指标 ID", "品目 ID", "事实 ID", "税说明来源 ID"),
        ("check-money", "price", "apple", "a" * 64, "tax-row"),
    ]
    assert book["材料引用分组"]["E2"].value == "body-row"
    assert tax.freeze_panes == "A2"
    assert tax.auto_filter.ref == "A1:E2"
    assert list(book["原文来源"].values) == [
        ("来源 ID", "文档 ID", "原件文件名", "PDF 页码", "工作表", "单元格", "原文"),
        ("price-row", "quote-money", "报价.xlsx", None, "报价", "A2",
         payload["sources"][0]["source_text"]),
        ("tax-row", "quote-money", "报价.xlsx", None, "报价", "A4",
         payload["sources"][1]["source_text"]),
        ("body-row", "report-document", "检测报告.xlsx", None, "报告", "A2",
         payload["sources"][2]["source_text"]),
    ]


def test_legacy_xlsx_export_has_no_tax_group_sheet():
    payload = money_export()
    del payload["result"]["checks"][0]["tax_source_groups"]
    payload["result"]["checks"][0]["rule_version"] = "dynamic-rules-v12"
    book = load_workbook(BytesIO(render_xlsx(payload)))
    assert "税口径引用分组" not in book.sheetnames
    assert "单价引用分组" not in book.sheetnames


@pytest.mark.parametrize("with_tax", [False, True])
def test_price_sources_export_in_independent_sheet_with_complete_original(with_tax):
    payload = money_export()
    check = payload["result"]["checks"][0]
    check["price_source_groups"] = [{"fact_id": "a" * 64, "evidence_refs": ["other-price-row"]}]
    check["evidence_refs"].append("other-price-row")
    payload["sources"].append({
        "evidence_id": "other-price-row", "document_id": "quote-money",
        "document_name": "报价.xlsx", "page": None, "sheet": "报价", "cell": "A5",
        "source_text": "苹果 不含税单价26元/箱",
    })
    if not with_tax:
        del check["tax_source_groups"]
        check["evidence_refs"].remove("tax-row")
        check["reason_evidence_refs"] = ["other-price-row"]
        payload["sources"] = [source for source in payload["sources"] if source["evidence_id"] != "tax-row"]
    assert json.loads(render_json(payload)) == payload
    book = load_workbook(BytesIO(render_xlsx(payload)))
    price = book["单价引用分组"]
    assert list(price.values) == [
        ("核验结果 ID", "指标 ID", "品目 ID", "事实 ID", "单价来源 ID"),
        ("check-money", "price", "apple", "a" * 64, "other-price-row"),
    ]
    assert price.freeze_panes == "A2"
    assert price.auto_filter.ref == "A1:E2"
    if with_tax:
        assert book["税口径引用分组"]["E2"].value == "tax-row"
        assert book["税口径引用分组"].max_row == 2
    else:
        assert "税口径引用分组" not in book.sheetnames
    original_sources = {row[0]: (row[5], row[6]) for row in list(book["原文来源"].values)[1:]}
    assert original_sources["price-row"] == ("A2", "苹果 不含税单价25元/箱")
    assert original_sources["other-price-row"] == ("A5", "苹果 不含税单价26元/箱")
