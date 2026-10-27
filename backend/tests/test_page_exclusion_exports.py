# ruff: noqa: F811 -- Shared pytest fixtures.
"""Frozen exclusions explain coverage in typed reports and literal exports."""

import io
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from backend.app.api.main import create_app
from backend.app.db.dynamic_repository import DynamicRepository
from backend.app.db.models import agent_runs
from backend.app.db.run_executor import execute_run
from backend.tests.test_dynamic_agent import ScriptProvider, call
from backend.tests.test_index_repository import indexed_db  # noqa: F401
from backend.tests.test_item_binding_api import put
from backend.tests.test_item_binding_runs import submit
from backend.tests.test_page_exclusion_repository import page_binding_db  # noqa: F401


class CapacityProvider(ScriptProvider):
    def extract(self, messages, schema, **options):
        data = json.loads(messages[-1].content)
        self.calls.append(("extract", data))
        block = next(b for b in data["evidence"] if b["text"] == "存储容量 512 GB")
        return call(schema["function"]["name"], {"facts": [{
            "criterion_id": "capacity", "item_id": "phone", "supplier": data["supplier"],
            "value": "512", "unit": "GB", "evidence_refs": [block["evidence_id"]], "quote": block["text"],
        }]})


def run(engine, key, monkeypatch, *, legacy=False):
    receipt = submit(engine, key)
    provider = CapacityProvider()
    with monkeypatch.context() as patch:
        if legacy:
            patch.setattr("backend.app.agent.dynamic.BINDING_RULE_VERSION", "dynamic-rules-v15")
        result = execute_run(engine, receipt.run_id, receipt.attempt_id, provider, heartbeat_interval_seconds=60)
    assert result.status == ("partial" if legacy else "completed"), result
    return receipt.run_id, provider


def test_new_scope_reports_and_exports_preserve_legacy_run(page_binding_db, monkeypatch):
    engine, body = page_binding_db
    client = TestClient(create_app(engine))
    old = put(engine, {k: v for k, v in body.items() if k not in {"exclusions", "exclusions_confirmed"}})
    old_id, _ = run(engine, "legacy", monkeypatch, legacy=True)
    old_payload = DynamicRepository(engine).verified(old_id)[1]
    old_export = client.get(f"/api/v1/agent-runs/{old_id}/export").json()
    assert "scope_exclusions" not in old_export["result"]
    assert "exclusions" not in old_export["item_attribution"]
    old_book = client.get(f"/api/v1/agent-runs/{old_id}/export?format=xlsx")
    assert "范围排除" not in load_workbook(io.BytesIO(old_book.content)).sheetnames

    current = put(engine, body | {"expected_binding_revision": 1}, key="explicit")["current_version"]
    new_id, provider = run(engine, "scoped", monkeypatch)
    endpoint = f"/api/v1/agent-runs/{new_id}"
    checks = client.get(endpoint + "/checks")
    assert checks.status_code == 200, checks.text
    checks = checks.json()
    assert checks["checks"][0]["business_status"] == "满足" and checks["checks"][0]["is_formal"]
    summary = checks["scope_exclusions"]
    expected = current["payload"]["exclusions"][0]
    assert len(summary) == 1
    assert summary[0] == {
        "document_id": "spec-A", "document_name": "供应商规格.pdf", "supplier": "A", "page": 2,
        "other_sku": "DEMO-OTHER", "reason": expected["reason"], "evidence_refs": expected["evidence_refs"],
        "sku_anchor": expected["sku_anchor"], "confirmed_by": current["confirmed_by"],
        "confirmed_at": current["confirmed_at"], "binding_version_id": current["id"],
        "sources": summary[0]["sources"],
    }
    assert {s["evidence_id"] for s in summary[0]["sources"]} == set(expected["evidence_refs"])
    assert any(s["source_text"] == "含税单价 6800 CNY/台" for s in summary[0]["sources"])
    extracted = {b["evidence_id"] for call in provider.calls if call[0] == "extract" for b in call[1]["evidence"]}
    assert extracted and not extracted.intersection(expected["evidence_refs"])
    with engine.connect() as connection:
        answer = connection.scalar(select(agent_runs.c.answer_text).where(agent_runs.c.id == new_id))
    assert "在已确认范围内" in answer and "供应商规格.pdf" in answer and "第2页" in answer
    assert "不作为已核验事实" in answer
    exported = client.get(endpoint + "/export")
    assert exported.status_code == 200, exported.text
    payload = exported.json()
    assert payload["result"] == checks
    assert payload["item_attribution"]["exclusions"] == summary
    book_response = client.get(endpoint + "/export?format=xlsx")
    assert book_response.status_code == 200, book_response.text
    book = load_workbook(io.BytesIO(book_response.content))
    rows = list(book["范围排除"].values)
    assert len(rows) == 5  # one row per source, never concatenate the entire page
    assert all(row[0:5] == ("spec-A", "供应商规格.pdf", "A", "2", "DEMO-OTHER") for row in rows[1:])
    assert {row[-2] for row in rows[1:]} == set(expected["evidence_refs"])
    assert any(row[-1] == "含税单价 6800 CNY/台" for row in rows[1:])
    assert all(cell.data_type != "f" for row in book["范围排除"] for cell in row)

    put(engine, body | {"expected_binding_revision": 2, "exclusions": []}, key="revoke")
    historical = client.get(endpoint + "/export").json()
    assert historical["item_attribution"]["exclusions"] == summary
    assert not historical["result"]["is_current"]
    assert not any(c["is_formal"] for c in historical["result"]["checks"])
    assert DynamicRepository(engine).verified(old_id)[1] == old_payload
    legacy = client.get(f"/api/v1/agent-runs/{old_id}/export").json()
    assert legacy == old_export | {"result": old_export["result"] | {"is_current": False}}
    assert legacy["item_attribution"]["version"]["id"] == old["current_version"]["id"]
