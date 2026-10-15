# ruff: noqa: F811 -- Shared pytest fixtures.
"""Run submission freezes authority independently from the mutable binding head."""

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from backend.app.api.main import create_app
from backend.app.db.assistant_repository import AssistantRepository
from backend.app.db.dynamic_repository import DynamicRepository
from backend.app.db.models import assistant_runs, dynamic_runs
from backend.app.db.run_repository import RunRepository, SubmitRun
from backend.app.runtime.lifecycle import RunStatus
from backend.tests import test_document_ingestion
from backend.tests.test_index_repository import indexed_db  # noqa: F401
from backend.tests.test_item_binding_api import binding_db, put  # noqa: F401

ingestion_case = test_document_ingestion.ingestion_case


def submit(engine, key, workflow="auto"):
    client = TestClient(create_app(engine))
    conv = client.post("/api/v1/cases/case/conversations").json()["id"]
    return RunRepository(engine).submit(SubmitRun("case", conv, "snap", 1, "审查", "standard", key,
                                                workflow=workflow))


@pytest.mark.parametrize("workflow,table", [("auto", dynamic_runs), ("assistant", assistant_runs)])
def test_new_runs_freeze_confirmed_binding_at_submission(binding_db, workflow, table):
    engine, _, body = binding_db
    version = put(engine, body)["current_version"]
    run = submit(engine, workflow, workflow)
    with engine.connect() as conn:
        row = conn.execute(select(table).where(table.c.run_id == run.run_id)).mappings().one()
    assert row["item_binding_version_id"] == version["id"]
    assert row["item_binding_digest"] == version["digest"]


def test_assistant_waiting_run_keeps_original_binding_for_formal_review(binding_db):
    engine, _, body = binding_db
    first = put(engine, body)["current_version"]
    run = submit(engine, "assistant", "assistant")
    runs = RunRepository(engine)
    claim = runs.claim(run.run_id, run.attempt_id)
    runs.finish(run.run_id, claim.active_attempt_id, claim.fence, claim.snapshot_id,
                RunStatus.WAITING_USER, answer_text="等待", memory_json={})
    second = put(engine, body | {"entries": [], "expected_binding_revision": 1}, key="clear")["current_version"]
    assert first["id"] != second["id"]
    with engine.begin() as conn:
        old = conn.execute(select(assistant_runs).where(assistant_runs.c.run_id == run.run_id)).mappings().one()
        AssistantRepository.validate_binding(conn, old)
    # Formal registration must use this submission's frozen association, not the new head.
    resumed = runs.resume(run.run_id, "snap", 1, "resume", "继续审查")
    claim = runs.claim(run.run_id, resumed.attempt_id)
    formal = AssistantRepository(engine).register_formal(claim, "review")
    assert formal["item_binding_version_id"] == first["id"]
    assert formal["item_bindings"]["digest"] == first["digest"]
    assert not AssistantRepository(engine).current(run.run_id)
    assert not DynamicRepository(engine).current(run.run_id)


def test_stopped_report_cannot_claim_item_coverage_without_manifest():
    from backend.app.db.run_repository import RepositoryError
    from backend.tests.test_item_binding_review import sample

    definition, repo, frozen = sample()
    payload = {"rule_version": "dynamic-rules-v4", "facts": [], "coverage": {"A": False},
               "item_coverage": {"A": {"phone": True}}, "batch_manifest": [], "batch_coverage": []}
    with pytest.raises(RepositoryError, match="dynamic_source_changed"):
        DynamicRepository(None)._validate_manifest({"id": "run"}, payload, definition, repo, context=frozen)


def test_persisted_retriever_loads_run_frozen_scope_instead_of_head(binding_db):
    from backend.app.db.agent_snapshot import load_agent_snapshot
    from backend.app.db.persisted_retrieval import persisted_retriever

    engine, _, body = binding_db
    first = put(engine, body)["current_version"]
    receipt = submit(engine, "rag", "assistant")
    runs = RunRepository(engine)
    claim = runs.claim(receipt.run_id, receipt.attempt_id)
    runs.finish(claim.run_id, claim.active_attempt_id, claim.fence, claim.snapshot_id,
                RunStatus.WAITING_USER, answer_text="等待", memory_json={})
    put(engine, body | {"entries": [], "expected_binding_revision": 1}, key="clear-rag")
    resumed = runs.resume(claim.run_id, "snap", 1, "resume-rag", "查数量")
    claim = runs.claim(claim.run_id, resumed.attempt_id)
    retriever = persisted_retriever(engine, load_agent_snapshot(engine, "case", "snap"), claim, target=None)
    hits = retriever.search("数量", item_id="rice")
    assert hits and "20件" not in str(hits)
    assert retriever.last_diagnostics["item_binding_version_id"] == first["id"]
    assert retriever.last_diagnostics["degraded_reason"] == "vector_disabled"


def test_persisted_retriever_rechecks_frozen_binding_before_disclosing(binding_db):
    from sqlalchemy import update

    from backend.app.db.agent_snapshot import load_agent_snapshot
    from backend.app.db.persisted_retrieval import persisted_retriever

    engine, _, body = binding_db
    put(engine, body)
    receipt = submit(engine, "rag-tamper", "assistant")
    claim = RunRepository(engine).claim(receipt.run_id, receipt.attempt_id)
    retriever = persisted_retriever(engine, load_agent_snapshot(engine, "case", "snap"), claim, target=None)
    with engine.begin() as conn:
        conn.execute(update(assistant_runs).where(assistant_runs.c.run_id == claim.run_id).values(item_binding_digest="b"*64))
    with pytest.raises(ValueError, match="item_binding_source_changed"):
        retriever.search("数量", item_id="rice")


@pytest.mark.parametrize("revoke", [False, True])
def test_frozen_scope_ready_manifest_and_permission_recheck_after_embedding(binding_db, revoke):
    from sqlalchemy import delete, insert, update

    from backend.app.api.auth import Principal
    from backend.app.db.agent_snapshot import load_agent_snapshot
    from backend.app.db.index_repository import IndexRepository
    from backend.app.db.index_worker import process_index
    from backend.app.db.models import (
        cases,
        conversations,
        index_manifests,
        project_members,
    )
    from backend.app.db.persisted_retrieval import persisted_retriever
    from backend.app.db.workspace import WorkspaceRepository
    from backend.tests.test_hybrid_retrieval import Embedding, Store
    from backend.tests.test_index_repository import NOW, TARGET

    engine, _, body = binding_db
    body["entries"][0]["selector"]["start_row"] = 1  # The header is explicitly in this confirmed range.
    confirmed = put(engine, body)["current_version"]
    actor = Principal("local", "scope-owner")
    project = WorkspaceRepository(engine).create_project(actor, "范围检索测试")
    with engine.begin() as conn:
        conn.execute(update(cases).where(cases.c.id == "case").values(project_id=project["id"]))
        conn.execute(insert(conversations).values(id="scope-conv", case_id="case", created_at=NOW))
    model, store = Embedding(), Store()
    model.revision = TARGET.revision
    indexes = IndexRepository(engine)
    indexes.schedule("case", "snap", 1, target=TARGET)
    assert process_index(indexes, indexes.claim(TARGET), embedding_factory=lambda t: model,
                         store_factory=lambda t: store) == "ready"
    receipt = RunRepository(engine).submit(SubmitRun("case", "scope-conv", "snap", 1, "查数量", "standard", "scope-run",
                                                    actor.key, workflow="assistant"))
    claim = RunRepository(engine).claim(receipt.run_id, receipt.attempt_id)
    retriever = persisted_retriever(engine, load_agent_snapshot(engine, "case", "snap"), claim, target=TARGET,
                                   embedding_factory=lambda t: model, store_factory=lambda t: store)
    if revoke:
        original = model.embed
        def embed(texts):
            vectors = original(texts)
            with engine.begin() as conn:
                conn.execute(delete(project_members).where(project_members.c.project_id == project["id"]))
            return vectors
        model.embed = embed
        with pytest.raises(ValueError, match="source_access_changed"):
            retriever.search("什么时候到", item_id="rice")
        assert not store.scopes  # The post-network permission check runs before vector search/disclosure.
    else:
        hits = retriever.search("什么时候到", item_id="rice")
        assert hits and "20件" not in str(hits)
        assert retriever.last_diagnostics["retrieval_mode"] == "hybrid"
        assert retriever.last_diagnostics["item_binding_version_id"] == confirmed["id"]
        assert store.scopes[-1].chunk_ids and store.scopes[-1].generations
        with engine.connect() as conn:
            assert len(conn.execute(select(index_manifests)).all()) == 1


@pytest.mark.parametrize("historical_version", [None, "dynamic-rules-v8", "dynamic-rules-v9", "dynamic-rules-v10"])
def test_uploaded_phone_review_history_export_and_empty_scope_isolation(ingestion_case, monkeypatch, historical_version):
    from copy import deepcopy
    from io import BytesIO

    from openpyxl import Workbook

    from backend.app.db.ingestion import DocumentIngestion
    from backend.app.db.item_binding_repository import ItemBindingRepository
    from backend.app.db.run_executor import execute_run
    from backend.tests.item_binding_helpers import confirm_single_item_bindings
    from backend.tests.test_item_binding_review import Provider, sample

    client, engine, storage, case_id = ingestion_case
    book = Workbook()
    book.active.title = "报价"
    for cell, text in (("B3", "苹果手机"), ("B4", "数量10台"), ("B5", "交期订单确认后18个自然日内交付")):
        book.active[cell] = text
    output = BytesIO()
    book.save(output)
    uploaded = client.post(f"/api/v1/cases/{case_id}/documents", headers={"Idempotency-Key": "phone"},
                          data={"role": "quote", "supplier_handle": "vendor-a", "supplier_name": "供应商甲", "model_name": "苹果手机", "expected_revision": "0"},
                          files={"file": ("手机报价.xlsx", output.getvalue(), "application/octet-stream")})
    assert uploaded.status_code == 201, uploaded.text
    receipt = uploaded.json()
    snapshot = receipt["snapshot_id"]
    assert DocumentIngestion(engine, storage).process_document(case_id, snapshot, receipt["document_id"])["parse_status"] == "complete"
    definition = sample()[0].model_dump(mode="json")
    for n in range(8):
        definition["criteria"].append(deepcopy(definition["criteria"][-1]) | {"criterion_id": f"manual-{n}", "name": f"待核对事项{n}"})
    url = f"/api/v1/cases/{case_id}/snapshots/{snapshot}/checklist"
    assert client.put(url, headers={"Idempotency-Key": "requirements"}, json=definition | {
        "expected_revision": 1, "expected_checklist_revision": 0, "reason": "买方确认"}).status_code == 200
    confirmed = client.post(url + "/confirm", headers={"Idempotency-Key": "confirm"}, json={
        "criterion_ids": [c["criterion_id"] for c in definition["criteria"]], "expected_revision": 1,
        "expected_checklist_revision": 1, "confirmed": True, "reason": "明确确认"}).json()
    binding = confirm_single_item_bindings(engine, case_id, snapshot, confirmed)["current_version"]
    conversation = client.post(f"/api/v1/cases/{case_id}/conversations").json()["id"]
    runs = RunRepository(engine)
    submitted = runs.submit(SubmitRun(case_id, conversation, snapshot, 1, "审查", "thorough", "review"))
    with monkeypatch.context() as patch:
        if historical_version:
            patch.setattr("backend.app.agent.dynamic.BINDING_RULE_VERSION", historical_version)
        result = execute_run(engine, submitted.run_id, submitted.attempt_id, Provider(), heartbeat_interval_seconds=60)
    assert result.status == "completed", result
    if historical_version:
        old_payload = DynamicRepository(engine).verified(submitted.run_id)[1]
        assert old_payload["rule_version"] == historical_version
        current = runs.submit(SubmitRun(case_id, conversation, snapshot, 1, "再次审查", "thorough", "upgraded-review"))
        upgraded = execute_run(engine, current.run_id, current.attempt_id, Provider(), heartbeat_interval_seconds=60)
        assert upgraded.status == "completed", upgraded
        new_payload = DynamicRepository(engine).verified(current.run_id)[1]
        assert new_payload["rule_version"] == "dynamic-rules-v11"
        assert not {e["input_digest"] for e in old_payload["batch_manifest"]}.intersection(
            e["input_digest"] for e in new_payload["batch_manifest"])
        assert DynamicRepository(engine).verified(submitted.run_id)[1] == old_payload
    checks = DynamicRepository(engine).checks(submitted.run_id)
    assert len(checks["checks"]) == 11
    assert [c["business_status"] for c in checks["checks"][:2]] == ["满足", "不满足"]
    assert all(c["business_status"] == "待人工确认" for c in checks["checks"][2:])
    exported = client.get(f"/api/v1/agent-runs/{submitted.run_id}/export")
    assert exported.status_code == 200, exported.text
    assert exported.json()["item_attribution"]["version"]["id"] == binding["id"]
    assert all(source["source_text"] != "苹果手机" for source in exported.json()["sources"])
    from openpyxl import load_workbook

    xlsx = client.get(f"/api/v1/agent-runs/{submitted.run_id}/export?format=xlsx")
    assert xlsx.status_code == 200, xlsx.text
    report = load_workbook(BytesIO(xlsx.content))
    assert {"品目对应依据", "对应锚点原文"} <= set(report.sheetnames)
    assert report["品目对应依据"]["B1"].value == binding["id"]
    assert report["对应锚点原文"]["G2"].value == "苹果手机"
    progress = DynamicRepository(engine).progress(submitted.run_id)
    assert progress["item_binding_digest"] == binding["digest"]
    assert progress["item_coverage"] == {"vendor-a": {"phone": True}}
    facts = DynamicRepository(engine).facts(submitted.run_id)
    assert facts["item_binding_digest"] == binding["digest"]
    ItemBindingRepository(engine).put(case_id, snapshot, [], actor="a"*64,
        checklist_version_id=confirmed["confirmed_version"]["id"], expected_revision=1,
        expected_checklist_revision=2, expected_binding_revision=1, confirmed=True,
        idempotency_key="clear", reason="重新核对")
    old = DynamicRepository(engine).checks(submitted.run_id)
    assert not old["is_current"] and not any(c["is_formal"] for c in old["checks"])
    assert old["checks"][0]["business_status"] == "满足"
    from backend.app.criteria.facts import FactCandidate

    fact = facts["fact_sets"][0]["facts"][0]
    rejected = client.post(f"/api/v1/agent-runs/{submitted.run_id}/fact-corrections",
        headers={"Idempotency-Key": "historical-correction"}, json={"fact_id": fact["fact_id"],
        "candidate": {k: fact[k] for k in FactCandidate.model_fields}, "reason": "测试旧版本只读",
        "expected_revision": 1, "expected_payload_digest": facts["payload_digest"]})
    assert rejected.status_code == 409, rejected.text
    historical = client.get(f"/api/v1/agent-runs/{submitted.run_id}/export")
    assert historical.status_code == 200
    assert historical.json()["item_attribution"]["version"]["id"] == binding["id"]
    next_run = runs.submit(SubmitRun(case_id, conversation, snapshot, 1, "重新审查", "thorough", "new-scope"))
    provider = Provider()
    new_result = execute_run(engine, next_run.run_id, next_run.attempt_id, provider, heartbeat_interval_seconds=60)
    assert new_result.status == "partial" and not provider.batches
    assert DynamicRepository(engine).checks(next_run.run_id)["checks"][0]["business_status"] is None
