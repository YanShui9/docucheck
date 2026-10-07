"""Source-backed parsing checks for the synthetic development pack."""

import json
import re
import shutil
import sys
from hashlib import sha256
from importlib import import_module
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.worksheet.table import Table
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
    TextStringObject,
)
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

PACK = Path(__file__).resolve().parents[2] / "examples" / "dev-pack-001"


def parsing():
    return import_module("backend.app.parsing")


def test_mineru_blocks_keep_original_page_and_bbox() -> None:
    result = parsing().parse_mineru_middle(
        PACK / "parsed" / "requirements.middle.json", document_id="requirements"
    )

    assert result.status == "complete"
    assert result.total_pages == 1
    assert result.processed_pages == (0,)
    memory = next(block for block in result.blocks if "已装内存至少 16 GB" in block.text)
    assert memory.document_id == "requirements"
    assert memory.page_idx == 0
    assert memory.page_width_pt > 500
    assert memory.page_height_pt > 800
    assert len(memory.bbox) == 4
    assert all(0 <= value <= 1 for value in memory.bbox)


def test_xlsx_cells_keep_sheet_and_cell_coordinates() -> None:
    result = parsing().parse_xlsx(PACK / "vendor-a-quote.xlsx", document_id="vendor-a-quote")

    assert result.status == "complete"
    assert result.processed_sheets == ("报价",)
    cells = {(block.sheet, block.cell): block for block in result.blocks}
    assert cells[("报价", "B6")].text == "已装内存 16 GB"
    assert cells[("报价", "B9")].text == "合同签订后 10 个工作日"
    assert len({block.evidence_id for block in result.blocks}) == len(result.blocks)


def test_xlsx_entity_declaration_is_rejected_before_fact_extraction(tmp_path):
    malicious = tmp_path / "entity.xlsx"
    source = PACK / "vendor-a-quote.xlsx"
    with ZipFile(source) as original, ZipFile(malicious, "w", compression=ZIP_DEFLATED) as changed:
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = b'<!DOCTYPE worksheet [<!ENTITY canary "not-an-extracted-fact">]>' + content
            changed.writestr(item, content)
    # The stdlib import is only ParseError; actual XML parsing uses defusedxml.
    from defusedxml.common import DefusedXmlException
    # openpyxl may reject the entity while reading dimensions, before our own iterparse.
    with pytest.raises(ValueError) as raised:
        parsing().parse_xlsx(malicious, document_id="entity-probe")
    error = raised.value
    while error is not None and not isinstance(error, DefusedXmlException):
        error = error.__cause__
    assert isinstance(error, DefusedXmlException)


def _with_declared_xlsx_dimension(source: Path, target: Path, dimension: str) -> None:
    with ZipFile(source) as original, ZipFile(target, "w", compression=ZIP_DEFLATED) as changed:
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content, replacements = re.subn(
                    rb'(<dimension ref=")[^"]+("\s*/>)',
                    lambda match: match.group(1) + dimension.encode("ascii") + match.group(2),
                    content,
                    count=1,
                )
                assert replacements == 1
            changed.writestr(item, content)


@pytest.mark.parametrize(
    ("declared_dimension", "first_missing_cell"),
    [("A1", "A6"), ("B6", "A1")],
)
def test_xlsx_underreported_dimension_recovers_cells_and_marks_partial(
    tmp_path: Path, declared_dimension: str, first_missing_cell: str,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "报价"
    sheet["A6"] = "已装内存"
    sheet["B6"] = "16 GB"
    original = tmp_path / "original.xlsx"
    incorrect = tmp_path / "incorrect.xlsx"
    workbook.save(original)
    _with_declared_xlsx_dimension(original, incorrect, declared_dimension)

    result = parsing().parse_xlsx(incorrect, document_id="incorrect")

    cells = {(block.sheet, block.cell): block.text for block in result.blocks}
    assert cells[(sheet.title, "A6")] == "已装内存"
    assert cells[(sheet.title, "B6")] == "16 GB"
    assert result.status == "partial"
    assert any(f"dimension_mismatch:{first_missing_cell}" in issue
               for issue in result.coverage_issues)


@pytest.mark.parametrize("hidden_cell", ["AE1", "A1001"])
def test_xlsx_underreported_dimension_cannot_hide_out_of_limit_cells(
    tmp_path: Path, hidden_cell: str,
) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "报价"
    workbook.active[hidden_cell] = "超出首版范围"
    original = tmp_path / "original.xlsx"
    incorrect = tmp_path / "incorrect.xlsx"
    workbook.save(original)
    _with_declared_xlsx_dimension(original, incorrect, "A1")

    with pytest.raises(parsing().UnsupportedContentError, match="dimensions exceed"):
        parsing().parse_xlsx(incorrect, document_id="incorrect")


def test_mineru_partial_flag_is_not_reported_as_complete(tmp_path: Path) -> None:
    source = json.loads(
        (PACK / "parsed" / "requirements.middle.json").read_text(encoding="utf-8")
    )
    source["is_full_document"] = False
    parsed_path = tmp_path / "partial.json"
    parsed_path.write_text(json.dumps(source), encoding="utf-8")

    result = parsing().parse_mineru_middle(parsed_path, document_id="requirements")
    assert result.status == "partial"


def test_formula_cells_are_rejected_instead_of_treated_as_values(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active["B2"] = "=SUM(1, 2)"
    path = tmp_path / "formula.xlsx"
    workbook.save(path)

    with pytest.raises(parsing().UnsupportedContentError, match="formula"):
        parsing().parse_xlsx(path, document_id="formula")


def test_merged_xlsx_cells_do_not_claim_complete_coverage(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A6"] = "已装内存"
    sheet.merge_cells("A6:A7")
    sheet["B6"] = "16 GB"
    sheet["B7"] = "32 GB"
    path = tmp_path / "merged.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id="merged")
    assert result.status == "partial"
    assert any("merged_cell" in issue for issue in result.coverage_issues)


def test_merged_xlsx_worksheet_at_relocated_part_is_still_partial(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A6"] = "已装内存"
    sheet.merge_cells("A6:A7")
    sheet["B6"] = "16 GB"
    sheet["B7"] = "32 GB"
    original = tmp_path / "original.xlsx"
    relocated = tmp_path / "relocated.xlsx"
    workbook.save(original)
    old_part = b"/xl/worksheets/sheet1.xml"
    new_part = b"/xl/custom/sheet1.xml"
    with ZipFile(original) as source, ZipFile(relocated, "w", compression=ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in {"xl/_rels/workbook.xml.rels", "[Content_Types].xml"}:
                data = data.replace(old_part, new_part)
            name = item.filename.replace("xl/worksheets/sheet1.xml", "xl/custom/sheet1.xml")
            target.writestr(name, data)

    result = parsing().parse_xlsx(relocated, document_id="relocated")
    assert result.status == "partial"
    assert any("merged_cell" in issue for issue in result.coverage_issues)


@pytest.mark.parametrize("kind", ["image", "chart"])
def test_xlsx_drawing_cannot_claim_complete(tmp_path: Path, kind: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "已装内存"
    sheet["B1"] = "16 GB"
    if kind == "image":
        image = Image.new("RGB", (220, 60), "white")
        ImageDraw.Draw(image).text((8, 20), "Installed RAM 32 GB", fill="black")
        image_bytes = BytesIO()
        image.save(image_bytes, format="PNG")
        image_bytes.seek(0)
        sheet.add_image(XlsxImage(image_bytes), "D4")
    else:
        sheet["C1"] = 16
        sheet["C2"] = 32
        chart = BarChart()
        chart.title = "Installed RAM"
        chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=2))
        sheet.add_chart(chart, "D4")
    path = tmp_path / f"{kind}.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id=kind)

    cells = {(block.sheet, block.cell): block.text for block in result.blocks}
    assert cells[("报价", "A1")] == "已装内存"
    assert cells[("报价", "B1")] == "16 GB"
    assert result.status == "partial"
    assert "sheet:报价:drawing_unverified:rId1" in result.coverage_issues


def test_xlsx_comment_cannot_claim_complete(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "已装内存"
    sheet["B1"] = "16 GB"
    sheet["B1"].comment = Comment("已装内存 32 GB", "供应商")
    path = tmp_path / "comment.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id="comment")

    cells = {(block.sheet, block.cell): block.text for block in result.blocks}
    assert cells[("报价", "A1")] == "已装内存"
    assert cells[("报价", "B1")] == "16 GB"
    assert result.status == "partial"
    assert "sheet:报价:legacy_drawing_unverified:anysvml" in result.coverage_issues


@pytest.mark.parametrize("relocated", [False, True])
def test_xlsx_comment_relationship_without_vml_cannot_claim_complete(
    tmp_path: Path, relocated: bool,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "已装内存 16 GB"
    sheet["A1"].comment = Comment("已装内存 32 GB", "供应商")
    original = tmp_path / "comment-original.xlsx"
    modified = tmp_path / "comment-only.xlsx"
    workbook.save(original)

    with ZipFile(original) as source, ZipFile(modified, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            if entry.filename == "xl/drawings/commentsDrawing1.vml":
                continue
            content = source.read(entry.filename)
            name = entry.filename
            if entry.filename == "xl/worksheets/sheet1.xml":
                content, removed = re.subn(rb"<legacyDrawing\b[^>]*/>", b"", content)
                assert removed == 1
            elif entry.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                content, removed = re.subn(
                    rb"<Relationship\b[^>]*vmlDrawing[^>]*/>", b"", content,
                )
                assert removed == 1
            if relocated:
                if entry.filename in {"xl/_rels/workbook.xml.rels", "[Content_Types].xml"}:
                    content = content.replace(
                        b"/xl/worksheets/sheet1.xml", b"/xl/custom/sheet1.xml",
                    )
                name = name.replace(
                    "xl/worksheets/_rels/sheet1.xml.rels",
                    "xl/custom/_rels/sheet1.xml.rels",
                ).replace("xl/worksheets/sheet1.xml", "xl/custom/sheet1.xml")
            target.writestr(name, content)

    readable = load_workbook(modified, read_only=False)
    try:
        assert readable["报价"]["A1"].comment.text == "已装内存 32 GB"
    finally:
        readable.close()
    result = parsing().parse_xlsx(modified, document_id="comment-only")

    assert result.status == "partial"
    assert "sheet:报价:comment_unverified:comments" in result.coverage_issues
    assert any(block.cell == "A1" and block.text == "已装内存 16 GB" for block in result.blocks)


def test_xlsx_threaded_comment_relationship_cannot_claim_complete(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "报价"
    workbook.active["A1"] = "已装内存 16 GB"
    original = tmp_path / "plain.xlsx"
    commented = tmp_path / "threaded-comment.xlsx"
    workbook.save(original)
    relation = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdThread1" '
        'Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
        'Target="../threadedComments/threadedComment1.xml"/>'
        '</Relationships>'
    )
    threaded_comments = (
        '<ThreadedComments '
        'xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments">'
        '<threadedComment ref="A1" id="{12345678-1234-1234-1234-123456789012}" '
        'personId="{12345678-1234-1234-1234-123456789013}">'
        '<text>已装内存 32 GB</text></threadedComment></ThreadedComments>'
    )
    with ZipFile(original) as source, ZipFile(commented, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "[Content_Types].xml":
                content = content.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/threadedComments/threadedComment1.xml" '
                    b'ContentType="application/vnd.ms-excel.threadedcomments+xml"/></Types>',
                )
            target.writestr(entry, content)
        target.writestr("xl/worksheets/_rels/sheet1.xml.rels", relation)
        target.writestr(
            "xl/threadedComments/threadedComment1.xml", threaded_comments,
        )

    result = parsing().parse_xlsx(commented, document_id="threaded-comment")

    assert any(block.cell == "A1" and block.text == "已装内存 16 GB" for block in result.blocks)
    assert result.status == "partial"
    assert "sheet:报价:threaded_comment_unverified:rIdThread1" in result.coverage_issues


def test_xlsx_header_text_cannot_claim_complete(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "已装内存 16 GB"
    sheet.oddHeader.center.text = "补充：已装内存 32 GB"
    path = tmp_path / "header-text.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id="header-text")

    assert any(block.cell == "A1" and block.text == "已装内存 16 GB" for block in result.blocks)
    assert result.status == "partial"
    assert "sheet:报价:header_footer_unverified:oddHeader" in result.coverage_issues


def test_xlsx_hyperlink_target_and_tooltip_cannot_claim_complete(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "规格详情"
    sheet["A1"].hyperlink = "https://example.invalid/spec?memory=32GB"
    sheet["A1"].hyperlink.tooltip = "实际已装内存 32 GB"
    path = tmp_path / "hyperlink.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id="hyperlink")

    assert any(block.cell == "A1" and block.text == "规格详情" for block in result.blocks)
    assert result.status == "partial"
    assert "sheet:报价:hyperlink_unverified:A1" in result.coverage_issues


def test_xlsx_referenced_table_relationship_cannot_claim_complete(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "参数"
    sheet["B1"] = "数值"
    sheet["A2"] = "已装内存"
    sheet["B2"] = "32 GB"
    sheet.add_table(Table(displayName="SupplementalDetails", ref="A1:B2"))
    path = tmp_path / "table-part.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id="table-part")

    assert any(block.cell == "B2" and block.text == "32 GB" for block in result.blocks)
    assert result.status == "partial"
    assert "sheet:报价:relationship_unverified:tablePart:rId1" in result.coverage_issues


def test_xlsx_unreferenced_table_part_does_not_change_coverage(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "参数"
    sheet["B1"] = "数值"
    sheet["A2"] = "已装内存"
    sheet["B2"] = "32 GB"
    sheet.add_table(Table(displayName="DetachedDetails", ref="A1:B2"))
    original = tmp_path / "table-original.xlsx"
    detached = tmp_path / "table-detached.xlsx"
    workbook.save(original)
    with ZipFile(original) as source, ZipFile(detached, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content, removed = re.subn(rb"<tableParts\b[^>]*>.*?</tableParts>", b"", content)
                assert removed == 1
            target.writestr(entry, content)

    result = parsing().parse_xlsx(detached, document_id="detached")

    assert result.status == "complete"
    assert result.coverage_issues == ()
    assert any(block.cell == "B2" and block.text == "32 GB" for block in result.blocks)


def test_xlsx_header_image_cannot_claim_complete(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "已装内存 16 GB"
    original = tmp_path / "plain.xlsx"
    annotated = tmp_path / "header-image.xlsx"
    workbook.save(original)

    image = Image.new("RGB", (220, 60), "white")
    ImageDraw.Draw(image).text((8, 20), "Installed RAM 32 GB", fill="black")
    picture = BytesIO()
    image.save(picture, format="PNG")
    sheet_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdHF1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" '
        'Target="/xl/drawings/vmlDrawing1.vml"/></Relationships>'
    )
    vml = (
        '<xml xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:o="urn:schemas-microsoft-com:office:office">'
        '<v:shapetype id="_x0000_t75" coordsize="21600,21600" o:spt="75"/>'
        '<v:shape id="LH" type="#_x0000_t75" style="width:220pt;height:60pt">'
        '<v:imagedata o:relid="rId1"/></v:shape></xml>'
    )
    vml_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
        'Target="../media/image1.png"/></Relationships>'
    )
    with ZipFile(original) as source, ZipFile(annotated, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                addition = (
                    b'<headerFooter><oddHeader>&amp;L&amp;G</oddHeader></headerFooter>'
                    b'<legacyDrawingHF xmlns:r="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships" r:id="rIdHF1"/></worksheet>'
                )
                content = content.replace(b"</worksheet>", addition)
            elif entry.filename == "[Content_Types].xml":
                addition = (
                    b'<Default Extension="vml" ContentType="application/'
                    b'vnd.openxmlformats-officedocument.vmlDrawing"/>'
                    b'<Default Extension="png" ContentType="image/png"/></Types>'
                )
                content = content.replace(b"</Types>", addition)
            target.writestr(entry, content)
        target.writestr("xl/worksheets/_rels/sheet1.xml.rels", sheet_rels)
        target.writestr("xl/drawings/vmlDrawing1.vml", vml)
        target.writestr("xl/drawings/_rels/vmlDrawing1.vml.rels", vml_rels)
        target.writestr("xl/media/image1.png", picture.getvalue())

    result = parsing().parse_xlsx(annotated, document_id="header-image")

    assert result.status == "partial"
    assert "sheet:Sheet:legacy_drawing_unverified:rIdHF1" in result.coverage_issues
    assert any(block.cell == "A1" and block.text == "已装内存 16 GB" for block in result.blocks)


@pytest.mark.parametrize(
    ("kind", "expected_issue"),
    [
        ("row", "sheet:报价:hidden_row:2"),
        ("column", "sheet:报价:hidden_columns:2-2"),
        ("column_range", "sheet:报价:hidden_columns:2-4"),
    ],
)
def test_xlsx_hidden_row_or_column_cannot_claim_complete(
    tmp_path: Path, kind: str, expected_issue: str,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "已装内存"
    sheet["B2"] = "32 GB"
    if kind == "row":
        sheet.row_dimensions[2].hidden = True
    elif kind == "column":
        sheet.column_dimensions["B"].hidden = True
    else:
        sheet.column_dimensions.group("B", "D", hidden=True)
    path = tmp_path / f"hidden-{kind}.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id=kind)

    cells = {(block.sheet, block.cell): block.text for block in result.blocks}
    assert cells[("报价", "A1")] == "已装内存"
    assert cells[("报价", "B2")] == "32 GB"
    assert result.status == "partial"
    assert expected_issue in result.coverage_issues


def test_xlsx_default_hidden_rows_cannot_claim_complete(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "已装内存"
    sheet.sheet_format.zeroHeight = True
    path = tmp_path / "default-hidden-rows.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id="default-hidden-rows")

    assert result.blocks[0].text == "已装内存"
    assert result.status == "partial"
    assert "sheet:报价:hidden_default_rows" in result.coverage_issues


@pytest.mark.parametrize(
    ("kind", "expected_issue"),
    [
        ("zero_height_row", "sheet:报价:hidden_row:2"),
        ("zero_width_column", "sheet:报价:hidden_columns:2-2"),
        ("zero_height_default_rows", "sheet:报价:hidden_default_rows"),
        ("zero_width_default_columns", "sheet:报价:hidden_default_columns"),
    ],
)
def test_xlsx_zero_sized_rows_or_columns_cannot_claim_complete(
    tmp_path: Path, kind: str, expected_issue: str,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "已装内存"
    sheet["B2"] = "32 GB"
    original = tmp_path / "original.xlsx"
    changed = tmp_path / f"{kind}.xlsx"
    workbook.save(original)
    with ZipFile(original) as source, ZipFile(changed, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                if kind == "zero_height_row":
                    content, count = re.subn(rb'(<row r="2")', rb'\1 ht="0"', content, count=1)
                elif kind == "zero_width_column":
                    content, count = re.subn(
                        rb"<sheetData>",
                        b'<cols><col min="2" max="2" width="0"/></cols><sheetData>',
                        content,
                        count=1,
                    )
                elif kind == "zero_height_default_rows":
                    content, count = re.subn(
                        rb'defaultRowHeight="[^"]+"', b'defaultRowHeight="0"',
                        content, count=1,
                    )
                else:
                    content, count = re.subn(
                        rb"<sheetFormatPr", b'<sheetFormatPr defaultColWidth="0"',
                        content, count=1,
                    )
                assert count == 1
            target.writestr(entry, content)

    result = parsing().parse_xlsx(changed, document_id=kind)

    cells = {(block.sheet, block.cell): block.text for block in result.blocks}
    assert cells[("报价", "B2")] == "32 GB"
    assert result.status == "partial"
    assert expected_issue in result.coverage_issues


@pytest.mark.parametrize("kind", ["row", "column"])
@pytest.mark.parametrize("hidden", [False, True])
def test_xlsx_negative_row_or_column_size_is_rejected(
    tmp_path: Path, kind: str, hidden: bool,
) -> None:
    workbook = Workbook()
    workbook.active["B2"] = "32 GB"
    original = tmp_path / "original.xlsx"
    changed = tmp_path / f"negative-{kind}-{hidden}.xlsx"
    workbook.save(original)
    with ZipFile(original) as source, ZipFile(changed, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                if kind == "row":
                    replacement = rb'\1 ht="-1" hidden="1"' if hidden else rb'\1 ht="-1"'
                    content, count = re.subn(rb'(<row r="2")', replacement, content, count=1)
                else:
                    hidden_attribute = b' hidden="1"' if hidden else b""
                    content, count = re.subn(
                        rb"<sheetData>",
                        b'<cols><col min="2" max="2" width="-1"'
                        + hidden_attribute + b"/></cols><sheetData>",
                        content,
                        count=1,
                    )
                assert count == 1
            target.writestr(entry, content)

    with pytest.raises(parsing().UnsupportedContentError, match="size is invalid"):
        parsing().parse_xlsx(changed, document_id=kind)


def test_xlsx_visible_row_and_column_remain_complete(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["B2"] = "32 GB"
    sheet.row_dimensions[2].hidden = False
    sheet.row_dimensions[2].height = 20
    sheet.column_dimensions["B"].hidden = False
    sheet.column_dimensions["B"].width = 12
    sheet.sheet_format.zeroHeight = False
    sheet.sheet_format.defaultColWidth = 10
    path = tmp_path / "visible.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id="visible")

    assert result.status == "complete"
    assert result.coverage_issues == ()
    assert any(block.cell == "B2" and block.text == "32 GB" for block in result.blocks)


@pytest.mark.parametrize("sheet_state", ["hidden", "veryHidden"])
def test_xlsx_hidden_worksheet_cannot_claim_complete(
    tmp_path: Path, sheet_state: str,
) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "报价首页"
    secret = workbook.create_sheet("补充参数")
    secret["B2"] = "已装内存 32 GB"
    secret.sheet_state = sheet_state
    path = tmp_path / f"{sheet_state}.xlsx"
    workbook.save(path)

    result = parsing().parse_xlsx(path, document_id=sheet_state)

    assert any(
        block.sheet == "补充参数" and block.cell == "B2" and block.text == "已装内存 32 GB"
        for block in result.blocks
    )
    assert result.status == "partial"
    assert f"sheet:补充参数:hidden_sheet:{sheet_state}" in result.coverage_issues


def test_pdf_entrypoint_checks_source_and_runs_mineru(tmp_path: Path) -> None:
    source_json = PACK / "parsed" / "requirements.middle.json"
    fake_mineru = tmp_path / "fake_mineru.py"
    fake_mineru.write_text(
        "import shutil, sys\n"
        "from pathlib import Path\n"
        f"source = Path({str(source_json)!r})\n"
        "assert sys.argv[1] == 'parse'\n"
        "assert '--format' in sys.argv and 'middle_json' in sys.argv\n"
        "shutil.copyfile(source, sys.argv[sys.argv.index('-o') + 1])\n",
        encoding="utf-8",
    )

    result = parsing().parse_pdf(
        PACK / "requirements.pdf",
        document_id="requirements",
        mineru_command=[sys.executable, str(fake_mineru)],
    )

    assert result.status == "complete"
    assert any("采购数量 10 台" in block.text for block in result.blocks)


def _middle_without_memory(path: Path) -> None:
    source = json.loads((PACK / "parsed" / "requirements.middle.json").read_text(encoding="utf-8"))
    source["pages"][0]["blocks"] = [
        block for block in source["pages"][0]["blocks"]
        if not any("已装内存至少 16 GB" in item.get("content", "") for item in block["content"])
    ]
    path.write_text(json.dumps(source, ensure_ascii=False), encoding="utf-8")


def test_pdf_entrypoint_marks_silently_omitted_source_text_partial(tmp_path: Path) -> None:
    incomplete_middle = tmp_path / "omitted.middle.json"
    _middle_without_memory(incomplete_middle)
    fake_mineru = tmp_path / "fake_mineru.py"
    fake_mineru.write_text(
        "import shutil, sys\n"
        "from pathlib import Path\n"
        f"source = Path({str(incomplete_middle)!r})\n"
        "shutil.copyfile(source, sys.argv[sys.argv.index('-o') + 1])\n",
        encoding="utf-8",
    )

    result = parsing().parse_pdf(
        PACK / "requirements.pdf", document_id="requirements",
        mineru_command=[sys.executable, str(fake_mineru)],
    )

    assert result.status == "partial"
    assert any(issue.startswith("page:0:source_line:") for issue in result.coverage_issues)


def test_mixed_pdf_with_omitted_image_cannot_claim_complete(tmp_path: Path) -> None:
    from backend.app.evidence.repository import bind_document

    image = Image.new("RGB", (220, 60), "white")
    ImageDraw.Draw(image).text((8, 20), "Installed RAM 32 GB", fill="black")
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")
    image_bytes.seek(0)
    pdf_path = tmp_path / "mixed.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(612, 792))
    pdf.drawString(72, 720, "Supplier specification")
    pdf.drawImage(ImageReader(image_bytes), 72, 600, width=220, height=60)
    pdf.save()

    middle = json.loads((PACK / "parsed" / "requirements.middle.json").read_text(encoding="utf-8"))
    middle["pages"][0]["blocks"] = [{
        "type": "text", "index": 0, "bbox": [0.1, 0.1, 0.5, 0.2],
        "content": [{"type": "text", "content": "Supplier specification"}],
    }]
    middle_path = tmp_path / "mixed.middle.json"
    middle_path.write_text(json.dumps(middle), encoding="utf-8")
    fake_mineru = tmp_path / "fake_mineru.py"
    fake_mineru.write_text(
        "import shutil, sys\n"
        "from pathlib import Path\n"
        f"source = Path({str(middle_path)!r})\n"
        "shutil.copyfile(source, sys.argv[sys.argv.index('-o') + 1])\n",
        encoding="utf-8",
    )

    parsed = parsing().parse_pdf(
        pdf_path, document_id="mixed", mineru_command=[sys.executable, str(fake_mineru)]
    )
    cached = parsing().parse_mineru_middle(middle_path, document_id="mixed")
    bound = bind_document(cached, source_path=pdf_path, role="spec", supplier="vendor-a")

    assert parsed.status == "partial"
    assert "page:0:embedded_image_unverified" in parsed.coverage_issues
    assert bound.parsed.status == "partial"
    assert "page:0:embedded_image_unverified" in bound.parsed.coverage_issues


@pytest.mark.parametrize("mode", ["direct", "form", "nested_form"])
def test_pdf_vector_table_lines_cannot_claim_complete(
    tmp_path: Path, mode: str,
) -> None:
    pdf_path = tmp_path / "vector-table.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(612, 792))
    pdf.drawString(72, 720, "Supplier specification")
    if mode in {"form", "nested_form"}:
        pdf.beginForm("table_lines")
        pdf.rect(0, 0, 220, 60, stroke=1, fill=0)
        pdf.line(0, 30, 220, 30)
        pdf.endForm()
        form_to_draw = "table_lines"
        if mode == "nested_form":
            pdf.beginForm("outer")
            pdf.doForm("table_lines")
            pdf.endForm()
            form_to_draw = "outer"
        pdf.saveState()
        pdf.translate(72, 600)
        pdf.doForm(form_to_draw)
        pdf.restoreState()
    else:
        pdf.rect(72, 600, 220, 60, stroke=1, fill=0)
        pdf.line(72, 630, 292, 630)
    pdf.save()

    middle = json.loads((PACK / "parsed" / "requirements.middle.json").read_text(encoding="utf-8"))
    middle["pages"][0]["blocks"] = [{
        "type": "text", "index": 0, "bbox": [0.1, 0.1, 0.5, 0.2],
        "content": [{"type": "text", "content": "Supplier specification"}],
    }]
    middle_path = tmp_path / "vector.middle.json"
    middle_path.write_text(json.dumps(middle), encoding="utf-8")

    parsed = parsing().parse_mineru_middle(middle_path, document_id="vector")
    result = parsing().verify_pdf_source_coverage(parsed, pdf_path)

    assert result.status == "partial"
    assert "page:0:vector_graphics_unverified" in result.coverage_issues


def test_pdf_borderless_text_columns_cannot_claim_complete(tmp_path: Path) -> None:
    pdf_path = tmp_path / "borderless-table.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(612, 792))
    for y, label, value in (
        (720, "Parameter", "Value"),
        (700, "Installed RAM", "32 GB"),
        (680, "Storage", "1 TB"),
    ):
        pdf.drawString(72, y, label)
        pdf.drawString(250, y, value)
    pdf.save()

    middle = json.loads((PACK / "parsed" / "requirements.middle.json").read_text(encoding="utf-8"))
    source_lines = PdfReader(pdf_path).pages[0].extract_text().splitlines()
    middle["pages"][0]["blocks"] = [
        {
            "type": "text", "index": index, "bbox": [0.1, 0.1, 0.8, 0.9],
            "content": [{"type": "text", "content": line}],
        }
        for index, line in enumerate(source_lines)
    ]
    middle_path = tmp_path / "borderless.middle.json"
    middle_path.write_text(json.dumps(middle), encoding="utf-8")

    parsed = parsing().parse_mineru_middle(middle_path, document_id="borderless")
    result = parsing().verify_pdf_source_coverage(parsed, pdf_path)

    assert result.status == "partial"
    assert "page:0:layout_columns_unverified" in result.coverage_issues
    assert not any("not_extracted" in issue for issue in result.coverage_issues)


def test_pdf_single_borderless_label_value_row_is_unverified(tmp_path: Path) -> None:
    pdf_path = tmp_path / "single-row.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(612, 792))
    pdf.drawString(72, 720, "Installed RAM")
    pdf.drawString(250, 720, "32 GB")
    pdf.save()

    middle = json.loads((PACK / "parsed" / "requirements.middle.json").read_text(encoding="utf-8"))
    middle["pages"][0]["blocks"] = [
        {
            "type": "text", "index": index, "bbox": [0.1, 0.1, 0.8, 0.9],
            "content": [{"type": "text", "content": line}],
        }
        for index, line in enumerate(PdfReader(pdf_path).pages[0].extract_text().splitlines())
    ]
    middle_path = tmp_path / "single-row.middle.json"
    middle_path.write_text(json.dumps(middle), encoding="utf-8")

    result = parsing().verify_pdf_source_coverage(
        parsing().parse_mineru_middle(middle_path, document_id="single-row"), pdf_path,
    )

    assert result.status == "partial"
    assert "page:0:layout_columns_unverified" in result.coverage_issues


def test_pdf_text_only_form_remains_complete(tmp_path: Path) -> None:
    pdf_path = tmp_path / "text-form.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(612, 792))
    pdf.drawString(72, 720, "Supplier specification")
    pdf.beginForm("text_only")
    pdf.drawString(0, 0, "Additional details")
    pdf.endForm()
    pdf.saveState()
    pdf.translate(72, 600)
    pdf.doForm("text_only")
    pdf.restoreState()
    pdf.save()

    middle = json.loads((PACK / "parsed" / "requirements.middle.json").read_text(encoding="utf-8"))
    middle["pages"][0]["blocks"] = [
        {"type": "text", "index": index, "bbox": [0.1, 0.1, 0.5, 0.2],
         "content": [{"type": "text", "content": value}]}
        for index, value in enumerate(("Supplier specification", "Additional details"))
    ]
    middle_path = tmp_path / "text-form.middle.json"
    middle_path.write_text(json.dumps(middle), encoding="utf-8")

    parsed = parsing().parse_mineru_middle(middle_path, document_id="text-form")
    result = parsing().verify_pdf_source_coverage(parsed, pdf_path)

    assert result.status == "complete"
    assert result.coverage_issues == ()


def test_unused_pdf_vector_form_does_not_change_coverage(tmp_path: Path) -> None:
    pdf_path = tmp_path / "unused-form.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(612, 792))
    pdf.drawString(72, 720, "Supplier specification")
    pdf.beginForm("unused_lines")
    pdf.rect(0, 0, 220, 60, stroke=1, fill=0)
    pdf.endForm()
    pdf.save()

    middle = json.loads((PACK / "parsed" / "requirements.middle.json").read_text(encoding="utf-8"))
    middle["pages"][0]["blocks"] = [{
        "type": "text", "index": 0, "bbox": [0.1, 0.1, 0.5, 0.2],
        "content": [{"type": "text", "content": "Supplier specification"}],
    }]
    middle_path = tmp_path / "unused-form.middle.json"
    middle_path.write_text(json.dumps(middle), encoding="utf-8")

    parsed = parsing().parse_mineru_middle(middle_path, document_id="unused-form")
    result = parsing().verify_pdf_source_coverage(parsed, pdf_path)

    assert result.status == "complete"
    assert result.coverage_issues == ()


@pytest.mark.parametrize("appearance", [True, False])
def test_pdf_annotation_cannot_claim_complete(tmp_path: Path, appearance: bool) -> None:
    source = BytesIO()
    pdf = canvas.Canvas(source, pagesize=(612, 792))
    pdf.drawString(72, 720, "Supplier specification")
    pdf.save()

    writer = PdfWriter()
    writer.append_pages_from_reader(PdfReader(BytesIO(source.getvalue())))
    annotation = DictionaryObject({
        NameObject("/Type"): NameObject("/Annot"),
        NameObject("/Subtype"): NameObject("/Stamp" if appearance else "/Text"),
        NameObject("/Rect"): ArrayObject([
            NumberObject(72), NumberObject(600), NumberObject(222), NumberObject(630),
        ]),
        NameObject("/F"): NumberObject(4),
    })
    if appearance:
        drawing = DecodedStreamObject()
        drawing.set_data(b"0 0 m 150 0 l S")
        drawing[NameObject("/Type")] = NameObject("/XObject")
        drawing[NameObject("/Subtype")] = NameObject("/Form")
        drawing[NameObject("/BBox")] = ArrayObject([
            NumberObject(0), NumberObject(0), NumberObject(150), NumberObject(30),
        ])
        annotation[NameObject("/AP")] = DictionaryObject({
            NameObject("/N"): writer._add_object(drawing),
        })
    else:
        annotation[NameObject("/Contents")] = TextStringObject("Installed RAM 32 GB")
    writer.pages[0][NameObject("/Annots")] = ArrayObject([writer._add_object(annotation)])
    pdf_path = tmp_path / "annotated.pdf"
    with pdf_path.open("wb") as output:
        writer.write(output)

    middle = json.loads((PACK / "parsed" / "requirements.middle.json").read_text(encoding="utf-8"))
    middle["pages"][0]["blocks"] = [{
        "type": "text", "index": 0, "bbox": [0.1, 0.1, 0.5, 0.2],
        "content": [{"type": "text", "content": "Supplier specification"}],
    }]
    middle_path = tmp_path / "annotated.middle.json"
    middle_path.write_text(json.dumps(middle), encoding="utf-8")

    parsed = parsing().parse_mineru_middle(middle_path, document_id="annotated")
    result = parsing().verify_pdf_source_coverage(parsed, pdf_path)

    assert result.status == "partial"
    assert "page:0:annotation:0:unverified" in result.coverage_issues


def test_pdf_source_coverage_does_not_reuse_one_match_for_repeated_lines(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    parsed = parsing().parse_mineru_middle(
        PACK / "parsed" / "requirements.middle.json", document_id="requirements"
    )

    class FakePage:
        images: tuple[()] = ()

        def get(self, key: str) -> None:
            return None

        def get_contents(self) -> None:
            return None

        def extract_text(self) -> str:
            return "已装内存至少 16 GB\n已装内存至少 16 GB\n"

    class FakeReader:
        def __init__(self, _source: Path) -> None:
            self.pages = [FakePage()]

    monkeypatch.setattr(parsing(), "PdfReader", FakeReader)
    result = parsing().verify_pdf_source_coverage(parsed, PACK / "requirements.pdf")

    assert result.status == "partial"
    assert "page:0:source_line:2:not_extracted" in result.coverage_issues


def test_dev_pack_cache_is_checked_against_original_pdf(tmp_path: Path) -> None:
    from backend.app.evidence.repository import load_case
    from backend.app.rules.memory import evaluate_memory

    case_dir = tmp_path / "case"
    shutil.copytree(PACK, case_dir)
    cached = case_dir / "parsed" / "requirements.middle.json"
    _middle_without_memory(cached)
    manifest_path = case_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["documents"][0]["parsed_sha256"] = sha256(cached.read_bytes()).hexdigest()
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    repo = load_case(case_dir)
    assert not repo.coverage("vendor-a")["complete"]
    requirement = next(document for document in repo.documents if document.role == "requirement")
    assert requirement.parsed.status == "partial"
    assert "/source-text:1" in requirement.parsed.parser_version
    assert any(issue.startswith("page:0:source_line:")
               for issue in requirement.parsed.coverage_issues)
    check = evaluate_memory(repo, "vendor-a", [], scan_complete=True)
    assert check.processing_status == "partial" and check.business_status is None


def test_all_labeled_evidence_can_be_located_in_parsed_sources() -> None:
    documents = {
        "requirements.pdf": parsing().parse_mineru_middle(
            PACK / "parsed" / "requirements.middle.json", document_id="requirements.pdf"
        ),
        "vendor-a-spec.pdf": parsing().parse_mineru_middle(
            PACK / "parsed" / "vendor-a-spec.middle.json", document_id="vendor-a-spec.pdf"
        ),
        "vendor-a-quote.xlsx": parsing().parse_xlsx(
            PACK / "vendor-a-quote.xlsx", document_id="vendor-a-quote.xlsx"
        ),
    }
    labels = json.loads((PACK / "labels.json").read_text(encoding="utf-8"))
    checked = 0
    for result in labels["results"]:
        for reference in result["requirement_evidence"] + result["supplier_evidence"]:
            parsed = documents[reference["document"]]
            assert parsed.status == "complete"
            assert any(
                reference["text"] in block.text
                and (
                    block.page_idx == reference["page"] - 1
                    if "page" in reference
                    else block.sheet == reference["sheet"] and block.cell == reference["cell"]
                )
                for block in parsed.blocks
            ), reference
            checked += 1
    assert checked == 16
