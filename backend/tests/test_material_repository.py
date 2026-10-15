"""Material body evidence survives SQL history, review, and snapshot exports."""

from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, update

from backend.app.agent.dynamic import review_plan
from backend.app.api.dynamic_schemas import DynamicCorrectionIn
from backend.app.criteria.facts import FactCandidate, evaluate_checklist, validate_fact
from backend.app.criteria.item_bindings import review_bindings
from backend.app.criteria.models import ChecklistDefinition
from backend.app.db.agent_snapshot import load_agent_snapshot
from backend.app.db.dynamic_repository import DynamicRepository, source_digest
from backend.app.db.ingestion import DocumentIngestion
from backend.app.db.models import dynamic_batches, dynamic_corrections, dynamic_runs
from backend.app.db.run_repository import RepositoryError, RunRepository, SubmitRun
from backend.app.evidence.repository import digest
from backend.app.runtime.lifecycle import RunStatus
from backend.tests import test_document_ingestion
from backend.tests.item_binding_helpers import confirm_single_item_bindings

ingestion_case = test_document_ingestion.ingestion_case


def correction_body(candidate):
    return {"fact_id": "a" * 64, "candidate": candidate, "reason": "原文复核",
            "expected_revision": 1, "expected_payload_digest": "b" * 64}


def test_material_correction_schema_accepts_independent_body_without_changing_legacy_payload():
    candidate = FactCandidate.model_validate({
        "criterion_id": "warranty", "item_id": "phone", "supplier": "vendor-a",
        "value": True, "unit": "", "evidence_refs": ["main"], "quote": "已附质保承诺书",
    }).model_dump(mode="json")
    legacy = correction_body(candidate)
    assert DynamicCorrectionIn.model_validate(legacy).model_dump(mode="json") == legacy
    material = correction_body(candidate | {"material_sources": [
        {"evidence_refs": ["body"], "quote": "质保承诺书：本产品保修一年。"}]})
    parsed = DynamicCorrectionIn.model_validate(material).model_dump(mode="json")
    assert parsed == material
    assert digest(parsed) == digest(material)


@pytest.fixture
def material_case(ingestion_case, request):
    client, engine, storage, case_id = ingestion_case
    for revision, (role, rows) in enumerate((
        ("quote", ["苹果手机", "检测报告 已提供检测报告"]),
        ("spec", ["苹果手机", "检测报告 检验结果为合格", "检测报告 检验结果为合格且已盖章"]),
    )):
        if role == "spec" and getattr(request, "param", None) == "large":
            rows = rows[:2] + ["普通附录说明" * 140 for _ in range(4)] + rows[2:]
        book = Workbook()
        book.active.title = "资料"
        for row, text in enumerate(rows, 1):
            book.active.cell(row, 1, text)
        output = BytesIO()
        book.save(output)
        response = client.post(f"/api/v1/cases/{case_id}/documents",
            headers={"Idempotency-Key": f"material-upload-{revision}"},
            data={"role": role, "supplier_handle": "vendor-a", "supplier_name": "供应商甲",
                  "expected_revision": str(revision)},
            files={"file": (f"{role}.xlsx", output.getvalue(), "application/octet-stream")})
        assert response.status_code == 201, response.text
        snapshot = response.json()["snapshot_id"]
    from backend.app.db.models import document_versions

    with engine.connect() as connection:
        ids = list(connection.scalars(select(document_versions.c.id).where(
            document_versions.c.snapshot_id == snapshot)))
    for document_id in ids:
        assert DocumentIngestion(engine, storage).process_document(case_id, snapshot, document_id)["parse_status"] == "complete"
    definition = ChecklistDefinition.model_validate({
        "items": [{"item_id": "phone", "name": "苹果手机"}],
        "criteria": [{"criterion_id": "report", "item_id": "phone", "name": "检测报告",
                      "rule": {"type": "presence", "description": "检测报告"},
                      "source": {"kind": "user_defined", "reason": "买方确认"}}],
    })
    url = f"/api/v1/cases/{case_id}/snapshots/{snapshot}/checklist"
    assert client.put(url, headers={"Idempotency-Key": "material-draft"}, json={
        **definition.model_dump(mode="json"), "expected_revision": 2,
        "expected_checklist_revision": 0, "reason": "采购要求"}).status_code == 200
    confirmed = client.post(url + "/confirm", headers={"Idempotency-Key": "material-confirm"}, json={
        "criterion_ids": ["report"], "expected_revision": 2, "expected_checklist_revision": 1,
        "confirmed": True, "reason": "确认"}).json()
    confirm_single_item_bindings(engine, case_id, snapshot, confirmed)
    conversation = client.post(f"/api/v1/cases/{case_id}/conversations").json()["id"]
    return client, engine, case_id, snapshot, conversation, definition


def persist_material_report(case, *, rule_version="dynamic-rules-v11", material=True):
    _, engine, case_id, snapshot, conversation, definition = case
    runs, store = RunRepository(engine), DynamicRepository(engine)
    receipt = runs.submit(SubmitRun(case_id, conversation, snapshot, 2, "审查", "standard",
                                   "material-" + rule_version, workflow="auto"))
    claim = runs.claim(receipt.run_id, receipt.attempt_id)
    frozen = store.binding(claim)["item_bindings"]
    repo = load_agent_snapshot(engine, case_id, snapshot)
    context = review_bindings(frozen, definition, repo)
    main = next(b for b in repo.blocks if b.text == "检测报告 已提供检测报告")
    body = next(b for b in repo.blocks if b.text == "检测报告 检验结果为合格")
    candidate = FactCandidate.model_validate({"criterion_id": "report", "item_id": "phone",
        "supplier": "vendor-a", "value": True, "unit": "", "evidence_refs": [main.evidence_id],
        "quote": main.text}).model_dump(mode="json")
    if material:
        candidate["material_sources"] = [{"evidence_refs": [body.evidence_id], "quote": body.text}]
    fact = validate_fact(candidate, definition, repo, "vendor-a", rule_version=rule_version, item_bindings=context)
    plans, _ = review_plan(repo, definition, rule_version=rule_version, item_bindings=context)
    manifest = [entry for batches in plans.values() for _, _, entry in batches]
    for batch, _, entry in plans["vendor-a"]:
        output = {"facts": [candidate] if main.evidence_id in {b["evidence_id"] for b in batch} else []}
        store.save_batch(claim, entry["batch_id"], entry["input_digest"], output)
    payload = {"version": 1, "rule_version": rule_version, "source_digest": source_digest(repo),
        "definition": definition.model_dump(mode="json"), "facts": [fact],
        "coverage": {"vendor-a": True}, "item_coverage": {"vendor-a": {"phone": True}},
        "batch_manifest": manifest,
        "batch_coverage": [{**entry, "complete": True} for entry in manifest],
        "needs_resolution": [], **context.identity()}
    payload["checks"] = evaluate_checklist(definition, payload["facts"], payload["coverage"],
        snapshot_id=snapshot, rule_version=rule_version, item_coverage=payload["item_coverage"], item_bindings=context)
    runs.finish(claim.run_id, claim.active_attempt_id, claim.fence, snapshot, RunStatus.COMPLETED,
                answer_text="完成", dynamic_payload=payload, memory_json={})
    return claim.run_id, payload, repo


def test_material_history_api_json_excel_and_original_locations(material_case):
    client, engine, case_id, snapshot, _, _ = material_case
    run_id, payload, repo = persist_material_report(material_case)
    assert DynamicRepository(engine).verified(run_id)[1] == payload
    facts = client.get(f"/api/v1/agent-runs/{run_id}/facts")
    assert facts.status_code == 200, facts.text
    assert facts.json()["fact_sets"][0]["facts"][0]["material_sources"] == payload["facts"][0]["material_sources"]
    exported = client.get(f"/api/v1/agent-runs/{run_id}/export")
    assert exported.status_code == 200, exported.text
    exported = exported.json()
    expected_groups = [{"fact_id": payload["facts"][0]["fact_id"],
                        "evidence_refs": payload["facts"][0]["material_sources"][0]["evidence_refs"]}]
    assert exported["result"]["checks"][0]["material_source_groups"] == expected_groups
    refs = {payload["facts"][0]["evidence_refs"][0], payload["facts"][0]["material_sources"][0]["evidence_refs"][0]}
    assert {s["evidence_id"] for s in exported["sources"]} == refs
    assert set(exported["result"]["checks"][0]["evidence_refs"]) == refs
    for source in exported["sources"]:
        evidence = client.get(f"/api/v1/cases/{case_id}/snapshots/{snapshot}/evidence/{source['evidence_id']}")
        assert evidence.status_code == 200, evidence.text
        assert evidence.json()["source_text"] == source["source_text"]
        original = client.get(f"/api/v1/cases/{case_id}/snapshots/{snapshot}/documents/{source['document_id']}/original")
        assert original.status_code == 200, original.text
        assert source["sheet"] == "资料" and source["cell"] == "A2"
    excel = client.get(f"/api/v1/agent-runs/{run_id}/export?format=xlsx")
    assert excel.status_code == 200, excel.text
    book = load_workbook(BytesIO(excel.content))
    assert "材料引用分组" in book.sheetnames
    assert book["材料引用分组"]["D2"].value == payload["facts"][0]["fact_id"]
    assert book["材料引用分组"]["E2"].value == "\n".join(expected_groups[0]["evidence_refs"])
    assert {book["原文来源"].cell(row, 1).value for row in range(2, 4)} == refs
    assert {book["原文来源"].cell(row, 7).value for row in range(2, 4)} == {
        b.text for b in repo.blocks if b.evidence_id in refs}


def test_legacy_v10_history_payload_and_export_shape_remain_unchanged(material_case):
    client, engine, *_ = material_case
    run_id, payload, _ = persist_material_report(material_case, rule_version="dynamic-rules-v10", material=False)
    assert DynamicRepository(engine).verified(run_id)[1] == payload
    assert "material_sources" not in payload["facts"][0]
    exported = client.get(f"/api/v1/agent-runs/{run_id}/export")
    assert exported.status_code == 200, exported.text
    assert set(exported.json()) == {"schema_version", "result", "sources", "item_attribution"}
    assert "material_source_groups" not in exported.json()["result"]["checks"][0]
    excel = client.get(f"/api/v1/agent-runs/{run_id}/export?format=xlsx")
    assert excel.status_code == 200, excel.text
    assert "材料引用分组" not in load_workbook(BytesIO(excel.content)).sheetnames


@pytest.mark.parametrize("change", ["omit_unchanged", "omit_false", "quote", "refs", "delete", "empty"])
def test_correction_cannot_drop_or_replace_original_material_sources(material_case, monkeypatch, change):
    from backend.app.criteria import facts as fact_module

    client, engine, *_ = material_case
    run_id, payload, repo = persist_material_report(material_case)
    fact = payload["facts"][0]
    candidate = {k: fact[k] for k in FactCandidate.model_fields}
    if change in {"omit_false", "delete"}:
        candidate["value"] = False
    if change in {"quote", "refs", "delete", "empty"}:
        candidate["material_sources"] = deepcopy(fact["material_sources"])
    if change == "quote":
        candidate["material_sources"][0]["quote"] = "检测报告 检验结果为合格且已盖章"
    elif change == "refs":
        other = next(b for b in repo.blocks if "已盖章" in b.text)
        candidate["material_sources"] = [{"evidence_refs": [other.evidence_id], "quote": other.text}]
    elif change in {"delete", "empty"}:
        candidate["material_sources"] = []
    observed = []
    original = fact_module.validate_fact

    def validate(candidate, *args, **kwargs):
        encoded = candidate.model_dump(mode="json") if isinstance(candidate, FactCandidate) else candidate
        observed.append(deepcopy(encoded))
        return original(candidate, *args, **kwargs)

    monkeypatch.setattr(fact_module, "validate_fact", validate)
    body = {"fact_id": fact["fact_id"], "candidate": candidate, "reason": "按原文复核",
            "expected_revision": 2, "expected_payload_digest": digest(payload)}
    response = client.post(f"/api/v1/agent-runs/{run_id}/fact-corrections",
        headers={"Idempotency-Key": "material-correct-" + change}, json=body)
    assert response.status_code in {409, 422}, response.text
    if change.startswith("omit"):
        assert observed[-1]["material_sources"] == fact["material_sources"]
    assert DynamicRepository(engine).verified(run_id)[1] == payload
    with engine.connect() as connection:
        assert connection.execute(select(dynamic_corrections).where(dynamic_corrections.c.run_id == run_id)).first() is None


@pytest.mark.parametrize("material_case", ["large"], indirect=True)
@pytest.mark.parametrize("tamper_cache", [False, True])
def test_nested_body_cannot_be_forged_into_report_or_cross_model_batch(material_case, tamper_cache):
    from backend.app.criteria.facts import fact_candidate_data

    _, engine, *_ = material_case
    run_id, payload, repo = persist_material_report(material_case, material=False)
    store = DynamicRepository(engine)
    row, _, _ = store.verified(run_id)
    definition = material_case[-1]
    context = review_bindings(row["item_bindings"], definition, repo)
    plans, _ = review_plan(repo, definition, rule_version=payload["rule_version"], item_bindings=context)
    main = payload["facts"][0]["evidence_refs"][0]
    batch, _, entry = next((batch, data, entry) for batch, data, entry in plans["vendor-a"]
                           if main in {b["evidence_id"] for b in batch})
    allowed = {b["evidence_id"] for b in batch}
    outside = next(b for b in repo.blocks if b.evidence_id not in allowed and "检验结果" in b.text)
    candidate = fact_candidate_data(payload["facts"][0]) | {
        "material_sources": [{"evidence_refs": [outside.evidence_id], "quote": outside.text}]}
    changed = deepcopy(payload)
    # Both spans are real and item-authorized. Recomputing every report/cache
    # digest still cannot make unseen body evidence part of this model batch.
    material_fact = validate_fact(candidate, definition, repo, "vendor-a",
        rule_version=payload["rule_version"], item_bindings=context)
    changed["facts"] = [material_fact] if tamper_cache else [*payload["facts"], material_fact]
    changed["checks"] = evaluate_checklist(definition, changed["facts"], changed["coverage"],
        snapshot_id=repo.snapshot_id, rule_version=payload["rule_version"],
        item_coverage=changed["item_coverage"], item_bindings=context)
    with engine.begin() as connection:
        connection.execute(update(dynamic_runs).where(dynamic_runs.c.run_id == run_id).values(
            payload_json=changed, payload_digest=digest(changed)))
        if tamper_cache:
            output = {"facts": [candidate]}
            connection.execute(update(dynamic_batches).where(dynamic_batches.c.run_id == run_id,
                dynamic_batches.c.batch_id == entry["batch_id"]).values(output_json=output, output_digest=digest(output)))
    with pytest.raises(RepositoryError, match="batch_source_changed"):
        store._validate_manifest(row, changed, definition, repo, context=context)
    with pytest.raises(RepositoryError, match="dynamic_source_changed"):
        store.verified(run_id)


def test_manual_previous_keeps_independent_body_sources(material_case):
    _, engine, case_id, snapshot, conversation, _ = material_case
    run_id, payload, _ = persist_material_report(material_case)
    reviewed = deepcopy(payload) | {"needs_review": True}
    with engine.begin() as connection:
        connection.execute(update(dynamic_runs).where(dynamic_runs.c.run_id == run_id).values(
            payload_json=reviewed, payload_digest=digest(reviewed)))
    runs = RunRepository(engine)
    receipt = runs.submit(SubmitRun(case_id, conversation, snapshot, 2, "重新审查", "standard",
                                   "material-inherit", workflow="auto"))
    claim = runs.claim(receipt.run_id, receipt.attempt_id)
    inherited, stamp = DynamicRepository(engine).manual_previous(claim)
    assert inherited == reviewed
    assert inherited["facts"][0]["material_sources"] == payload["facts"][0]["material_sources"]
    assert len(stamp) == 64
