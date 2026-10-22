"""Repacked requests keep persisted coverage, original evidence and exports valid."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from backend.app.agent.dynamic import (
    _extraction_messages,
    _facts_schema,
    _repair_envelope,
    review_plan,
)
from backend.app.agent.policy import input_bound
from backend.app.criteria.facts import BINDING_RULE_VERSION
from backend.app.criteria.models import ChecklistDefinition
from backend.app.db.agent_snapshot import load_agent_snapshot
from backend.app.db.dynamic_repository import DynamicRepository
from backend.app.db.item_binding_repository import load_run_bindings
from backend.app.db.models import dynamic_batches, dynamic_runs
from backend.app.providers.base import ProviderError
from backend.tests import test_document_ingestion, test_dynamic_workflow
from backend.tests.test_dynamic_agent import ScriptProvider

ingestion_case = test_document_ingestion.ingestion_case


def matrix_workbook(text):
    book = Workbook()
    book.active.title = "采购资料"
    book.active["A1"] = text
    if "不低于" not in text:
        for row in range(2, 7):
            book.active.cell(row, 1, "苹果 其他说明 " + "备注" * 100)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.fixture
def dynamic_case(ingestion_case, request, monkeypatch):
    monkeypatch.setattr(test_dynamic_workflow, "workbook", matrix_workbook)
    return test_dynamic_workflow.dynamic_case.__wrapped__(ingestion_case, request)


def small_window(case, run_id, confirmed):
    rules = ChecklistDefinition.model_validate(confirmed["confirmed_version"]["definition"])
    repo = load_agent_snapshot(case[1], case[2], case[3])
    with case[1].connect() as connection:
        row = connection.execute(select(dynamic_runs).where(dynamic_runs.c.run_id == run_id)).mappings().one()
        frozen = load_run_bindings(connection, row, rules)
    plans, _ = review_plan(repo, rules, rule_version=BINDING_RULE_VERSION, item_bindings=frozen)
    sizes = [input_bound(_repair_envelope(_extraction_messages(data, bound=True, rule_version=BINDING_RULE_VERSION), True),
                         [_facts_schema(supplier, material_sources=True)])
             for supplier in plans for _, data, _ in plans[supplier]]
    # Make the real complete-request estimate miss by ~100 bytes, so splitting
    # must happen. Derive the synthetic model window with 5% safety +1536 output.
    return (min(sizes) - 100 + 1536) * 100 // 95


def test_repacked_review_roundtrips_sql_history_progress_and_exports(dynamic_case):
    confirmed, _ = test_dynamic_workflow.confirm(dynamic_case)
    run_id = test_dynamic_workflow.message(dynamic_case, key="repacked-review")
    provider = ScriptProvider()
    provider.context_window = small_window(dynamic_case, run_id, confirmed)
    report = test_dynamic_workflow.run(dynamic_case, run_id, provider)
    assert report.status == "completed", report
    row, payload, repo = DynamicRepository(dynamic_case[1]).verified(run_id)
    rules = ChecklistDefinition.model_validate(payload["definition"])
    # Rebuild the original plan independently of the Provider's small window.
    expected, _ = review_plan(repo, rules, rule_version=payload["rule_version"], item_bindings=row["item_bindings"])
    assert payload["batch_manifest"] == [m for supplier in sorted(expected) for _, _, m in expected[supplier]]
    assert all(p["request_parts"] > 1 and p["completed_parts"] == p["request_parts"]
               for p in payload["batch_coverage"])
    with dynamic_case[1].connect() as connection:
        caches = list(connection.execute(select(dynamic_batches).where(dynamic_batches.c.run_id == run_id)).mappings())
    assert len(caches) == len(payload["batch_manifest"]) == 2
    assert {c["batch_id"] for c in caches} == {m["batch_id"] for m in payload["batch_manifest"]}
    url = f"/api/v1/agent-runs/{run_id}"
    checks = dynamic_case[0].get(url + "/checks")
    assert checks.status_code == 200, checks.text
    assert all(c["is_formal"] for c in checks.json()["checks"])
    assert dynamic_case[0].get(url + "/progress").json()["validated_batches"] == 2
    assert dynamic_case[0].get(url + "/export").status_code == 200
    excel = dynamic_case[0].get(url + "/export?format=xlsx")
    assert excel.status_code == 200
    assert load_workbook(BytesIO(excel.content))["核验结果"]["B2"].value == "数量"


def test_incomplete_parent_is_not_cached_or_formal_and_retry_revalidates(dynamic_case):
    confirmed, _ = test_dynamic_workflow.confirm(dynamic_case)
    run_id = test_dynamic_workflow.message(dynamic_case, key="partial-repacked-review")

    class UnknownSecondChild(ScriptProvider):
        def extract(self, *args, **kwargs):
            if self.extractions:
                raise ProviderError("timeout", outcome_unknown=True)
            return super().extract(*args, **kwargs)

    provider = UnknownSecondChild()
    provider.context_window = small_window(dynamic_case, run_id, confirmed)
    report = test_dynamic_workflow.run(dynamic_case, run_id, provider)
    assert report.status == "partial" and report.error_code == "timeout", report
    _, partial, _ = DynamicRepository(dynamic_case[1]).verified(run_id)
    assert partial["facts"] and not any(partial["coverage"].values())
    with dynamic_case[1].connect() as connection:
        assert connection.scalar(select(dynamic_batches.c.batch_id).where(dynamic_batches.c.run_id == run_id)) is None
    url = f"/api/v1/agent-runs/{run_id}"
    checks = dynamic_case[0].get(url + "/checks")
    assert checks.status_code == 200 and all(not c["is_formal"] for c in checks.json()["checks"])
    retry = dynamic_case[0].post(url + "/retry", headers={"Idempotency-Key": "explicit-repack-retry"},
                                json={"expected_revision": 3})
    assert retry.status_code == 202, retry.text
    provider = ScriptProvider()
    assert test_dynamic_workflow.run(dynamic_case, run_id, provider).status == "completed"
    assert provider.extractions == 2
    assert DynamicRepository(dynamic_case[1]).verified(run_id)[1]["facts"]
