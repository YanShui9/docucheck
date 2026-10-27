"""Independent tax provenance survives the real offline SQL/API/worker path."""

from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import update

from backend.app.db.agent_snapshot import load_agent_snapshot
from backend.app.db.dynamic_repository import DynamicRepository
from backend.app.db.ingestion import DocumentIngestion
from backend.app.db.models import dynamic_runs
from backend.app.db.run_executor import execute_run
from backend.app.db.run_repository import RepositoryError
from backend.app.evidence.repository import digest
from backend.tests import test_document_ingestion
from backend.tests.item_binding_helpers import confirm_single_item_bindings
from backend.tests.test_material_agent import MaterialProvider

ingestion_case = test_document_ingestion.ingestion_case


@pytest.fixture
def money_case(ingestion_case, request):
    client, engine, storage, case_id = ingestion_case
    book = Workbook()
    sheet = book.active
    sheet.title = "报价"
    equivalent = getattr(request, "param", None) == "equivalent"
    rows = ([("苹果",), ("不含税单价", 25, "元/箱"), ("不含税单价", "25.00", "元/箱")]
            if equivalent else [("苹果",), ("不含税单价", 25, "元/箱"), ("税费口径", "含税"),
                                ("报价日期", "2026-10-19 00:00:00")])
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    uploaded = client.post(f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "money-upload"},
        data={"role": "quote", "supplier_handle": "vendor-a", "supplier_name": "供应商甲",
              "expected_revision": "0"},
        files={"file": ("苹果报价.xlsx", stream.getvalue(), "application/octet-stream")})
    assert uploaded.status_code == 201, uploaded.text
    snapshot = uploaded.json()["snapshot_id"]
    document = uploaded.json()["document_id"]
    assert DocumentIngestion(engine, storage).process_document(case_id, snapshot, document)["parse_status"] == "complete"
    definition = {"items": [{"item_id": "apple", "name": "苹果"}], "criteria": [
        {"criterion_id": "price", "item_id": "apple", "name": "单价上限",
         "rule": {"type": "money", "operator": "lte", "value": "25", "currency": "CNY",
                  "per_unit": "箱", "tax_basis": "excluded"},
         "source": {"kind": "user_defined", "reason": "买方确认单价上限"}}]}
    url = f"/api/v1/cases/{case_id}/snapshots/{snapshot}/checklist"
    saved = client.put(url, headers={"Idempotency-Key": "money-draft"}, json={
        **definition, "expected_revision": 1, "expected_checklist_revision": 0, "reason": "采购要求"})
    assert saved.status_code == 200, saved.text
    confirmed = client.post(url + "/confirm", headers={"Idempotency-Key": "money-confirm"}, json={
        "criterion_ids": ["price"], "expected_revision": 1, "expected_checklist_revision": 1,
        "confirmed": True, "reason": "确认"})
    assert confirmed.status_code == 200, confirmed.text
    confirm_single_item_bindings(engine, case_id, snapshot, confirmed.json())
    conversation = client.post(f"/api/v1/cases/{case_id}/conversations").json()["id"]
    return client, engine, case_id, snapshot, conversation


def execute_money(case):
    client, engine, case_id, snapshot, conversation = case
    posted = client.post(f"/api/v1/conversations/{conversation}/messages",
        headers={"Idempotency-Key": "money-review"}, json={
            "content": "正式审查供应商报价",
            "expected_revision": 1, "mode": "standard", "operation": "review"})
    assert posted.status_code == 202, posted.text
    run_id = posted.json()["run_id"]
    repo = load_agent_snapshot(engine, case_id, snapshot)
    price_blocks = [b for b in repo.blocks if b.cell in {"A2", "B2", "C2"}]
    candidate = {"criterion_id": "price", "item_id": "apple", "supplier": "vendor-a",
                 "value": "25", "unit": "元/箱", "evidence_refs": [b.evidence_id for b in price_blocks],
                 "quote": " ".join(b.text for b in price_blocks)}
    attempt = client.get(f"/api/v1/agent-runs/{run_id}").json()["active_attempt_id"]
    result = execute_run(engine, run_id, attempt, MaterialProvider([candidate]), heartbeat_interval_seconds=60)
    assert result.status == "completed", result
    payload = DynamicRepository(engine).verified(run_id)[1]
    return run_id, payload, repo


def test_worker_persists_tax_conflict_and_api_exports_both_independent_sources(money_case):
    client, _, *_ = money_case
    run_id, payload, repo = execute_money(money_case)
    assert payload["rule_version"] == "dynamic-rules-v16"
    assert payload["checks"][0]["business_status"] == "资料冲突"
    tax_refs = {b.evidence_id for b in repo.blocks if b.cell in {"A3", "B3"}}
    fact = payload["facts"][0]
    assert tax_refs == set(fact["tax_sources"][0]["evidence_refs"])
    response = client.get(f"/api/v1/agent-runs/{run_id}/checks")
    assert response.status_code == 200, response.text
    assert response.json()["checks"][0]["tax_source_groups"][0]["evidence_refs"] == fact["tax_sources"][0]["evidence_refs"]
    exported = client.get(f"/api/v1/agent-runs/{run_id}/export")
    assert exported.status_code == 200, exported.text
    assert {s["evidence_id"] for s in exported.json()["sources"]} == set(fact["evidence_refs"]) | tax_refs
    xlsx = client.get(f"/api/v1/agent-runs/{run_id}/export?format=xlsx")
    assert xlsx.status_code == 200, xlsx.text
    book = load_workbook(BytesIO(xlsx.content))
    assert "税口径引用分组" in book.sheetnames
    texts = {row[6] for row in book["原文来源"].iter_rows(min_row=2, values_only=True)}
    assert {"不含税单价", "25", "元/箱", "税费口径", "含税"} <= texts


@pytest.mark.parametrize("money_case", ["equivalent"], indirect=True)
def test_worker_history_and_exports_preserve_all_equivalent_price_sources(money_case):
    client, engine, *_ = money_case
    run_id, payload, repo = execute_money(money_case)
    check = payload["checks"][0]
    assert check["business_status"] == "满足" and check["is_formal"]
    assert check["processing_status"] == "complete"
    fact = payload["facts"][0]
    assert fact["price_scope_state"] == fact["tax_scope_state"] == "resolved"
    assert fact["price_sources"] and fact["tax_sources"]
    assert DynamicRepository(engine).verified(run_id)[1] == payload
    exported = client.get(f"/api/v1/agent-runs/{run_id}/export").json()
    expected_refs = {b.evidence_id for b in repo.blocks if b.cell and b.cell[1:] in {"2", "3"}}
    assert {s["evidence_id"] for s in exported["sources"]} == expected_refs
    exported_check = exported["result"]["checks"][0]
    assert {key: exported_check[key] for key in check} == check
    assert exported_check["checklist_version_id"]
    xlsx = client.get(f"/api/v1/agent-runs/{run_id}/export?format=xlsx")
    book = load_workbook(BytesIO(xlsx.content))
    assert {"单价引用分组", "税口径引用分组"} <= set(book.sheetnames)
    assert {row[0] for row in book["原文来源"].iter_rows(min_row=2, values_only=True)} == expected_refs


@pytest.mark.parametrize("mutation", ["delete", "quote", "basis", "scope"])
def test_verified_history_rejects_changed_tax_metadata_even_after_rehash(money_case, mutation):
    _, engine, *_ = money_case
    run_id, payload, _ = execute_money(money_case)
    changed = deepcopy(payload)
    fact = changed["facts"][0]
    if mutation == "delete":
        del fact["tax_sources"]
    elif mutation == "quote":
        fact["tax_sources"][0]["quote"] = "税费口径 不含税"
    elif mutation == "basis":
        fact["tax_sources"][0]["tax_basis"] = "excluded"
    else:
        fact["tax_scope_state"] = "unclear"
    with engine.begin() as connection:
        connection.execute(update(dynamic_runs).where(dynamic_runs.c.run_id == run_id).values(
            payload_json=changed, payload_digest=digest(changed)))
    with pytest.raises(RepositoryError, match="dynamic_source_changed"):
        DynamicRepository(engine).verified(run_id)
