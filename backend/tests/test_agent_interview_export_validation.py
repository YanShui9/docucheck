"""The downloaded workbook must faithfully expose the complete JSON artifact."""

import importlib
from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.app.api.run_export import render_xlsx


def validator():
    return importlib.import_module('evals.agent_interview.export_validation').validate_review_artifacts


def fixture(*, dynamic=True, detailed=True):
    run = {'id': 'run-1', 'case_id': 'case-1', 'snapshot_id': 'snapshot-1',
           'revision': 4, 'status': 'partial'}
    check = {'supplier': '供应商甲', 'criterion_name': '含税单价', 'field_key': 'unit_price',
             'processing_status': 'complete', 'business_status': '资料冲突',
             'reason_code': 'money_tax_conflict', 'rule_version': 'dynamic-rules-v16',
             'is_formal': False, 'evidence_refs': ['source-1', 'source-2'],
             'reason_evidence_refs': ['source-2'], 'fact_ids': ['fact-1', 'fact-2'],
             'check_result_id': 'check-1', 'item_id': 'item-1', 'criterion_id': 'criterion-1',
             'item_name': '手机', 'material_source_groups': [{'fact_id': 'fact-1',
                 'evidence_refs': ['source-1', 'source-2']}],
             'tax_source_groups': [{'fact_id': 'fact-2', 'evidence_refs': ['source-2']}],
             'price_source_groups': [{'fact_id': 'fact-1', 'evidence_refs': ['source-1']}]}
    checks = {'run_id': run['id'], 'case_id': run['case_id'], 'snapshot_id': run['snapshot_id'],
              'revision': run['revision'], 'run_status': run['status'], 'is_current': True,
              'checks': [check, dict(deepcopy(check), check_result_id='check-2',
                                   is_formal=True, business_status=None, criterion_name='')]}
    if dynamic:
        checks['checklist_version_id'] = 'checklist-1'
    sources = [
        {'evidence_id': 'source-1', 'document_id': 'document-1', 'document_name': '=报价.xlsx',
         'page': None, 'sheet': '报价表', 'cell': 'B2',
         'source_text': '=1+1\r\n单价 10\r含税\n\x00\x0b\x1f\t尾部'},
        {'evidence_id': 'source-2', 'document_id': 'document-2', 'document_name': '税说明.pdf',
         'page': 2, 'sheet': None, 'cell': '', 'source_text': '不含税 10 元'},
    ]
    payload = {'schema_version': 1, 'result': deepcopy(checks), 'sources': sources}
    if detailed:
        payload['item_attribution'] = {
            'version': {'id': 'binding-1', 'digest': 'digest-1', 'confirmed_by': 'user-1',
                        'confirmed_at': '2026-10-28T00:00:00Z', 'reason': '确认\r\n对应',
                        'payload': {'entries': [
                            {'document_id': 'document-1', 'item_id': 'item-1',
                             'identifier_class': 'sku', 'supplier_sku': 'SKU-001',
                             'selector': {'sheet': '报价表', 'cell_range': 'A2:C2'},
                             'anchor': {'evidence_refs': ['source-1'], 'quote': '手机\rSKU-001'},
                             'sku_anchor': {'evidence_refs': ['source-1'], 'quote': 'SKU-001'}},
                            {'document_id': 'document-2', 'item_id': 'item-1',
                             'identifier_class': 'name', 'supplier_sku': None,
                             'selector': {'page_start': 1, 'page_end': 2},
                             'anchor': {'evidence_refs': ['source-2'], 'quote': '手机'},
                             'sku_anchor': None},
                        ]}},
            'notice': '品目对应锚点仅说明归属；数值结论以核验结果引用的原文为准。',
            'anchors': deepcopy(sources),
            'exclusions': [{'document_id': 'document-2', 'document_name': '税说明.pdf',
                'supplier': '供应商甲', 'page': 3, 'other_sku': 'SKU-OTHER', 'reason': '其他型号',
                'sku_anchor': {'quote': 'SKU-OTHER'}, 'confirmed_by': 'user-1',
                'confirmed_at': '2026-10-28T00:00:00Z', 'binding_version_id': 'binding-1',
                'sources': [{'evidence_id': 'excluded-1', 'source_text': '其他手机\r\nSKU-OTHER'},
                            {'evidence_id': 'excluded-2', 'source_text': '其他型号参数'}]}],
        }
    else:
        for row in checks['checks']:
            for key in ('material_source_groups', 'tax_source_groups', 'price_source_groups'):
                row.pop(key)
        payload['result'] = deepcopy(checks)
    return run, checks, payload


def edit(binary, mutation):
    book = load_workbook(BytesIO(binary))
    try:
        mutation(book)
        stream = BytesIO()
        book.save(stream)
        return stream.getvalue()
    finally:
        book.close()


def test_validator_is_an_independent_module():
    assert importlib.util.find_spec('evals.agent_interview.export_validation') is not None


@pytest.mark.parametrize(('dynamic', 'detailed'), [(False, False), (True, False), (True, True)])
def test_real_renderer_contract_is_accepted(dynamic, detailed):
    run, checks, payload = fixture(dynamic=dynamic, detailed=detailed)
    validator()(run, checks, payload, render_xlsx(payload))


# Every exported cell, including header/metadata blanks, is covered independently.
SHAPES = {'运行信息': (10, 2), '核验结果': (3, 15), '原文来源': (3, 7),
          '品目对应依据': (9, 8), '对应锚点原文': (3, 7), '范围排除': (3, 13),
          '材料引用分组': (3, 5), '税口径引用分组': (3, 5), '单价引用分组': (3, 5)}
CELLS = [(sheet, row, column) for sheet, (rows, columns) in SHAPES.items()
         for row in range(1, rows + 1) for column in range(1, columns + 1)]


@pytest.mark.parametrize(('sheet', 'row', 'column'), CELLS)
def test_every_exported_cell_mutation_is_rejected(sheet, row, column):
    run, checks, payload = fixture()
    binary = edit(render_xlsx(payload), lambda book: setattr(
        book[sheet].cell(row, column), 'value', '篡改内容'))
    with pytest.raises(ValueError, match='evaluation_xlsx_invalid'):
        validator()(run, checks, payload, binary)


@pytest.mark.parametrize('sheet', SHAPES)
@pytest.mark.parametrize('mutation', ['extra_row', 'missing_row', 'extra_column', 'missing_column', 'missing_sheet'])
def test_sheet_shapes_are_exact(sheet, mutation):
    run, checks, payload = fixture()
    def change(book):
        target = book[sheet]
        if mutation == 'extra_row':
            target.cell(target.max_row + 1, 1, 'extra')
        elif mutation == 'missing_row':
            target.delete_rows(target.max_row)
        elif mutation == 'extra_column':
            target.cell(1, target.max_column + 1, 'extra')
        elif mutation == 'missing_column':
            target.delete_cols(target.max_column)
        else:
            book.remove(target)
    with pytest.raises(ValueError, match='evaluation_xlsx_invalid'):
        validator()(run, checks, payload, edit(render_xlsx(payload), change))


def test_extra_sheet_is_rejected():
    run, checks, payload = fixture()
    with pytest.raises(ValueError, match='evaluation_xlsx_invalid'):
        validator()(run, checks, payload, edit(render_xlsx(payload),
                    lambda book: book.create_sheet('额外表')))


@pytest.mark.parametrize('sheet', SHAPES)
def test_formula_is_rejected_in_every_sheet(sheet):
    run, checks, payload = fixture()
    binary = edit(render_xlsx(payload), lambda book: setattr(book[sheet]['A2'], 'value', '=1+1'))
    with pytest.raises(ValueError, match='evaluation_xlsx_invalid'):
        validator()(run, checks, payload, binary)


@pytest.mark.parametrize('field', ['run_id', 'case_id', 'snapshot_id', 'revision', 'run_status', 'is_current'])
def test_checks_identity_is_bound_to_run(field):
    run, checks, payload = fixture()
    checks[field] = False if field == 'is_current' else 'wrong'
    payload['result'] = deepcopy(checks)
    with pytest.raises(ValueError, match='evaluation_result_identity'):
        validator()(run, checks, payload, render_xlsx(payload))


@pytest.mark.parametrize('field', ['item_name', 'criterion_id', 'evidence_refs', 'reason_evidence_refs',
                                 'fact_ids', 'material_source_groups', 'tax_source_groups', 'price_source_groups'])
def test_json_result_is_entire_checks_object(field):
    run, checks, payload = fixture()
    payload['result']['checks'][0][field] = 'changed'
    with pytest.raises(ValueError, match='evaluation_export_mismatch'):
        validator()(run, checks, payload, render_xlsx(fixture()[2]))


def test_empty_optional_sheets_and_absent_checklist_have_exact_contract():
    run, checks, payload = fixture(dynamic=False, detailed=False)
    checks['checks'] = []
    payload['result'] = deepcopy(checks)
    payload['sources'] = []
    validator()(run, checks, payload, render_xlsx(payload))


def test_optional_attribution_without_entries_or_exclusions_is_accepted():
    run, checks, payload = fixture()
    attribution = payload['item_attribution']
    attribution['version']['id'] = None
    attribution['version']['confirmed_by'] = None
    attribution['version']['confirmed_at'] = None
    attribution['version']['payload']['entries'] = []
    attribution['anchors'] = []
    attribution['exclusions'] = []
    validator()(run, checks, payload, render_xlsx(payload))


@pytest.mark.parametrize('binary', [b'not-a-workbook', b'', b'PK\x03\x04broken'])
def test_invalid_workbook_is_rejected(binary):
    run, checks, payload = fixture()
    with pytest.raises(ValueError, match='evaluation_xlsx_invalid'):
        validator()(run, checks, payload, binary)


def test_validation_does_not_call_application_renderer(monkeypatch):
    run, checks, payload = fixture()
    binary = render_xlsx(payload)
    def unavailable(*args, **kwargs):
        pytest.fail('validator must independently compute the contract')
    monkeypatch.setattr('backend.app.api.run_export.render_xlsx', unavailable)
    validator()(run, checks, payload, binary)
