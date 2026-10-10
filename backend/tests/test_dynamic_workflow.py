"""Uploaded buyer requirements -> candidate -> confirmation -> versioned review."""

from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, update

from backend.app.db.dynamic_repository import DynamicRepository
from backend.app.db.ingestion import DocumentIngestion
from backend.app.db.models import (
    document_versions,
    dynamic_corrections,
    dynamic_runs,
    task_memory,
)
from backend.app.db.run_executor import execute_run
from backend.app.db.run_repository import RepositoryError
from backend.tests import test_document_ingestion
from backend.tests.test_dynamic_agent import ScriptProvider

ingestion_case = test_document_ingestion.ingestion_case


def workbook(text):
    book = Workbook()
    book.active.title = "采购资料"
    book.active["A1"] = text
    output = BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.mark.parametrize('old_version', ['dynamic-rules-v1', 'dynamic-rules-v2'])
def test_v1_persisted_report_still_reads_and_new_plan_invalidates_its_cache(dynamic_case, monkeypatch, old_version):
    import backend.app.agent.dynamic as workflow
    from backend.app.criteria.facts import evaluate_checklist, validate_fact

    confirm(dynamic_case)
    old_id = message(dynamic_case, key='legacy-review')
    plan = workflow.review_plan
    with monkeypatch.context() as patch:
        patch.setattr(workflow, 'RULE_VERSION', old_version)
        patch.setattr(workflow, 'review_plan', lambda repo, definition: plan(repo, definition, rule_version=old_version))
        patch.setattr(workflow, 'validate_fact', lambda *args: validate_fact(*args, rule_version=old_version))
        patch.setattr(workflow, 'evaluate_checklist', lambda *args, **kwargs: evaluate_checklist(*args, **kwargs, rule_version=old_version))
        assert run(dynamic_case, old_id, ScriptProvider()).status == 'completed'
    row, payload, repo = DynamicRepository(dynamic_case[1]).verified(old_id)
    assert row['status'] == 'completed' and payload['rule_version'] == old_version
    assert dynamic_case[0].get(f'/api/v1/agent-runs/{old_id}/checks').status_code == 200
    assert dynamic_case[0].get(f'/api/v1/agent-runs/{old_id}/export').status_code == 200
    new_plan, _ = plan(repo, payload['definition'])
    new_digests = {entry['input_digest'] for batches in new_plan.values() for _, _, entry in batches}
    assert not new_digests.intersection(entry['input_digest'] for entry in payload['batch_manifest'])
    new_ids = {entry['batch_id'] for batches in new_plan.values() for _, _, entry in batches}
    assert not new_ids.intersection(entry['batch_id'] for entry in payload['batch_manifest'])


@pytest.mark.parametrize('old_version', ['dynamic-rules-v1', 'dynamic-rules-v2'])
def test_retry_after_rules_upgrade_uses_new_sql_batch_namespace(dynamic_case, monkeypatch, old_version):
    import backend.app.agent.dynamic as workflow
    from backend.app.criteria.facts import evaluate_checklist, validate_fact
    from backend.app.providers.base import ProviderError

    class SecondBatchFails(ScriptProvider):
        def extract(self, *args, **kwargs):
            if self.extractions == 1:
                self.failure = ProviderError('unavailable')
            return super().extract(*args, **kwargs)

    confirm(dynamic_case)
    run_id = message(dynamic_case, key='old-version-partial')
    plan = workflow.review_plan
    with monkeypatch.context() as patch:
        patch.setattr(workflow, 'RULE_VERSION', old_version)
        patch.setattr(workflow, 'review_plan', lambda repo, definition: plan(repo, definition, rule_version=old_version))
        patch.setattr(workflow, 'validate_fact', lambda *args: validate_fact(*args, rule_version=old_version))
        patch.setattr(workflow, 'evaluate_checklist', lambda *args, **kwargs: evaluate_checklist(*args, **kwargs, rule_version=old_version))
        assert run(dynamic_case, run_id, SecondBatchFails()).status == 'partial'
    progress_url = f'/api/v1/agent-runs/{run_id}/progress'
    old_progress = dynamic_case[0].get(progress_url).json()
    assert (old_progress['validated_batches'], old_progress['planned_batches']) == (1, 2)
    response = dynamic_case[0].post(f'/api/v1/agent-runs/{run_id}/retry',
        headers={'Idempotency-Key': 'upgrade-retry'}, json={'expected_revision': 3})
    assert response.status_code == 202
    provider = ScriptProvider()
    result = run(dynamic_case, run_id, provider)
    assert result.status == 'completed', result
    assert provider.extractions == 2
    assert DynamicRepository(dynamic_case[1]).verified(run_id)[1]['rule_version'] == workflow.RULE_VERSION
    current_progress = dynamic_case[0].get(progress_url).json()
    assert (current_progress['validated_batches'], current_progress['planned_batches']) == (2, 2)


@pytest.fixture
def dynamic_case(ingestion_case, request):
    client, engine, storage, case_id = ingestion_case
    for revision, (role, supplier, text) in enumerate(
        [
            ("requirement", None, "苹果 数量不低于10箱"),
            ("quote", "vendor-a", getattr(request, "param", "苹果 数量12箱")),
            ("quote", "vendor-b", "苹果 数量8箱"),
        ]
    ):
        data = {"role": role, "expected_revision": str(revision)}
        if supplier:
            data.update(supplier_handle=supplier, supplier_name=supplier)
        response = client.post(
            f"/api/v1/cases/{case_id}/documents",
            data=data,
            headers={"Idempotency-Key": f"upload-{revision}"},
            files={
                "file": (
                    f"{supplier or 'buyer'}.xlsx",
                    workbook(text),
                    "application/octet-stream",
                )
            },
        )
        assert response.status_code == 201, response.text
        snapshot = response.json()["snapshot_id"]
    with engine.connect() as connection:
        ids = list(
            connection.scalars(
                select(document_versions.c.id).where(
                    document_versions.c.snapshot_id == snapshot
                )
            )
        )
    ingestion = DocumentIngestion(engine, storage)
    for document_id in ids:
        assert (
            ingestion.process_document(case_id, snapshot, document_id)["parse_status"]
            == "complete"
        )
    conversation = client.post(f"/api/v1/cases/{case_id}/conversations").json()["id"]
    return client, engine, case_id, snapshot, conversation


def message(case, *, proposal=False, key="message-1", content="完整审查"):
    client, _, _, _, conversation = case
    response = client.post(
        f"/api/v1/conversations/{conversation}/messages",
        headers={"Idempotency-Key": key},
        json={
            "content": content,
            "expected_revision": 3,
            "operation": "propose_checklist" if proposal else "review",
        },
    )
    assert response.status_code == 202, response.text
    return response.json()["run_id"]


def run(case, run_id, provider):
    attempt_id = case[0].get(f"/api/v1/agent-runs/{run_id}").json()["active_attempt_id"]
    return execute_run(
        case[1], run_id, attempt_id, provider, heartbeat_interval_seconds=60
    )


def confirm(case):
    client, _, case_id, snapshot, _ = case
    proposal_id = message(case, proposal=True)
    result = run(case, proposal_id, ScriptProvider(decision="propose_checklist"))
    assert result.status == "completed", result
    preview = client.get(f"/api/v1/agent-runs/{proposal_id}/proposal")
    assert preview.status_code == 200, preview.text
    preview = preview.json()
    assert preview["requires_user_confirmation"] is True
    assert (
        client.get(f"/api/v1/cases/{case_id}/snapshots/{snapshot}/checklist").json()[
            "confirmed_version"
        ]
        is None
    )
    url = f"/api/v1/cases/{case_id}/snapshots/{snapshot}/checklist"
    saved = client.put(
        url,
        headers={"Idempotency-Key": "draft-1"},
        json={
            **preview["definition"],
            "expected_revision": 3,
            "expected_checklist_revision": 0,
            "reason": "核对需求原文",
        },
    )
    assert saved.status_code == 200, saved.text
    accepted = client.post(
        url + "/confirm",
        headers={"Idempotency-Key": "confirm-1"},
        json={
            "expected_revision": 3,
            "expected_checklist_revision": saved.json()["checklist_revision"],
            "criterion_ids": [
                c["criterion_id"] for c in preview["definition"]["criteria"]
            ],
            "confirmed": True,
            "reason": "买方确认",
        },
    )
    assert accepted.status_code == 200, accepted.text
    return accepted.json(), preview["definition"]


def test_proposal_confirmation_review_export_and_task_memory(dynamic_case):
    confirmed, _ = confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="review-1")
    report = run(dynamic_case, run_id, ScriptProvider())
    assert report.status == "completed"
    response = client.get(f"/api/v1/agent-runs/{run_id}/checks")
    assert response.status_code == 200, response.text
    checks = response.json()
    assert checks["checklist_version_id"] == confirmed["confirmed_version"]["id"]
    assert {c["business_status"] for c in checks["checks"]} == {"满足", "不满足"}
    assert all(
        c["is_formal"] and c["criterion_name"] == "数量" and c["item_name"] == "苹果" for c in checks["checks"]
    )
    facts = client.get(f"/api/v1/agent-runs/{run_id}/facts").json()
    assert len(facts["fact_sets"]) == 2 and facts["payload_digest"]
    assert checks["payload_digest"] == facts["payload_digest"]
    assert all(g["rule"]["type"] == "numeric" for g in facts["fact_sets"])
    audit = client.get(f"/api/v1/agent-runs/{run_id}/fact-corrections").json()
    assert audit["payload_digest"] == facts["payload_digest"]
    exported = client.get(f"/api/v1/agent-runs/{run_id}/export").json()
    assert exported["result"]["checklist_version_id"] == checks["checklist_version_id"]
    excel = client.get(f"/api/v1/agent-runs/{run_id}/export?format=xlsx")
    assert excel.status_code == 200
    book = load_workbook(BytesIO(excel.content))
    assert book["核验结果"]["B2"].value == "数量"
    assert "清单版本" in {book["运行信息"].cell(i, 1).value for i in range(1, 12)}
    with engine.connect() as connection:
        state = connection.scalar(
            select(task_memory.c.state_json).where(
                task_memory.c.source_run_id == run_id
            )
        )
    assert (
        state["checkpoint"]["version"] == 4
        and state["checkpoint"]["source_run_id"] == run_id
    )
    assert state["checkpoint"]["checklist_version_id"] == checks["checklist_version_id"]
    progress = client.get(f"/api/v1/agent-runs/{run_id}/progress")
    assert progress.status_code == 200 and progress.json()["validated_batches"] == 2


def test_result_tampering_and_checklist_change_cannot_promote_results(dynamic_case):
    _, definition = confirm(dynamic_case)
    client, engine, case_id, snapshot, _ = dynamic_case
    run_id = message(dynamic_case, key="review-1")
    run(dynamic_case, run_id, ScriptProvider())
    saved = client.put(
        f"/api/v1/cases/{case_id}/snapshots/{snapshot}/checklist",
        headers={"Idempotency-Key": "change-1"},
        json={
            **definition,
            "expected_revision": 3,
            "expected_checklist_revision": 2,
            "reason": "重新核对",
        },
    )
    assert saved.status_code == 200, saved.text
    checks = client.get(f"/api/v1/agent-runs/{run_id}/checks").json()
    assert not checks["is_current"] and not any(
        c["is_formal"] for c in checks["checks"]
    )
    with engine.begin() as connection:
        payload = deepcopy(
            connection.scalar(
                select(dynamic_runs.c.payload_json).where(
                    dynamic_runs.c.run_id == run_id
                )
            )
        )
        payload["facts"][0]["value"] = "999"
        connection.execute(
            update(dynamic_runs)
            .where(dynamic_runs.c.run_id == run_id)
            .values(payload_json=payload)
        )
    with pytest.raises(RepositoryError, match="dynamic_source_changed"):
        DynamicRepository(engine).checks(run_id)


def test_impossible_correction_rejected_without_writes(dynamic_case):
    confirm(dynamic_case)
    client, _, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="review-1")
    run(dynamic_case, run_id, ScriptProvider())
    facts = client.get(f"/api/v1/agent-runs/{run_id}/facts").json()
    old = facts["fact_sets"][0]["facts"][0]
    from backend.app.criteria.facts import FactCandidate

    candidate = {k: old[k] for k in FactCandidate.model_fields}
    candidate["value"] = "999"
    response = client.post(
        f"/api/v1/agent-runs/{run_id}/fact-corrections",
        headers={"Idempotency-Key": "correct-1"},
        json={
            "fact_id": old["fact_id"],
            "candidate": candidate,
            "reason": "复核",
            "expected_revision": 3,
            "expected_payload_digest": facts["payload_digest"],
        },
    )
    assert response.status_code == 422, response.text
    assert (
        client.get(f"/api/v1/agent-runs/{run_id}/facts").json()["payload_digest"]
        == facts["payload_digest"]
    )


@pytest.mark.parametrize("dynamic_case", ["苹果 数量12箱或8箱"], indirect=True)
def test_same_source_two_values_keep_conflict_and_correction_audit(dynamic_case):
    confirm(dynamic_case)
    client, _, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="review-1")

    class AlternativeProvider(ScriptProvider):
        def extract(self, *args, **kwargs):
            output = super().extract(*args, **kwargs)
            for fact in output.tool_calls[0]["args"].get("facts", []):
                fact["unit"] = "箱"
            return output

    run(dynamic_case, run_id, AlternativeProvider())
    facts = client.get(f"/api/v1/agent-runs/{run_id}/facts").json()
    group = next(g for g in facts["fact_sets"] if g["supplier"] == "vendor-a")
    old = group["facts"][0]
    from backend.app.criteria.facts import FactCandidate

    candidate = {k: old[k] for k in FactCandidate.model_fields}
    candidate["value"] = "8"
    body = {
        "fact_id": old["fact_id"],
        "candidate": candidate,
        "reason": "原文也支持另一值",
        "expected_revision": 3,
        "expected_payload_digest": facts["payload_digest"],
    }
    url = f"/api/v1/agent-runs/{run_id}/fact-corrections"
    headers = {"Idempotency-Key": "correct-1"}
    response = client.post(url, headers=headers, json=body)
    assert response.status_code == 201, response.text
    assert response.json()["review_status"] == "needs_review"
    detail = client.get(f"/api/v1/agent-runs/{run_id}").json()
    assert "指标审查完成" in detail["answer_text"]
    assert "结论仍需人工复核" in detail["answer_text"]
    assert "未覆盖项" not in detail["answer_text"]
    assert client.post(url, headers=headers, json=body).json() == response.json()
    changed = client.get(f"/api/v1/agent-runs/{run_id}/facts").json()
    assert {
        f["value"]
        for g in changed["fact_sets"]
        if g["supplier"] == "vendor-a"
        for f in g["facts"]
    } == {"12", "8"}
    checks = client.get(f"/api/v1/agent-runs/{run_id}/checks").json()
    assert (
        next(c for c in checks["checks"] if c["supplier"] == "vendor-a")[
            "business_status"
        ]
        == "资料冲突"
    )
    assert not any(c["is_formal"] for c in checks["checks"])
    assert client.get(url).json()["corrections"][0]["old_fact"]["value"] == "12"
    assert checks["payload_digest"] == changed["payload_digest"]
    assert client.get(url).json()["payload_digest"] == response.json()["payload_digest"]
    exported = client.get(f"/api/v1/agent-runs/{run_id}/export").json()
    assert not any(c["is_formal"] for c in exported["result"]["checks"])
    assert (
        client.post(url, headers={"Idempotency-Key": "stale"}, json=body).status_code
        == 409
    )
    review_id = message(dynamic_case, key="review-after-correction")
    assert run(dynamic_case, review_id, AlternativeProvider()).status == "completed"
    rereview = client.get(f"/api/v1/agent-runs/{review_id}/checks").json()
    assert (
        next(c for c in rereview["checks"] if c["supplier"] == "vendor-a")[
            "business_status"
        ]
        == "资料冲突"
    )
    assert not any(c["is_formal"] for c in rereview["checks"])


@pytest.mark.parametrize("dynamic_case", ["苹果 数量12箱或8箱"], indirect=True)
def test_correction_preserves_incomplete_extraction_warning(dynamic_case):
    from backend.app.criteria.facts import FactCandidate
    from backend.app.providers.base import ProviderError

    class IncompleteProvider(ScriptProvider):
        def extract(self, *args, **kwargs):
            if self.extractions:
                raise ProviderError("rate_limit")
            output = super().extract(*args, **kwargs)
            for fact in output.tool_calls[0]["args"].get("facts", []):
                fact["unit"] = "箱"
            return output

    confirm(dynamic_case)
    client, _, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="partial-review")
    assert run(dynamic_case, run_id, IncompleteProvider()).status == "partial"
    facts = client.get(f"/api/v1/agent-runs/{run_id}/facts").json()
    old = next(g for g in facts["fact_sets"] if g["supplier"] == "vendor-a")["facts"][0]
    candidate = {k: old[k] for k in FactCandidate.model_fields}
    candidate.update(value="8", unit="箱")
    response = client.post(f"/api/v1/agent-runs/{run_id}/fact-corrections",
        headers={"Idempotency-Key": "partial-correction"}, json={
            "fact_id": old["fact_id"], "candidate": candidate, "reason": "同源备选",
            "expected_revision": 3, "expected_payload_digest": facts["payload_digest"],
        })
    assert response.status_code == 201, response.text
    detail = client.get(f"/api/v1/agent-runs/{run_id}").json()
    assert detail["status"] == "partial"
    assert "部分完成，未覆盖项仍为处理未完成" in detail["answer_text"]
    assert "结论仍需人工复核" in detail["answer_text"]
    checks = client.get(f"/api/v1/agent-runs/{run_id}/checks").json()["checks"]
    assert not any(c["is_formal"] for c in checks)
    assert any(c["processing_status"] != "complete" for c in checks)


@pytest.mark.parametrize("dynamic_case", ["苹果 数量12箱或8箱或10箱"], indirect=True)
def test_consecutive_corrections_record_the_selected_old_fact(dynamic_case):
    from backend.app.criteria.facts import FactCandidate

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="review-consecutive")
    run(dynamic_case, run_id, ScriptProvider())
    url = f"/api/v1/agent-runs/{run_id}/fact-corrections"
    prior_id = None
    for index, value in enumerate(("8", "10")):
        facts = client.get(f"/api/v1/agent-runs/{run_id}/facts").json()
        group = next(g for g in facts["fact_sets"] if g["supplier"] == "vendor-a")
        old = next((f for f in group["facts"] if f["fact_id"] == prior_id), group["facts"][0])
        candidate = {k: old[k] for k in FactCandidate.model_fields}
        candidate["value"] = value
        candidate["unit"] = "箱"
        response = client.post(url, headers={"Idempotency-Key": f"consecutive-{index}"}, json={
            "fact_id": old["fact_id"], "candidate": candidate, "reason": "同源备选",
            "expected_revision": 3, "expected_payload_digest": facts["payload_digest"],
        })
        assert response.status_code == 201, response.text
        audit = client.get(url).json()
        record = next(c for c in audit["corrections"] if c["correction_id"] == response.json()["correction_id"])
        assert record["old_fact"]["fact_id"] == old["fact_id"]
        prior_id = record["corrected_fact"]["fact_id"]
    # Historical records without the target ID must not guess the old value.
    with engine.begin() as connection:
        records = connection.execute(select(dynamic_corrections).where(
            dynamic_corrections.c.run_id == run_id)).mappings().all()
        for record in records:
            old_payload = dict(record["old_payload_json"])
            old_payload.pop("correction_target_fact_id")
            connection.execute(update(dynamic_corrections).where(
                dynamic_corrections.c.id == record["id"]).values(old_payload_json=old_payload))
    legacy = client.get(url).json()["corrections"]
    assert next(c for c in legacy if c["corrected_fact"]["fact_id"] == prior_id)["old_fact"] is None


def test_queued_dynamic_run_facts_are_readable_without_an_extraction_payload(dynamic_case):
    confirm(dynamic_case)
    client, _, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="review-queued")
    response = client.get(f"/api/v1/agent-runs/{run_id}/facts")
    assert response.status_code == 200, response.text
    assert response.json()["available"] is False
    assert response.json()["fact_sets"] == []


@pytest.mark.parametrize("dynamic_case", ["苹果 数量12箱或8箱"], indirect=True)
@pytest.mark.parametrize("reader", ["facts", "corrections"])
def test_report_read_rejects_a_correction_committed_between_queries(dynamic_case, monkeypatch, reader):
    from backend.app.criteria.facts import FactCandidate

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="review-read-race")
    run(dynamic_case, run_id, ScriptProvider())
    repository = DynamicRepository(engine)
    verified = repository.verified
    inserted = False

    def concurrent_verified(identifier):
        nonlocal inserted
        result = verified(identifier)
        if not inserted:
            inserted = True
            row, payload, _ = result
            old = next(f for f in payload["facts"] if f["supplier"] == "vendor-a")
            candidate = {k: old[k] for k in FactCandidate.model_fields}
            candidate.update(value="8", unit="箱")
            response = client.post(f"/api/v1/agent-runs/{identifier}/fact-corrections",
                headers={"Idempotency-Key": "concurrent-correction"}, json={
                    "fact_id": old["fact_id"], "candidate": candidate, "reason": "同源备选",
                    "expected_revision": 3, "expected_payload_digest": row["payload_digest"],
                })
            assert response.status_code == 201, response.text
        return result

    monkeypatch.setattr(repository, "verified", concurrent_verified)
    with pytest.raises(RepositoryError, match="run_changed"):
        getattr(repository, reader)(run_id)


def test_source_followup_reuses_validated_report_only_in_same_conversation(
    dynamic_case,
):
    confirm(dynamic_case)
    client, _, case_id, _, _ = dynamic_case
    run_id = message(dynamic_case, key="review-1")
    run(dynamic_case, run_id, ScriptProvider())
    followup = message(
        dynamic_case, key="followup", content="为什么数量不满足？依据在哪里？"
    )
    provider = ScriptProvider(decision="get_checklist")
    assert run(dynamic_case, followup, provider).status == "completed"
    assert provider.extractions == 0
    assert (
        len(client.get(f"/api/v1/agent-runs/{followup}/checks").json()["checks"]) == 2
    )
    new_conversation = client.post(f"/api/v1/cases/{case_id}/conversations").json()[
        "id"
    ]
    other = (*dynamic_case[:4], new_conversation)
    new_id = message(other, key="new-conversation")
    isolated = ScriptProvider(decision="get_checklist")
    run(other, new_id, isolated)
    assert isolated.extractions == 2


def test_retry_keeps_verified_batches_and_source_manifest_is_not_optional(dynamic_case):
    from backend.app.evidence.repository import digest
    from backend.app.providers.base import ProviderError

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case

    class SecondBatchFails(ScriptProvider):
        def extract(self, *args, **kwargs):
            if self.extractions == 1:
                self.failure = ProviderError("unavailable")
            return super().extract(*args, **kwargs)

    run_id = message(dynamic_case, key="partial")
    assert run(dynamic_case, run_id, SecondBatchFails()).status == "partial"
    retry = client.post(
        f"/api/v1/agent-runs/{run_id}/retry",
        headers={"Idempotency-Key": "retry"},
        json={"expected_revision": 3},
    )
    assert retry.status_code == 202, retry.text
    provider = ScriptProvider()
    assert run(dynamic_case, run_id, provider).status == "completed"
    assert provider.extractions == 1
    with engine.begin() as connection:
        payload = deepcopy(
            connection.scalar(
                select(dynamic_runs.c.payload_json).where(
                    dynamic_runs.c.run_id == run_id
                )
            )
        )
        payload["batch_manifest"].pop()
        connection.execute(
            update(dynamic_runs)
            .where(dynamic_runs.c.run_id == run_id)
            .values(payload_json=payload, payload_digest=digest(payload))
        )
    assert client.get(f"/api/v1/agent-runs/{run_id}/export").status_code == 409


def test_concurrent_cancel_keeps_known_usage_without_publishing(dynamic_case):
    from backend.app.db.models import run_attempts

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="cancel")

    class CancellingProvider(ScriptProvider):
        def extract(self, *args, **kwargs):
            reply = super().extract(*args, **kwargs)
            assert client.post(f"/api/v1/agent-runs/{run_id}/cancel").status_code == 200
            return reply

    report = run(dynamic_case, run_id, CancellingProvider())
    assert report.status == "cancelled"
    with engine.connect() as connection:
        usage = connection.scalar(
            select(run_attempts.c.usage_json).where(run_attempts.c.run_id == run_id)
        )
        assert (
            connection.scalar(
                select(dynamic_runs.c.payload_json).where(
                    dynamic_runs.c.run_id == run_id
                )
            )
            is None
        )
    assert usage["model_calls"] == 2 and usage["actual_input"] == 200


def test_cancel_after_first_terminal_check_still_retains_usage(
    dynamic_case, monkeypatch
):
    from backend.app.db.models import run_attempts

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="late-cancel")
    original = DynamicRepository.retain_cancelled_usage
    calls = 0

    def cancel_after_check(self, claim, result, provider):
        nonlocal calls
        retained = original(self, claim, result, provider)
        calls += 1
        if calls == 1:
            assert not retained
            assert client.post(f"/api/v1/agent-runs/{run_id}/cancel").status_code == 200
        return retained

    monkeypatch.setattr(DynamicRepository, "retain_cancelled_usage", cancel_after_check)
    report = run(dynamic_case, run_id, ScriptProvider())
    assert report.status == "cancelled" and calls == 2
    with engine.connect() as connection:
        usage = connection.scalar(
            select(run_attempts.c.usage_json).where(run_attempts.c.run_id == run_id)
        )
    assert usage["model_calls"] == 3 and usage["actual_input"] == 300


def test_first_model_failure_preserves_readable_incomplete_matrix_and_can_retry(
    dynamic_case,
):
    from backend.app.providers.base import ProviderError

    confirm(dynamic_case)
    client, _, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="first-failure")

    class FailedDecision(ScriptProvider):
        def decide(self, *args, **kwargs):
            raise ProviderError("timeout", outcome_unknown=True)

    assert run(dynamic_case, run_id, FailedDecision()).status == "partial"
    checks = client.get(f"/api/v1/agent-runs/{run_id}/checks")
    assert checks.status_code == 200, checks.text
    assert len(checks.json()["checks"]) == 2
    assert all(
        not c["is_formal"] and c["business_status"] is None
        for c in checks.json()["checks"]
    )
    assert (
        client.post(
            f"/api/v1/agent-runs/{run_id}/retry",
            headers={"Idempotency-Key": "retry-first"},
            json={"expected_revision": 3},
        ).status_code
        == 202
    )
    assert run(dynamic_case, run_id, ScriptProvider()).status == "completed"


def test_completed_batch_facts_cannot_be_omitted_from_recomputed_payload(dynamic_case):
    from backend.app.criteria.facts import evaluate_checklist
    from backend.app.evidence.repository import digest

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="missing-fact")
    run(dynamic_case, run_id, ScriptProvider())
    row, payload, repo = DynamicRepository(engine).verified(run_id)
    payload = deepcopy(payload)
    payload["facts"].pop()
    payload["checks"] = evaluate_checklist(
        payload["definition"],
        payload["facts"],
        payload["coverage"],
        snapshot_id=repo.snapshot_id,
    )
    with engine.begin() as connection:
        connection.execute(
            update(dynamic_runs)
            .where(dynamic_runs.c.run_id == row["id"])
            .values(payload_json=payload, payload_digest=digest(payload))
        )
    response = client.get(f"/api/v1/agent-runs/{run_id}/checks")
    assert response.status_code == 409, response.text


@pytest.mark.parametrize('completed_claim', [False, True])
def test_progress_rejects_batch_claim_without_manifest(dynamic_case, completed_claim):
    from backend.app.evidence.repository import digest
    from backend.app.providers.base import ProviderError

    class FailedDecision(ScriptProvider):
        def decide(self, *args, **kwargs):
            raise ProviderError('timeout', outcome_unknown=True)

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key='unplanned-progress')
    assert run(dynamic_case, run_id, FailedDecision()).status == 'partial'
    _, payload, _ = DynamicRepository(engine).verified(run_id)
    assert not payload['facts'] and not any(payload['coverage'].values())
    payload = deepcopy(payload)
    payload['batch_manifest'] = []
    payload['batch_coverage'] = [{'batch_id': 'x' * 64, 'complete': completed_claim}]
    with engine.begin() as connection:
        connection.execute(update(dynamic_runs).where(dynamic_runs.c.run_id == run_id)
                           .values(payload_json=payload, payload_digest=digest(payload)))
    response = client.get(f'/api/v1/agent-runs/{run_id}/progress')
    assert response.status_code == 409, response.text


def test_export_rejects_copied_source_even_if_database_changes_back(
    dynamic_case, monkeypatch
):
    from backend.app.db.models import evidence_blocks

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="export-aba")
    run(dynamic_case, run_id, ScriptProvider())
    original = DynamicRepository.verified
    calls, evidence_id = 0, None

    def source_aba(self, requested_id):
        nonlocal calls, evidence_id
        calls += 1
        if calls == 2:
            with engine.begin() as connection:
                connection.execute(
                    update(evidence_blocks)
                    .where(evidence_blocks.c.id == evidence_id)
                    .values(source_text="苹果 数量12箱")
                )
        result = original(self, requested_id)
        if calls == 1:
            _, payload, _ = result
            evidence_id = next(
                f for f in payload["facts"] if f["supplier"] == "vendor-a"
            )["evidence_refs"][0]
            with engine.begin() as connection:
                connection.execute(
                    update(evidence_blocks)
                    .where(evidence_blocks.c.id == evidence_id)
                    .values(source_text="苹果 数量8箱")
                )
        return result

    monkeypatch.setattr(DynamicRepository, "verified", source_aba)
    response = client.get(f"/api/v1/agent-runs/{run_id}/export")
    assert response.status_code == 409, response.text


def test_cancel_before_interrupt_cas_keeps_usage(dynamic_case, monkeypatch):
    from backend.app.db.models import run_attempts

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="interrupt-cancel")
    retain, validate = (
        DynamicRepository.retain_cancelled_usage,
        DynamicRepository.validate,
    )
    calls = 0

    def retain_then_cancel(self, claim, result, provider):
        nonlocal calls
        kept = retain(self, claim, result, provider)
        calls += 1
        if calls == 2:
            assert client.post(f"/api/v1/agent-runs/{run_id}/cancel").status_code == 200
        return kept

    def fail_terminal_validation(self, claim):
        if calls:
            raise RepositoryError("memory_context_changed")
        return validate(self, claim)

    monkeypatch.setattr(DynamicRepository, "retain_cancelled_usage", retain_then_cancel)
    monkeypatch.setattr(DynamicRepository, "validate", fail_terminal_validation)
    assert run(dynamic_case, run_id, ScriptProvider()).status == "cancelled"
    assert calls == 3
    with engine.connect() as connection:
        usage = connection.scalar(
            select(run_attempts.c.usage_json).where(run_attempts.c.run_id == run_id)
        )
    assert usage["model_calls"] == 3 and usage["actual_input"] == 300


def test_export_rejects_source_drift_between_checks_and_source_read(
    dynamic_case, monkeypatch
):
    from backend.app.db.models import evidence_blocks

    confirm(dynamic_case)
    client, engine, _, _, _ = dynamic_case
    run_id = message(dynamic_case, key="export-race")
    run(dynamic_case, run_id, ScriptProvider())
    original = DynamicRepository.verified
    calls = 0

    def change_after_verification(self, requested_id):
        nonlocal calls
        result = original(self, requested_id)
        calls += 1
        if calls == 1:
            _, payload, _ = result
            fact = next(f for f in payload["facts"] if f["supplier"] == "vendor-a")
            with engine.begin() as connection:
                connection.execute(
                    update(evidence_blocks)
                    .where(evidence_blocks.c.id == fact["evidence_refs"][0])
                    .values(source_text="苹果 数量8箱")
                )
        return result

    monkeypatch.setattr(DynamicRepository, "verified", change_after_verification)
    response = client.get(f"/api/v1/agent-runs/{run_id}/export")
    assert response.status_code == 409, response.text
