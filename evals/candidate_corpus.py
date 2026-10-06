"""Offline synthetic candidates and original-source label integrity checks.

No production decision engine, Provider, parser subprocess, or .env is used.
Proposed gold is not ground truth until independently reviewed by a human.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from datetime import UTC, datetime
from hashlib import sha256
from importlib.metadata import version
from pathlib import Path
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field, model_validator
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

from backend.app.field_contract import BusinessStatus, FieldKey
from backend.app.parsing import MAX_FILE_BYTES, parse_xlsx

from .acceptance.candidate_blueprints import FIELDS, NAMES, PACKS

DESIGN = Path(__file__).parent / "acceptance"
DEFAULT_ROOT = Path("data/m5-candidates")
FILENAME = r"^[a-z0-9][a-z0-9-]*\.(pdf|xlsx)$"
ID = r"^[a-z0-9][a-z0-9-]{0,63}$"


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)


class Document(StrictModel):
    filename: str = Field(pattern=FILENAME)
    role: str = Field(pattern=r"^(requirement|quote|spec)$")
    supplier: str | None
    model: str | None
    sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    bytes: int = Field(gt=0, le=MAX_FILE_BYTES)
    pages: int = Field(ge=0, le=30)


class Supplier(StrictModel):
    id: str = Field(pattern=ID)
    model: str = Field(min_length=1, max_length=80)


class InputManifest(StrictModel):
    schema_version: int = Field(ge=1, le=1)
    pack_id: str = Field(pattern=ID)
    dataset_kind: str = Field(pattern=r"^synthetic_candidate$")
    suppliers: list[Supplier] = Field(min_length=1, max_length=3)
    documents: list[Document] = Field(min_length=3, max_length=7)
    question: str = Field(min_length=1, max_length=500)

    @model_validator(mode="after")
    def check_scope(self):
        suppliers = {s.id: s.model for s in self.suppliers}
        if len(suppliers) != len(self.suppliers):
            raise ValueError("duplicate_supplier")
        if len({d.filename for d in self.documents}) != len(self.documents):
            raise ValueError("duplicate_document")
        if sum(d.pages for d in self.documents) > 30:
            raise ValueError("case_page_limit")
        if sum(d.role == "requirement" for d in self.documents) != 1:
            raise ValueError("requirement_scope")
        for d in self.documents:
            if d.role == "requirement":
                if d.supplier is not None or d.model is not None:
                    raise ValueError("requirement_scope")
            elif d.supplier not in suppliers or d.model != suppliers[d.supplier]:
                raise ValueError("supplier_scope")
        for sid in suppliers:
            if Counter(d.role for d in self.documents if d.supplier == sid) != {"quote": 1, "spec": 1}:
                raise ValueError("supplier_documents")
        return self


class Anchor(StrictModel):
    document: str = Field(pattern=FILENAME)
    text: str = Field(min_length=1, max_length=2000)
    page: int | None = Field(default=None, ge=1, le=4096)
    sheet: str | None = Field(default=None, min_length=1, max_length=31)
    cell: str | None = Field(default=None, pattern=r"^[A-Z]{1,2}[1-9][0-9]{0,3}$")


class GoldResult(StrictModel):
    supplier: str = Field(pattern=ID)
    model: str = Field(min_length=1, max_length=80)
    field_key: FieldKey
    processing_status: str = Field(pattern=r"^complete$")
    business_status: BusinessStatus
    rationale: str = Field(min_length=1, max_length=1000)
    evidence: list[Anchor] = Field(min_length=1, max_length=12)
    absence_documents: list[str]


class ProposedGold(StrictModel):
    schema_version: int = Field(ge=1, le=1)
    pack_id: str = Field(pattern=ID)
    review_state: str = Field(pattern=r"^pending_independent_review$")
    author: str = Field(pattern=r"^agent_proposed$")
    results: list[GoldResult] = Field(min_length=8, max_length=24)


class PackEntry(StrictModel):
    id: str = Field(pattern=ID)
    partition: str = Field(pattern=r"^(calibration|holdout_candidate)$")
    family: str = Field(pattern=ID)
    theme: str = Field(min_length=1, max_length=200)
    manifest_sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    gold_sha256: str = Field(pattern=r"^[a-f0-9]{64}$")


class CorpusIndex(StrictModel):
    schema_version: int = Field(ge=1, le=1)
    dataset_kind: str = Field(pattern=r"^synthetic_candidates$")
    independence: str = Field(pattern=r"^author_seen_requires_human_review$")
    packs: list[PackEntry] = Field(min_length=1, max_length=1000)
    provenance: dict
    performance_pack: str = Field(pattern=ID)


def _digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _safe(root: Path, name: str) -> Path:
    path = (root / name).resolve()
    if not path.is_relative_to(root.resolve()):
        raise ValueError("asset_path")
    return path


def _lines(text: str, font: str, width: float) -> list[str]:
    lines, current = [], ""
    for char in text:
        if pdfmetrics.stringWidth(current + char, font, 11) > width:
            lines.append(current)
            current = char
        else:
            current += char
    return lines + [current]


def _pdf(path: Path, title: str, metadata: list[str], fields: dict, notes: list[str], annex: bool):
    """Return source anchors; keep annotations out of the source document."""
    items = list(fields.items())
    groups = [items[:3], items[3:]] if annex else [items]
    if not groups[-1]:
        groups = groups[:1]
    page = Canvas(str(path), pagesize=A4, invariant=1, pageCompression=1)
    page.setAuthor("DocuCheck synthetic candidate - human review pending")
    page.setTitle(title)
    anchors = {}
    for number, group in enumerate(groups, 1):
        width, height = A4
        page.setFont("M5CJK", 17)
        page.drawString(48, height - 60, title)
        page.setFont("M5CJK", 10)
        page.drawString(48, height - 84, "合成候选资料 / 虚构组织及产品 / 仅供测试，不代表真实采购")
        y = height - 116
        page.setFont("M5CJK", 11)
        for line in metadata:
            page.drawString(48, y, line)
            y -= 23
        if annex:
            page.drawString(48, y, f"附件 {number}：{'配置说明' if number == 1 else '服务与商务约定'}")
            y -= 36
        else:
            y -= 20
        for key, text in group:
            for line in _lines(text, "M5CJK", width - 96):
                if y < 100:
                    raise ValueError("layout_overflow")
                page.drawString(48, y, line)
                y -= 21
            anchors[key] = {"document": path.name, "page": number, "text": text}
            y -= 19
        if number == len(groups):
            for text in notes:
                for line in _lines(text, "M5CJK", width - 96):
                    if y < 100:
                        raise ValueError("layout_overflow")
                    page.drawString(48, y, line)
                    y -= 21
                y -= 12
        page.setFont("M5CJK", 9)
        page.drawString(48, 42, "此文件是完整的候选原件；未附未声明的配置或承诺。")
        page.drawRightString(width - 48, 42, f"第 {number} 页 / 共 {len(groups)} 页")
        page.showPage()
    page.save()
    return anchors


def _xlsx(path: Path, supplier: str, model: str, fields: dict, annex: bool):
    workbook = Workbook()
    workbook.properties.creator = "DocuCheck synthetic candidate"
    workbook.properties.created = workbook.properties.modified = datetime(2026, 10, 1, tzinfo=UTC)
    sheet = workbook.active
    sheet.title = "商务报价" if annex else "报价明细"
    rows = [
        ["合成候选报价单", "虚构组织及产品，仅供测试"],
        ["供应商", supplier], ["产品型号", model],
    ]
    if annex:
        rows += [["适用说明", "以下均为该型号本次供货承诺"], ["报价版本", "2026-10-01 / 唯一有效版本"], []]
    rows += [["项目", "本次承诺原文"]]
    anchors = {}
    for key, text in fields.items():
        rows.append([NAMES[FIELDS.index(key)], text])
        anchors[key] = {"document": path.name, "sheet": sheet.title, "cell": f"B{len(rows)}", "text": text}
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 19
    sheet.column_dimensions["B"].width = 68
    sheet.freeze_panes = "B5"
    for row in sheet:
        sheet.row_dimensions[row[0].row].height = 33
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row[0].fill = PatternFill("solid", fgColor="E8EEF4")
    workbook.save(path)
    workbook.close()
    # openpyxl uses wall-clock ZIP timestamps and a fresh modified core timestamp.
    # Fix both to make hash comparisons meaningful in the same pinned environment.
    with ZipFile(path) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
    members["docProps/core.xml"] = re.sub(
        rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
        b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" xsi:type="dcterms:W3CDTF">2026-10-01T00:00:00Z</dcterms:modified>',
        members["docProps/core.xml"],
    )
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in sorted(members.items()):
            info = ZipInfo(name, date_time=(2026, 10, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, content)
    return anchors


def build_corpus(root: Path, *, font_path: Path) -> None:
    """Create a new directory; never overwrite or self-seal a reviewed corpus."""
    blueprint = json.loads((DESIGN / "candidate_gold.json").read_text(encoding="utf-8"))
    if blueprint["field_order"] != list(FIELDS):
        raise ValueError("gold_field_order")
    gold_rows = {(r["pack"], r["supplier"]): r for r in blueprint["rows"]}
    pdfmetrics.registerFont(TTFont("M5CJK", str(font_path)))
    root.mkdir(parents=True, exist_ok=False)
    index = {"schema_version": 1, "dataset_kind": "synthetic_candidates",
             "independence": "author_seen_requires_human_review", "packs": [],
             "provenance": {"blueprints_sha256": _digest(DESIGN / "candidate_blueprints.py"),
                            "proposed_gold_sha256": _digest(DESIGN / "candidate_gold.json"),
                            "font_sha256": _digest(font_path),
                            "dependencies": {p: version(p) for p in ("reportlab", "openpyxl", "pypdf")}}}
    review_rows = []
    guide = ["# M5-A 人工复核工作表", "", "全部答案是 Agent 拟定标签，尚未独立复核。原件均为新合成资料。",
             "逐格打开原件核对原文、型号、口径与状态；尤其注意缺失范围和冲突双方。",
             "在 review.csv 填 review_decision=accepted 或 changed、真实复核人/日期/意见；changed 必须先修正对应 gold，再重做完整校验。",
             "校验工具仅核对资产与原文定位，不能证明状态标注正确。用户确认封存前 sealed 始终为0。", ""]
    for pack in PACKS:
        pack_id = pack["id"]
        directory = root / "inputs" / pack_id
        directory.mkdir(parents=True)
        annex = pack["family"] == "annex-v2"
        documents, suppliers, results = [], [], []
        req_path = directory / "requirements.pdf"
        req_anchors = _pdf(req_path, "采购需求说明", [f"采购任务：{pack_id}"], pack["requirements"], [], annex)
        documents.append(_document(req_path, "requirement", None, None))
        for offset, raw in enumerate(pack["suppliers"]):
            sid = f"vendor-{chr(ord('a') + offset)}"
            model = f"MC-{pack_id[3:].upper()}-{offset + 1}"
            suppliers.append({"id": sid, "model": model})
            spec_path = directory / f"{sid}-spec.pdf"
            spec_fields = {**raw["spec"], **raw.get("spec_extra", {})}
            spec_anchors = _pdf(spec_path, "供货规格与服务说明", [f"供应商：{sid}", f"本次产品型号：{model}"],
                                spec_fields, raw.get("notes", []), annex)
            quote_path = directory / f"{sid}-quote.xlsx"
            quote_anchors = _xlsx(quote_path, sid, model, {**raw["quote"], **raw.get("quote_extra", {})}, annex)
            documents += [_document(quote_path, "quote", sid, model), _document(spec_path, "spec", sid, model)]
            proposed = gold_rows[(pack_id, sid)]
            if len(proposed["statuses"]) != 8 or len(proposed["reasons"]) != 8:
                raise ValueError("gold_field_coverage")
            for field, code, reason in zip(FIELDS, proposed["statuses"], proposed["reasons"], strict=True):
                evidence = [req_anchors[field]]
                evidence += [a[field] for a in (spec_anchors, quote_anchors) if field in a]
                results.append({"supplier": sid, "model": model, "field_key": field,
                                "processing_status": "complete", "business_status": blueprint["codes"][code],
                                "rationale": reason, "evidence": evidence,
                                "absence_documents": [quote_path.name, spec_path.name] if code == "I" else []})
                review_rows.append({"pack": pack_id, "supplier": sid, "field": field,
                                    "proposed_status": blueprint["codes"][code], "review_decision": "pending",
                                    "reviewer": "", "reviewed_on": "", "comment": ""})
        manifest_path = directory / "manifest.json"
        gold_path = root / "gold" / f"{pack_id}.json"
        manifest = InputManifest(schema_version=1, pack_id=pack_id, dataset_kind="synthetic_candidate",
                                 suppliers=suppliers, documents=documents,
                                 question="请完整审查全部供应商的八项指标，并列出每家每项结果与来源。")
        gold = ProposedGold(schema_version=1, pack_id=pack_id, review_state="pending_independent_review",
                            author="agent_proposed", results=results)
        _json(manifest_path, manifest.model_dump(mode="json"))
        _json(gold_path, gold.model_dump(mode="json"))
        guide += [f"## {pack_id} · {pack['theme']}", "", f"分组：{pack['partition']}；版式家族：{pack['family']}。",
                  "", f"[输入清单](../inputs/{pack_id}/manifest.json) · [拟定gold](../gold/{pack_id}.json)", "",
                  "| 供应商/字段 | 拟定状态 | 必要原文及定位 | 拟定理由 |", "| --- | --- | --- | --- |"]
        for result in results:
            sources = []
            for anchor in result["evidence"]:
                location = f"第{anchor['page']}页" if "page" in anchor else f"{anchor['sheet']}!{anchor['cell']}"
                sources.append(f"[{anchor['document']}](../inputs/{pack_id}/{anchor['document']}) {location}：{anchor['text']}")
            if result["absence_documents"]:
                sources.append("缺失须完整检查：" + "、".join(result["absence_documents"]))
            guide.append(f"| {result['supplier']}/{NAMES[FIELDS.index(result['field_key'])]} | {result['business_status']} | "
                         + "<br>".join(sources) + f" | {result['rationale']} |")
        guide.append("")
        index["packs"].append({"id": pack_id, "partition": pack["partition"], "family": pack["family"],
                               "theme": pack["theme"], "manifest_sha256": _digest(manifest_path),
                               "gold_sha256": _digest(gold_path)})
    index["performance_pack"] = "m5-c01"
    _json(root / "corpus.json", index)
    review_dir = root / "review"
    review_dir.mkdir()
    (review_dir / "review-guide.md").write_text("\n".join(guide) + "\n", encoding="utf-8")
    with (review_dir / "review.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(review_rows[0]))
        writer.writeheader()
        writer.writerows(review_rows)
    _json(review_dir / "attestation-template.json", {
        "state": "pending", "reviewer": "", "reviewed_on": "", "code_baseline": "",
        "corpus_sha256": _digest(root / "corpus.json"),
        "contamination_history": "Agent authored and inspected all candidates; no model runs or prompt/rule tuning on these candidates yet.",
        "instructions": "Complete every row against originals. User must explicitly confirm independence and freezing. This template is not a seal.",
    })


def _document(path: Path, role: str, supplier: str | None, model: str | None):
    return {"filename": path.name, "role": role, "supplier": supplier, "model": model,
            "sha256": _digest(path), "bytes": path.stat().st_size,
            "pages": len(PdfReader(path, strict=True).pages) if path.suffix == ".pdf" else 0}


def runtime_manifest(pack_dir: Path) -> dict:
    """Input-only interface: no traversal into corpus index, review, or gold."""
    manifest = InputManifest.model_validate_json(_safe(pack_dir, "manifest.json").read_text(encoding="utf-8"))
    for doc in manifest.documents:
        path = _safe(pack_dir, doc.filename)
        if path.stat().st_size != doc.bytes or _digest(path) != doc.sha256:
            raise ValueError("asset_hash")
    return manifest.model_dump(mode="json")


def validate_corpus(root: Path) -> dict:
    """Validate source anchoring, not the semantic correctness of proposed gold."""
    errors, statuses, counts = [], Counter(), Counter()
    def error(pack, code):
        errors.append({"pack": pack, "code": code})

    try:
        index = CorpusIndex.model_validate_json(_safe(root, "corpus.json").read_text(encoding="utf-8")).model_dump()
    except (ValueError, OSError):
        error("corpus", "index_schema")
        return _report(counts, statuses, errors)
    families, ids, seen_assets = {}, set(), {}
    development_hashes = {_digest(p) for p in (Path(__file__).resolve().parents[1] / "examples").rglob("*")
                          if p.suffix in {".pdf", ".xlsx"} and p.is_file()}
    for entry in index["packs"]:
        pack_id = entry["id"]
        if not re.fullmatch(ID, pack_id) or pack_id in ids:
            error(pack_id, "pack_identity")
            continue
        ids.add(pack_id)
        partition, family = entry["partition"], entry["family"]
        if partition not in {"calibration", "holdout_candidate"}:
            error(pack_id, "partition")
        if family in families and families[family] != partition:
            error(pack_id, "family_split")
        families[family] = partition
        try:
            pack_dir = _safe(root, f"inputs/{pack_id}")
            manifest_path = _safe(pack_dir, "manifest.json")
            raw_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            if (not isinstance(raw_manifest, dict) or not isinstance(raw_manifest.get("documents"), list)
                    or any(not isinstance(d, dict) for d in raw_manifest["documents"])):
                raise ValueError("manifest_structure")
            if any(not re.fullmatch(FILENAME, d.get("filename", "")) for d in raw_manifest.get("documents", [])):
                error(pack_id, "asset_path")
            manifest = InputManifest.model_validate(raw_manifest)
        except (ValueError, OSError, TypeError):
            error(pack_id, "input_schema")
            continue
        if manifest.pack_id != pack_id:
            error(pack_id, "pack_identity")
        if _digest(manifest_path) != entry["manifest_sha256"]:
            error(pack_id, "manifest_hash")
        counts["packs"] += 1
        counts[partition] += 1
        counts["documents"] += len(manifest.documents)
        counts["pages"] += sum(d.pages for d in manifest.documents)
        counts["suppliers"] += len(manifest.suppliers)
        docs = {d.filename: d for d in manifest.documents}
        texts = {}
        for doc in manifest.documents:
            try:
                path = _safe(pack_dir, doc.filename)
                if path.stat().st_size != doc.bytes or _digest(path) != doc.sha256:
                    error(pack_id, "asset_hash")
                    continue
                if doc.sha256 in development_hashes:
                    error(pack_id, "development_duplicate")
                if doc.sha256 in seen_assets:
                    error(pack_id, "duplicate_asset")
                seen_assets[doc.sha256] = pack_id
                if path.suffix == ".pdf":
                    with path.open("rb") as stream:
                        if stream.read(5) != b"%PDF-":
                            raise ValueError("pdf_signature")
                    reader = PdfReader(path, strict=True)
                    if reader.is_encrypted or len(reader.pages) != doc.pages or not doc.pages:
                        raise ValueError("pdf_pages")
                    texts[doc.filename] = {i: p.extract_text() for i, p in enumerate(reader.pages, 1)}
                else:
                    parse_xlsx(path, document_id=doc.filename)
                    with path.open("rb") as stream:
                        book = load_workbook(stream, read_only=True, data_only=False)
                        texts[doc.filename] = {(s.title, c.coordinate): c.value
                                               for s in book for row in s for c in row if c.value is not None}
                        book.close()
                    if doc.pages != 0:
                        raise ValueError("xlsx_pages")
                if doc.model and doc.model not in " ".join(str(t) for t in texts[doc.filename].values()):
                    error(pack_id, "model_source")
            except (ValueError, OSError, KeyError, PdfReadError, BadZipFile):
                error(pack_id, "asset_structure")
        try:
            gold_path = _safe(root, f"gold/{pack_id}.json")
            gold = ProposedGold.model_validate_json(gold_path.read_text(encoding="utf-8"))
        except (ValueError, OSError):
            error(pack_id, "gold_schema")
            continue
        if _digest(gold_path) != entry["gold_sha256"]:
            error(pack_id, "gold_hash")
        if gold.pack_id != pack_id:
            error(pack_id, "pack_identity")
        expected = {(s.id, k) for s in manifest.suppliers for k in FIELDS}
        actual = [(r.supplier, r.field_key) for r in gold.results]
        if len(actual) != len(expected) or set(actual) != expected:
            error(pack_id, "field_coverage")
        models = {s.id: s.model for s in manifest.suppliers}
        for result in gold.results:
            counts["results"] += 1
            statuses[result.business_status] += 1
            if result.model != models.get(result.supplier):
                error(pack_id, "model_scope")
            applicable = {d.filename for d in docs.values() if d.supplier == result.supplier}
            anchors = []
            for anchor in result.evidence:
                doc = docs.get(anchor.document)
                if doc is None or (doc.role != "requirement" and doc.supplier != result.supplier):
                    error(pack_id, "evidence_scope")
                    continue
                anchors.append((doc.role, anchor))
                if anchor.document not in texts:
                    continue
                if doc.filename.endswith(".pdf"):
                    location = anchor.page
                    valid = location is not None and anchor.sheet is None and anchor.cell is None
                else:
                    location = (anchor.sheet, anchor.cell)
                    valid = anchor.page is None and anchor.sheet is not None and anchor.cell is not None
                if not valid or location not in texts[anchor.document]:
                    error(pack_id, "evidence_location")
                    continue
                original = str(texts[anchor.document][location])
                normalized = lambda text: re.sub(r"\s+", "", text)
                matches = normalized(anchor.text) in normalized(original) if doc.filename.endswith(".pdf") else anchor.text == original
                if not matches:
                    error(pack_id, "evidence_text")
            supplier_anchors = [a for role, a in anchors if role != "requirement"]
            if not any(role == "requirement" for role, a in anchors):
                error(pack_id, "requirement_evidence")
            if result.business_status == "证据不足":
                if supplier_anchors or set(result.absence_documents) != applicable or len(result.absence_documents) != len(applicable):
                    error(pack_id, "absence_scope")
            elif result.absence_documents or not supplier_anchors:
                error(pack_id, "supplier_evidence")
            if result.business_status == "资料冲突" and len({(a.document, a.page, a.sheet, a.cell) for a in supplier_anchors}) < 2:
                error(pack_id, "conflict_sources")
    if counts["packs"] < 12 or counts["results"] < 104 or counts["holdout_candidate"] < 6:
        error("corpus", "cohort_coverage")
    performance = next((e for e in index["packs"] if e["id"] == index["performance_pack"]), None)
    if performance is None:
        error("corpus", "performance_pack")
    else:
        try:
            perf = runtime_manifest(_safe(root, f"inputs/{performance['id']}"))
            if len(perf["documents"]) != 5 or len(perf["suppliers"]) != 2 or performance["partition"] != "calibration":
                error("corpus", "performance_pack")
        except (ValueError, OSError):
            error("corpus", "performance_pack")
    return _report(counts, statuses, errors)


def _report(counts, statuses, errors):
    return {"scope": "candidate_asset_and_proposed_label_integrity", "candidate_packs": counts["packs"],
            "holdout_candidates": counts["holdout_candidate"], "supplier_field_cases": counts["results"],
            "suppliers": counts["suppliers"], "documents": counts["documents"], "pdf_pages": counts["pages"],
            "proposed_status_distribution": dict(statuses), "errors": errors,
            "candidate_preparation_ready": not errors,
            "reviewed_packs": 0, "sealed_packs": 0, "business_acceptance_measured": False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument("--build", action="store_true", help="Create a new candidate directory without overwriting")
    parser.add_argument("--font", type=Path, default=Path("C:/Windows/Fonts/simhei.ttf"))
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.build:
        build_corpus(args.root, font_path=args.font)
    report = validate_corpus(args.root)
    if args.output:
        _json(args.output, report)
    print(json.dumps(report, ensure_ascii=False))
    if report["errors"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
