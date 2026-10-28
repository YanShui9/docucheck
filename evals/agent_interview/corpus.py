"""Generate new synthetic originals and verify frozen handwritten expectations offline.

This module never imports product extraction/rules, loads credentials or calls models.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

from .contracts import CorpusManifest, Dataset, FileRecord, safe_relative

DATASET_PATH = Path(__file__).with_name("dataset.json")
NOTICE = "合成评估资料 / 虚构供应商和产品 / CC0"
FONT_NAME = "AgentInterviewCJK"
MAX_JSON_BYTES = 8 * 1024 * 1024
PUBLIC_CONTEXT_KEYS = frozenset({
    "selected_supplier", "selected_item", "retry_operation", "existing_report",
    "unparsed_bound_page", "coverage_complete", "current_rule_version", "actor_role",
    "can_write", "authorized_case_ids", "requested_external_case", "previous_dispatch_state",
    "historical_todo_executed", "draft_rules",
})


def _json_bytes(value) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def _read_json(path: Path):
    if (path.is_symlink() or path.is_junction() or not path.is_file()
            or path.stat().st_size > MAX_JSON_BYTES):
        raise ValueError("unsafe_json_asset")
    return json.loads(path.read_text(encoding="utf-8"))


def load_dataset(path: Path | str | None = None) -> Dataset:
    """Read the public protocol; validation never evaluates the product's rules."""
    return Dataset.model_validate(_read_json(Path(path) if path is not None else DATASET_PATH))


def _safe_path(root: Path, relative: str) -> Path:
    safe_relative(relative)
    path = root
    if root.is_symlink() or root.is_junction():
        raise ValueError("unsafe_corpus_root")
    for component in relative.split("/"):
        path = path / component
        if path.is_symlink() or path.is_junction():
            raise ValueError("unsafe_asset_path")
    if not path.resolve().is_relative_to(root.resolve()):
        raise ValueError("unsafe_asset_path")
    return path


def _font(font_path: Path | str | None) -> Path:
    if font_path is None:
        font_path = os.getenv("DOCUCHECK_CJK_FONT", "C:/Windows/Fonts/simhei.ttf")
    path = Path(font_path)
    if not path.is_file():
        raise ValueError("cjk_font_unavailable")
    pdfmetrics.registerFont(TTFont(FONT_NAME, str(path)))
    return path


def _pdf(path: Path, pages: list[list[str]]) -> None:
    document = Canvas(str(path), pagesize=A4, invariant=1)
    document.setTitle("DocuCheck synthetic procurement original")
    document.setAuthor("DocuCheck evaluation authors / public synthetic")
    for number, lines in enumerate(pages, 1):
        for index, line in enumerate(lines):
            size = 14 if index == 0 else 11
            if pdfmetrics.stringWidth(line, FONT_NAME, size) > A4[0] - 88:
                size = 9
            if pdfmetrics.stringWidth(line, FONT_NAME, size) > A4[0] - 88:
                raise ValueError("pdf_line_too_wide")
            document.setFont(FONT_NAME, size)
            document.drawString(44, A4[1] - 58 - index * 30, line)
        document.setFont(FONT_NAME, 8)
        document.drawString(44, 30, f"{NOTICE} | 第 {number} 页")
        document.showPage()
    document.save()


def _stable_xlsx(payload: bytes) -> bytes:
    """Freeze both core XML timestamps and ZIP member timestamps."""
    result = BytesIO()
    with ZipFile(BytesIO(payload)) as original, ZipFile(result, "w") as stable:
        for name in sorted(original.namelist()):
            data = original.read(name)
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>).*?(</dcterms:(?:created|modified)>)",
                              rb"\g<1>2026-10-28T00:00:00Z\g<2>", data)
            info = ZipInfo(name, date_time=(2026, 10, 28, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            stable.writestr(info, data)
    return result.getvalue()


def _xlsx(path: Path, sheets: dict[str, list[list[str]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in "ABCDEFG":
            sheet.column_dimensions[column].width = 25
        sheet.column_dimensions["H"].width = 60
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=11, bold=cell.row == 1)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="DDE6ED")
            sheet.row_dimensions[row[0].row].height = 34
    memory = BytesIO()
    workbook.save(memory)
    workbook.close()
    path.write_bytes(_stable_xlsx(memory.getvalue()))


def _agent_scenarios(dataset: Dataset) -> tuple[dict, dict]:
    """Adapter fixtures have no task id, expected answer, gold anchors or scores.

    scenario_NNN is an opaque public fixture key, mapped to tasks only in gold.
    Setup history is fixture input, not a claim of pre-existing real user history.
    """
    setups = {setup.setup_id: setup for setup in dataset.setups}
    scenarios, mapping = [], {}
    for index, task in enumerate(dataset.tasks, 1):
        key = f"scenario_{index:03d}"
        setup = setups[task.setup_id].model_dump(mode="json", exclude={"setup_id", "preconditions"})
        # Context is explicit business fixture state, never arbitrary author metadata.
        # The runtime must still replace these declarations with actual API observations.
        setup["context"] = {key: value for key, value in setup["context"].items()
                            if key in PUBLIC_CONTEXT_KEYS}
        mapping[task.task_id] = key
        scenarios.append({"scenario_key": key, "prompt": task.prompt,
                          "turns": [turn.model_dump(mode="json") for turn in task.turns],
                          "setup": setup})
    return {"schema_version": 1, "revision": dataset.revision,
            "synthetic": True, "scenarios": scenarios}, mapping


def build_corpus(output_root: Path | str, *, font_path: Path | str | None = None) -> CorpusManifest:
    """Create a fresh directory exclusively; never overwrite an existing corpus."""
    root = Path(output_root)
    if root.exists() or root.is_symlink() or root.is_junction():
        raise FileExistsError(str(root))
    dataset = load_dataset()
    _font(font_path)
    root.mkdir(parents=True, exist_ok=False)
    for source in dataset.sources:
        path = _safe_path(root, source.path)
        path.parent.mkdir(parents=True, exist_ok=True)
        if source.format == "pdf":
            _pdf(path, source.pages)
        else:
            _xlsx(path, source.sheets)
    scenarios, mapping = _agent_scenarios(dataset)
    (root / "gold").mkdir()
    _safe_path(root, "inputs/scenarios.json").write_bytes(_json_bytes(scenarios))
    dataset_bytes = DATASET_PATH.read_bytes()
    _safe_path(root, "gold/dataset.json").write_bytes(dataset_bytes)
    _safe_path(root, "gold/scenario_map.json").write_bytes(_json_bytes(mapping))
    paths = [s.path for s in dataset.sources] + [
        "inputs/scenarios.json", "gold/dataset.json", "gold/scenario_map.json"]
    files = []
    for relative in sorted(paths):
        payload = _safe_path(root, relative).read_bytes()
        files.append(FileRecord(path=relative, sha256=hashlib.sha256(payload).hexdigest(),
                                size=len(payload), role="input" if relative.startswith("inputs/") else "gold"))
    manifest = CorpusManifest(schema_version=1, revision=dataset.revision, dataset_id=dataset.dataset_id,
                              dataset_sha256=hashlib.sha256(dataset_bytes).hexdigest(),
                              files=files, task_ids=[t.task_id for t in dataset.tasks])
    (root / "manifest.json").write_bytes(_json_bytes(manifest.model_dump(mode="json")))
    return validate_corpus(root)


def validate_corpus(root: Path | str) -> CorpusManifest:
    """Verify every hash and every PDF page/line or XLSX cell anchor, not samples."""
    root = Path(root)
    manifest = CorpusManifest.model_validate(_read_json(_safe_path(root, "manifest.json")))
    dataset_path = _safe_path(root, "gold/dataset.json")
    if hashlib.sha256(dataset_path.read_bytes()).hexdigest() != manifest.dataset_sha256:
        raise ValueError("dataset_hash_mismatch")
    dataset = load_dataset(dataset_path)
    expected = {s.path for s in dataset.sources} | {
        "inputs/scenarios.json", "gold/dataset.json", "gold/scenario_map.json"}
    actual = {f.path for f in manifest.files}
    if (actual != expected or manifest.task_ids != [t.task_id for t in dataset.tasks]
            or manifest.revision != dataset.revision):
        raise ValueError("manifest_completeness_mismatch")
    discovered = set()
    for directory, folders, filenames in os.walk(root, followlinks=False):
        for name in folders + filenames:
            path = Path(directory) / name
            if path.is_symlink() or path.is_junction():
                raise ValueError("unsafe_asset_path")
        for name in filenames:
            discovered.add((Path(directory) / name).relative_to(root).as_posix())
    if discovered != expected | {"manifest.json"}:
        raise ValueError("manifest_unlisted_assets")
    for record in manifest.files:
        payload = _safe_path(root, record.path).read_bytes()
        if len(payload) != record.size or hashlib.sha256(payload).hexdigest() != record.sha256:
            raise ValueError("asset_hash_mismatch:" + record.path)
    scenarios, mapping = _agent_scenarios(dataset)
    if (_read_json(_safe_path(root, "inputs/scenarios.json")) != scenarios
            or _read_json(_safe_path(root, "gold/scenario_map.json")) != mapping):
        raise ValueError("manifest_scenario_mapping_mismatch")
    anchors = {}
    for anchor in dataset.anchors:
        anchors.setdefault(anchor.source_id, []).append(anchor)
    for source in dataset.sources:
        path = _safe_path(root, source.path)
        if source.format == "pdf":
            reader = PdfReader(path)
            if len(reader.pages) != len(source.pages):
                raise ValueError("pdf_page_count_mismatch")
            pages = [page.extract_text().splitlines() for page in reader.pages]
            for index, (actual_lines, business_lines) in enumerate(zip(pages, source.pages), 1):
                if actual_lines != business_lines + [f"{NOTICE} | 第 {index} 页"]:
                    raise ValueError("pdf_full_content_mismatch:" + source.source_id)
            for anchor in anchors.get(source.source_id, []):
                lines = pages[anchor.page - 1]
                if anchor.line > len(lines) or lines[anchor.line - 1] != anchor.text:
                    raise ValueError("pdf_anchor_text_mismatch:" + anchor.anchor_id)
        else:
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                if workbook.sheetnames != list(source.sheets):
                    raise ValueError("xlsx_sheet_set_mismatch:" + source.source_id)
                for title, rows in source.sheets.items():
                    sheet = workbook[title]
                    if (sheet.max_row != len(rows) or sheet.max_column != max(map(len, rows))):
                        raise ValueError("xlsx_matrix_dimensions_mismatch:" + source.source_id)
                    matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
                    padded = [row + [None] * (sheet.max_column - len(row)) for row in rows]
                    if matrix != padded:
                        raise ValueError("xlsx_anchor_matrix_content_mismatch:" + source.source_id)
                for anchor in anchors.get(source.source_id, []):
                    if workbook[anchor.sheet][anchor.cell].value != anchor.text:
                        raise ValueError("xlsx_anchor_text_mismatch:" + anchor.anchor_id)
            finally:
                workbook.close()
    return manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=("build", "validate"))
    parser.add_argument("root", type=Path)
    parser.add_argument("--font", type=Path)
    args = parser.parse_args(argv)
    result = build_corpus(args.root, font_path=args.font) if args.action == "build" else validate_corpus(args.root)
    print(json.dumps({"dataset_id": result.dataset_id, "revision": result.revision,
                      "tasks": len(result.task_ids),
                      "assets": len(result.files), "dataset_sha256": result.dataset_sha256,
                      "valid": True, "models_called": 0}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
