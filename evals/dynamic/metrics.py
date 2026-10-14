"""Score frozen dynamic labels against independently collected output surfaces."""

from __future__ import annotations

import unicodedata
from collections import Counter, defaultdict
from collections.abc import Mapping
from hashlib import sha256
from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from .contracts import (
    Anchor,
    Gold,
    Manifest,
    Observation,
    ObservedCell,
    ObservedSource,
    contains,
    validate_pair,
)

DETERMINISTIC = frozenset({"满足", "不满足", "资料冲突"})
SURFACES = ("api_cells", "visible_cells", "export_cells", "xlsx_export_cells")


def _text(text: str) -> str:
    return "".join(unicodedata.normalize("NFKC", text).split())


def _location(anchor: Anchor) -> tuple:
    return anchor.document, anchor.page, anchor.sheet, anchor.cell


def read_originals(manifest: Manifest, originals: Mapping[str, bytes]) -> tuple[dict, list[dict]]:
    """Read supplied bytes only; no filesystem, environment, network or subprocess access."""
    texts, errors = {}, []
    for document in manifest.documents:
        data = originals.get(document.filename)
        code = None
        if not isinstance(data, bytes):
            code = "missing_original"
        elif len(data) != document.bytes or sha256(data).hexdigest() != document.sha256:
            code = "original_hash_or_size_mismatch"
        if code:
            errors.append({"document": document.filename, "code": code})
            continue
        try:
            current = {}
            if document.pages:
                if not data.startswith(b"%PDF-"):
                    raise ValueError("not_pdf")
                reader = PdfReader(BytesIO(data), strict=True)
                if reader.is_encrypted or len(reader.pages) != document.pages:
                    raise ValueError("invalid_pdf_pages")
                for index, page in enumerate(reader.pages, 1):
                    current[document.filename, index, None, None] = page.extract_text() or ""
            else:
                with ZipFile(BytesIO(data)) as archive:
                    if len(archive.infolist()) > 2000 or sum(e.file_size for e in archive.infolist()) > 50 * 1024 * 1024:
                        raise ValueError("xlsx_archive_limit")
                    if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(archive.namelist()):
                        raise ValueError("not_xlsx")
                    if archive.testzip() is not None:
                        raise ValueError("xlsx_crc_error")
                book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet in book:
                        if sheet.max_row > 10000 or sheet.max_column > 256:
                            raise ValueError("xlsx_dimensions_limit")
                        for row in sheet.iter_rows():
                            for cell in row:
                                if cell.value is not None:
                                    current[document.filename, None, sheet.title, cell.coordinate] = str(cell.value)
                finally:
                    book.close()
            texts.update(current)
        except (BadZipFile, OSError, ValueError, KeyError, TypeError, RuntimeError, ParseError,
                PdfReadError, InvalidFileException, DefusedXmlException):
            errors.append({"document": document.filename, "code": "invalid_original"})
    return texts, errors


def _matches(anchor: Anchor, source: ObservedSource) -> bool:
    return (_location(anchor) == _location(source) and anchor.role == source.role
            and anchor.supplier_handle == source.supplier_handle and anchor.item_id == source.item_id
            and _text(anchor.text) in _text(source.text)
            and (anchor.bbox is None or anchor.bbox == source.bbox))


def _valid_source(source: ObservedSource, key: tuple, manifest: Manifest, texts: dict) -> bool:
    document = next((doc for doc in manifest.documents if doc.filename == source.document), None)
    original_text = texts.get(_location(source))
    return bool(document and document.role == source.role and document.supplier_handle == source.supplier_handle
                and document.sha256 == source.sha256 and source.supplier_handle in {None, key[1]}
                and source.item_id == key[2] and original_text is not None
                and _text(source.text) in _text(original_text)
                and (source.role == "requirement" or any(
                    binding.document == source.document and binding.item_id == key[2]
                    and contains(binding.scope, source) for binding in manifest.bindings)))


def _state_correct(cell: ObservedCell | None, expected, *, visible: bool = False) -> bool:
    if cell is None or cell.processing_status != "complete" or cell.business_status != expected.expected_business:
        return False
    if visible and cell.surface == "status_only":
        return True
    return (cell.surface == "full" and cell.is_formal == expected.expected_formal
            and cell.needs_review == expected.expected_needs_review)


def _baseline_sources(observed, manifest, sources, texts, key):
    """Buyer provenance belongs to the frozen checklist, not supplier facts."""
    confirmed = observed.confirmed_checklist
    if (confirmed is None or confirmed.version_id != observed.checklist_version_id
            or confirmed.case_id != observed.case_id or confirmed.snapshot_id != observed.snapshot_id):
        return []
    expected = manifest.checklist_definition.model_dump(mode="json")
    actual = confirmed.definition.model_dump(mode="json")
    expected_by_id = {row["criterion_id"]: row for row in expected["criteria"]}
    actual_by_id = {row["criterion_id"]: row for row in actual["criteria"]}
    if set(expected_by_id) != set(actual_by_id):
        return []
    # Only the runtime source context can differ, because its IDs come from
    # the actual parsed snapshot. Do not accept a changed rule or item.
    normalized = {**actual, "criteria": []}
    for row in expected["criteria"]:
        current = actual_by_id[row["criterion_id"]]
        normalized["criteria"].append({**current, "source": {
            **current["source"], "context": row["source"]["context"]}})
    if normalized != expected:
        return []
    criterion = actual_by_id.get(key[3])
    if criterion is None or criterion["item_id"] != key[2]:
        return []
    context = criterion["source"]["context"]
    if not context:
        return []
    refs = context["evidence_refs"]
    # This evaluation protocol resolves each buyer anchor to one real block.
    # Do not join unrelated PDF blocks or use a broad source pool as a quote.
    if len(refs) != 1:
        return []
    found = [sources.get(ref) for ref in refs]
    if any(source is None or source.role != "requirement" or source.supplier_handle is not None
           or not _valid_source(source, key, manifest, texts) for source in found):
        return []
    # The quoted baseline must be supported by a real original block; a pool
    # containing a matching block is not itself a confirmed source citation.
    return ([source.model_copy(update={"text": context["quote"]}) for source in found]
            if any(_text(context["quote"]) in _text(source.text) for source in found) else [])


def score_pack(manifest: Manifest, gold: Gold, observation: Observation | dict,
               originals: Mapping[str, bytes]) -> dict:
    """Keep the manifest's denominator through missing, partial and failed observations."""
    validate_pair(manifest, gold)
    errors = []
    raw = observation.model_dump(mode="json") if isinstance(observation, Observation) else observation
    if not isinstance(raw, dict):
        raw = {}
    try:
        observed = observation if isinstance(observation, Observation) else Observation.model_validate(observation)
    except ValidationError:
        observed = Observation(pack_id=manifest.pack_id, run_status="failed")
        errors.append("invalid_observation")
    if observed.pack_id != manifest.pack_id:
        errors.append("observation_pack_mismatch")
    if observed.error_code:
        errors.append(observed.error_code)
    errors.extend(observed.surface_errors)
    texts, original_errors = read_originals(manifest, originals)
    expected_keys = set(manifest.primary_keys())
    indexed, identity_errors = {}, []
    required_surfaces = SURFACES
    for surface in SURFACES:
        groups = defaultdict(list)
        for cell in getattr(observed, surface):
            groups[cell.key()].append(cell)
        indexed[surface] = {key: rows[0] for key, rows in groups.items() if len(rows) == 1}
        if surface not in required_surfaces:
            continue
        for key in sorted(expected_keys - set(groups)):
            identity_errors.append({"surface": surface, "key": list(key), "code": "missing_cell"})
        for key, rows in groups.items():
            if key not in expected_keys:
                identity_errors.append({"surface": surface, "key": list(key), "code": "extra_cell"})
            if len(rows) != 1:
                identity_errors.append({"surface": surface, "key": list(key), "code": "duplicate_cell"})
    source_counts = Counter(source.evidence_id for source in observed.sources)
    sources = {source.evidence_id: source for source in observed.sources if source_counts[source.evidence_id] == 1}
    citation_claims = {}
    for surface in SURFACES:
        for cell in getattr(observed, surface):
            ref_counts = Counter(cell.evidence_refs)
            for ref in cell.evidence_refs:
                source = sources.get(ref)
                valid = bool(source and ref_counts[ref] == 1 and _valid_source(source, cell.key(), manifest, texts))
                claim = (cell.key(), ref)
                citation_claims[claim] = citation_claims.get(claim, True) and valid
    citation_total = len(citation_claims)
    citation_bad = sum(not valid for valid in citation_claims.values())
    dangerous_keys, dangerous_surfaces = set(), dict.fromkeys(SURFACES, 0)
    cells, confusion = [], defaultdict(Counter)
    status_numerator = necessary_numerator = absence_numerator = 0
    rules = {criterion.criterion_id: criterion.rule.type for criterion in manifest.checklist_definition.criteria}
    completed = (observed.run_status == "completed" and observed.pack_id == manifest.pack_id
                 and "invalid_observation" not in errors)
    for expected in gold.cells:
        key = expected.key()
        api = indexed["api_cells"].get(key)
        state_correct = completed and _state_correct(api, expected)
        output_correct = all(_state_correct(indexed[surface].get(key), expected, visible=surface == "visible_cells")
                             for surface in required_surfaces)
        if api:
            output_correct = output_correct and all(
                set(indexed[surface][key].evidence_refs) == set(api.evidence_refs)
                for surface in required_surfaces if surface != "visible_cells" and key in indexed[surface])
        status_numerator += int(state_correct)
        confusion[expected.expected_business][api.business_status if api and api.processing_status == "complete"
                                              else "未完成或无结果"] += 1
        for surface in SURFACES:
            rows = raw.get(surface, [])
            if not isinstance(rows, list):
                rows = []
            meets = any(isinstance(cell, dict) and tuple(cell.get(field) for field in
                        ("pack_id", "supplier_handle", "item_id", "criterion_id")) == key
                        and cell.get("business_status") == "满足" for cell in rows)
            if expected.expected_business != "满足" and meets:
                dangerous_keys.add(key)
                dangerous_surfaces[surface] += 1
        valid_refs, invalid_refs = [], []
        refs = api.evidence_refs if api else []
        ref_counts = Counter(refs)
        for ref in refs:
            source = sources.get(ref)
            if source is None or ref_counts[ref] != 1 or not _valid_source(source, key, manifest, texts):
                invalid_refs.append(ref)
            else:
                valid_refs.append(source)
        has_baseline_channel = any((observed.confirmed_checklist, observed.case_id,
                                    observed.snapshot_id, observed.checklist_version_id))
        baseline_refs = (_baseline_sources(observed, manifest, sources, texts, key)
                         if has_baseline_channel else valid_refs)
        missing_groups = [group.group_id for group in expected.evidence_groups
                          if not any(_matches(anchor, source) for anchor in group.alternatives
                                     for source in (baseline_refs if group.purpose == "buyer_baseline" else valid_refs))]
        # Gold itself must point into the frozen original; altered labels cannot validate an invented citation.
        bad_gold_groups = [group.group_id for group in expected.evidence_groups
                           if not any(_text(anchor.text) in _text(texts.get(_location(anchor), ""))
                                      for anchor in group.alternatives)]
        missing_groups = list(dict.fromkeys([*missing_groups, *bad_gold_groups]))
        missing_coverage = []
        for scope in expected.required_coverage:
            matching = [record for record in observed.coverage
                        if (record.document, record.item_id, record.scope) == (scope.document, scope.item_id, scope.scope)]
            if (observed.coverage_unknown or len(matching) != 1 or matching[0].processing_status != "complete"
                or not any(location[0] == scope.document for location in texts)):
                missing_coverage.append(scope.model_dump(mode="json"))
        evidence_correct = state_correct and not missing_groups and not invalid_refs
        if expected.expected_business in DETERMINISTIC:
            necessary_numerator += int(evidence_correct)
        absence_correct = state_correct and not missing_coverage and not missing_groups and not invalid_refs
        if expected.expected_business == "证据不足":
            absence_numerator += int(absence_correct)
        cells.append({"key": list(key), "expected_business": expected.expected_business,
                      "category": manifest.category, "rule_type": rules[expected.criterion_id],
                      "observed_business": api.business_status if api else None,
                      "processing_status": api.processing_status if api else "missing",
                      "is_formal": api.is_formal if api else None, "needs_review": api.needs_review if api else None,
                      "status_correct": state_correct, "output_consistent": output_correct,
                      "evidence_correct": evidence_correct, "missing_evidence_groups": missing_groups,
                      "evidence_group_hits": len(expected.evidence_groups) - len(missing_groups),
                      "evidence_group_total": len(expected.evidence_groups),
                      "invalid_evidence_refs": invalid_refs, "missing_coverage": missing_coverage,
                      "strict_correct": state_correct and output_correct and not missing_groups
                                        and not invalid_refs and not missing_coverage})
    result = {
        "pack_id": manifest.pack_id, "protocol_id": manifest.protocol_id, "data_class": manifest.data_class,
        "category": manifest.category, "run_status": observed.run_status,
        "status_accuracy": {"numerator": status_numerator, "denominator": len(gold.cells)},
        "necessary_evidence": {"numerator": necessary_numerator,
                               "denominator": sum(c.expected_business in DETERMINISTIC for c in gold.cells)},
        "necessary_evidence_definition": "grounded_correct_cells",
        "anchor_group_recall": {"numerator": sum(cell["evidence_group_hits"] for cell in cells),
                                "denominator": sum(cell["evidence_group_total"] for cell in cells)},
        "absence_coverage": {"numerator": absence_numerator,
                             "denominator": sum(c.expected_business == "证据不足" for c in gold.cells)},
        "citation_errors": {"numerator": citation_bad, "denominator": citation_total},
        "citation_counting": "unique_cell_and_evidence_id_across_surfaces",
        "citation_error_rate": citation_bad / citation_total if citation_total else None,
        "dangerous_meets": len(dangerous_keys), "dangerous_meet_surfaces": dangerous_surfaces,
        "identity_errors": identity_errors, "original_errors": original_errors, "errors": errors,
        "processing_counts": dict(Counter(cell["processing_status"] for cell in cells)),
        "formal_counts": dict(Counter("formal" if cell["is_formal"] is True else "nonformal_or_unknown" for cell in cells)),
        "needs_review_count": sum(cell["needs_review"] is True for cell in cells),
        "confusion_matrix": {key: dict(counts) for key, counts in confusion.items()}, "cells": cells,
        "task_success": bool(completed and not errors and not identity_errors and not original_errors
                             and not dangerous_keys and all(cell["strict_correct"] for cell in cells)),
        "formal_independent_acceptance": False,
    }
    result["grounded_correct_cells"] = dict(result["necessary_evidence"])
    result["by_rule"] = _breakdown(cells, "rule_type")
    return result


def _breakdown(cells: list[dict], field: str) -> dict:
    groups = defaultdict(list)
    for cell in cells:
        groups[cell[field]].append(cell)
    return {label: {"status_accuracy": {"numerator": sum(c["status_correct"] for c in rows), "denominator": len(rows)},
                    "grounded_correct_cells": {"numerator": sum(c["evidence_correct"] for c in rows
                                                               if c["expected_business"] in DETERMINISTIC),
                                               "denominator": sum(c["expected_business"] in DETERMINISTIC for c in rows)},
                    "anchor_group_recall": {"numerator": sum(c["evidence_group_hits"] for c in rows),
                                            "denominator": sum(c["evidence_group_total"] for c in rows)}}
            for label, rows in groups.items()}


def aggregate_scores(scores: list[dict]) -> dict:
    """Aggregate only one data class at a time; absent attempts must be supplied as failed scores."""
    if len({score["data_class"] for score in scores}) > 1:
        raise ValueError("mixed_data_classes")
    if len({score["pack_id"] for score in scores}) != len(scores):
        raise ValueError("duplicate_pack")
    result = {"pack_count": len(scores), "data_class": scores[0]["data_class"] if scores else None,
              "agent_tasks": {"numerator": sum(score["task_success"] for score in scores), "denominator": len(scores)},
              "dangerous_meets": sum(score["dangerous_meets"] for score in scores),
              "formal_independent_acceptance": False}
    for metric in ("status_accuracy", "necessary_evidence", "grounded_correct_cells", "anchor_group_recall",
                   "absence_coverage", "citation_errors"):
        result[metric] = {name: sum(score[metric][name] for score in scores) for name in ("numerator", "denominator")}
    total = result["citation_errors"]["denominator"]
    result["citation_error_rate"] = result["citation_errors"]["numerator"] / total if total else None
    cells = [cell for score in scores for cell in score["cells"]]
    result["by_category"] = _breakdown(cells, "category")
    result["by_rule"] = _breakdown(cells, "rule_type")
    return result
