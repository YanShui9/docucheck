"""Create and verify development originals; handwritten labels stay outside inputs.

No model, environment file, database or production decision engine is used here.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from .contracts import Anchor, Gold, Manifest, validate_pair
from .metrics import read_originals

HERE = Path(__file__).parent
IDENTIFIER = re.compile(r"^[a-z0-9][a-z0-9_-]{0,63}$")
NOTICE = "合成开发评估资料：内容均为虚构，不代表真实产品、证书或供应商。"
SHEET = "采购资料"
FONT_NAME = "DynamicEvalChinese"
MAX_JSON_BYTES = 4 * 1024 * 1024
_FONT_SOURCE: tuple[str, int, int] | None = None


def _json_bytes(value: dict) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def _stable_xlsx(content: bytes) -> bytes:
    """Remove save-time metadata from our own freshly generated workbooks."""
    result = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(result, "w", compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>).*?(</dcterms:(?:created|modified)>)",
                                 rb"\g<1>2026-10-11T00:00:00Z\g<2>", payload)
            entry = ZipInfo(name, date_time=(2026, 10, 11, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            target.writestr(entry, payload)
    return result.getvalue()


def _read_json(path: Path) -> dict:
    if path.is_symlink() or not path.is_file() or path.stat().st_size > MAX_JSON_BYTES:
        raise ValueError("invalid_json_asset")
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise TypeError("invalid_json_asset")
    return value


def _safe_path(root: Path, *parts: str) -> Path:
    path = root
    for part in parts:
        path = path / part
        if path.is_symlink() or path.is_junction():
            raise ValueError("unsafe_asset_path")
    if not path.resolve().is_relative_to(root.resolve()):
        raise ValueError("unsafe_asset_path")
    return path


def _normal(text: str) -> str:
    return "".join(unicodedata.normalize("NFKC", text).split())


def _buyer_text(criterion: dict) -> str:
    name, rule = criterion["name"], criterion["rule"]
    operation = {"eq": "为", "gte": "至少", "lte": "不超过", "in": "属于"}
    kind = rule["type"]
    if kind == "numeric":
        return f"{name}{operation[rule['operator']]} {rule['value']} {rule['unit']}"
    if kind == "money":
        tax = {"included": "含税", "excluded": "不含税", "unspecified": "税费口径未指定"}[rule["tax_basis"]]
        return f"{name}（{tax}）{operation[rule['operator']]} {rule['value']} {rule['currency']}/{rule['per_unit']}"
    if kind == "duration":
        unit = {"natural_days": "自然日", "working_days": "工作日", "months": "个月"}[rule["calendar"]]
        return f"{name}：{rule['start_event']}后{operation[rule['operator']]} {rule['value']} {unit}"
    if kind == "enum":
        return f"{name}{operation[rule['operator']]} {'、'.join(rule['values'])}"
    if kind == "presence":
        return f"{name}：必须提供{rule['description']}"
    return f"{name}：{rule['instruction']}"


def _identity(item: dict) -> str:
    return f"品目：{item['name']}；型号：{item['specification'].removeprefix('型号 ')}"


def _pdf(pages: list[list[str]]) -> bytes:
    buffer = BytesIO()
    document = canvas.Canvas(buffer, pagesize=A4, invariant=1)
    document.setTitle("DocuCheck 合成开发评估资料")
    for lines in pages:
        document.setFont(FONT_NAME, 9)
        document.drawString(36, 807, NOTICE)
        document.setFont(FONT_NAME, 11)
        y = 773
        for line in lines:
            # Wrap by measured width, keeping every source character available.
            current = ""
            for character in line:
                if pdfmetrics.stringWidth(current + character, FONT_NAME, 11) > 515:
                    document.drawString(36, y, current)
                    y -= 20
                    current = ""
                current += character
            if current:
                document.drawString(36, y, current)
                y -= 26
            if y < 60:
                raise ValueError("pdf_page_overflow")
        document.showPage()
    document.save()
    return buffer.getvalue()


def _anchor(document: str, role: str, supplier: str | None, item: str, text: str,
            *, page: int | None = None, cell: str | None = None) -> dict:
    return Anchor(document=document, role=role, supplier_handle=supplier, item_id=item,
                  text=text, page=page, sheet=SHEET if cell else None, cell=cell).model_dump(mode="json")


def _build_pack(blueprint: dict, examples: dict) -> tuple[Manifest, Gold, dict[str, bytes]]:
    definition, suppliers = blueprint["checklist_definition"], blueprint["supplier_handles"]
    items, criteria = definition["items"], definition["criteria"]
    originals, documents, bindings, cells = {}, [], [], []
    buyer_anchors, claims, material_bodies = {}, {}, {}
    pages = []
    for page, item in enumerate(items, 1):
        lines = [_identity(item), "买方确认的采购要求"]
        for criterion in criteria:
            if criterion["item_id"] == item["item_id"]:
                text = _buyer_text(criterion)
                lines.append(text)
                buyer_anchors[criterion["criterion_id"]] = _anchor(
                    "requirements.pdf", "requirement", None, item["item_id"], text, page=page)
        pages.append(lines)
    originals["requirements.pdf"] = _pdf(pages)
    documents.append({"filename": "requirements.pdf", "role": "requirement", "supplier_handle": None,
                      "pages": len(items)})
    for supplier, examples_for_supplier in zip(suppliers, examples["suppliers"], strict=True):
        if len(examples_for_supplier) != len(criteria):
            raise ValueError("example_criterion_count")
        book = Workbook()
        sheet = book.active
        sheet.title = SHEET
        sheet.append([NOTICE])
        sheet.append([f"供应商：{supplier}"])
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 100
        spec_pages = []
        quote_filename, spec_filename = f"{supplier}-quote.xlsx", f"{supplier}-spec.pdf"
        for page, item in enumerate(items, 1):
            item_id, start = item["item_id"], sheet.max_row + 1
            identity = _identity(item)
            sheet.append([identity])
            spec_lines = [identity, f"供应商：{supplier}", "供应商响应"]
            for criterion, example in zip(criteria, examples_for_supplier, strict=True):
                if criterion["item_id"] != item_id:
                    continue
                quote, spec = example["quote"], example["spec"] or example["quote"]
                group = []
                if quote:
                    sheet.append([criterion["name"], quote])
                    group.append(_anchor(quote_filename, "quote", supplier, item_id, quote,
                                         cell=f"B{sheet.max_row}"))
                if spec:
                    spec_lines.append(spec)
                    group.append(_anchor(spec_filename, "spec", supplier, item_id, spec, page=page))
                    if criterion["rule"]["type"] == "presence" and example["status"] == "满足":
                        material_text = f"附件材料正文 — {criterion['name']}：适用型号与本页一致；合成记录编号 {blueprint['pack_id']}。"
                        spec_lines.append(material_text)
                        material_bodies[supplier, criterion["criterion_id"]] = _anchor(
                            spec_filename, "spec", supplier, item_id, material_text, page=page)
                claims[supplier, criterion["criterion_id"]] = group
            # Empty response rows are represented by omission, not by invented denials.
            if sheet.max_row == start:
                sheet.append(["已列出采购品目；未写明其他响应条款。"])
            end = sheet.max_row
            quote_anchor = _anchor(quote_filename, "quote", supplier, item_id, identity, cell=f"A{start}")
            spec_anchor = _anchor(spec_filename, "spec", supplier, item_id, identity, page=page)
            sku = item["specification"].removeprefix("型号 ")
            bindings.extend([
                {"document": quote_filename, "item_id": item_id, "sku": sku,
                 # The source inventory includes the synthetic notice and
                 # supplier header. Assign them once, with the first item;
                 # later item ranges remain disjoint and keep their anchors.
                 "scope": {"sheet": SHEET, "cell_range": f"A{1 if page == 1 else start}:B{end}"},
                 "anchor": quote_anchor, "sku_anchor": quote_anchor},
                {"document": spec_filename, "item_id": item_id, "sku": sku,
                 "scope": {"page_start": page, "page_end": page},
                 "anchor": spec_anchor, "sku_anchor": spec_anchor},
            ])
            spec_pages.append(spec_lines)
        if blueprint["pack_id"] == "dyn-d03":
            spec_pages.append(["品目：其他手机；型号：DEMO-OTHER", "此页产品不在当前采购范围。",
                               "存储容量 512 GB", "含税单价 6800 CNY/台"])
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(name="Microsoft YaHei", size=11)
            sheet.row_dimensions[row[0].row].height = 40
        sheet.freeze_panes = "B4"
        sheet["A1"].fill = PatternFill("solid", fgColor="FFF1CC")
        buffer = BytesIO()
        book.save(buffer)
        book.close()
        originals[quote_filename], originals[spec_filename] = _stable_xlsx(buffer.getvalue()), _pdf(spec_pages)
        documents.extend([
            {"filename": quote_filename, "role": "quote", "supplier_handle": supplier, "pages": 0},
            {"filename": spec_filename, "role": "spec", "supplier_handle": supplier, "pages": len(spec_pages)},
        ])
    for document in documents:
        content = originals[document["filename"]]
        document.update(sha256=sha256(content).hexdigest(), bytes=len(content))
    manifest = Manifest.model_validate({
        "schema_version": 1, "protocol_id": "dynamic-dev-v1", "pack_id": blueprint["pack_id"],
        "data_class": "ai_authored_seen_development", "category": blueprint["category"],
        "source_family_ids": blueprint["source_family_ids"], "supplier_handles": suppliers,
        "documents": documents, "checklist_definition": definition, "bindings": bindings,
        "buyer_anchors": buyer_anchors,
        "authors": ["codex_ai"], "reviewers": ["codex_ai_source_review_not_independent"],
    })
    for supplier, examples_for_supplier in zip(suppliers, examples["suppliers"], strict=True):
        for criterion, example in zip(criteria, examples_for_supplier, strict=True):
            criterion_id, item_id = criterion["criterion_id"], criterion["item_id"]
            buyer = buyer_anchors[criterion_id]
            groups = [{"group_id": "buyer", "purpose": "buyer_baseline", "alternatives": [buyer]}]
            anchors = claims[supplier, criterion_id]
            if example["status"] == "资料冲突":
                groups.extend({"group_id": f"claim-{i}", "purpose": "supplier_claim", "alternatives": [a]}
                              for i, a in enumerate(anchors, 1))
            elif anchors:
                groups.append({"group_id": "supplier", "purpose": "supplier_claim", "alternatives": anchors})
            material = material_bodies.get((supplier, criterion_id))
            if material is not None:
                groups.append({"group_id": "supplier-material", "purpose": "supplier_claim", "alternatives": [material]})
            applicable = [{"document": b.document, "item_id": b.item_id, "scope": b.scope.model_dump(mode="json")}
                          for b in manifest.bindings if b.item_id == item_id and b.anchor.supplier_handle == supplier]
            source_clauses = "；".join(a["text"] for a in anchors) or "适用范围的两份供应商原件均未写明该指标"
            cells.append({
                "pack_id": manifest.pack_id, "supplier_handle": supplier, "item_id": item_id,
                "criterion_id": criterion_id, "expected_processing": "complete",
                "expected_business": example["status"], "expected_formal": True, "expected_needs_review": False,
                "reason": f"买方：{buyer['text']}。供应商：{source_clauses}。预定：{example['status']}。{example['reason']}",
                "evidence_groups": groups, "required_coverage": applicable,
                "rule_basis": json.dumps(criterion["rule"], ensure_ascii=False, sort_keys=True),
                "buyer_baseline": buyer["text"], "source_version": "sha256_manifest_originals",
            })
    gold = Gold.model_validate({"schema_version": 1, "protocol_id": "dynamic-dev-v1", "pack_id": manifest.pack_id,
                               "data_class": "ai_authored_seen_development", "author": "codex_ai",
                               "reviewer": "codex_ai_source_review_not_independent",
                               "review_state": "ai_reviewed_development", "cells": cells})
    validate_pair(manifest, gold)
    return manifest, gold, originals


def build_corpus(root: Path, *, font_path: Path, pack_ids: list[str] | None = None) -> dict:
    global _FONT_SOURCE  # Reuse the public font registry to bound repeated test builds.
    root, font_path = Path(root), Path(font_path)
    if root.exists() or root.is_symlink():
        raise ValueError("destination_exists")
    blueprint = _read_json(HERE / "case_matrix.json")
    examples = _read_json(HERE / "development_examples.json")
    by_id = {pack["pack_id"]: pack for pack in blueprint["packs"]}
    selected = list(by_id) if pack_ids is None else pack_ids
    if not selected or len(set(selected)) != len(selected) or any(p not in by_id for p in selected):
        raise ValueError("pack_selection")
    if not font_path.is_file():
        raise ValueError("font_unavailable")
    stat = font_path.stat()
    font_source = str(font_path.resolve()), stat.st_mtime_ns, stat.st_size
    if _FONT_SOURCE != font_source or FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(FONT_NAME, str(font_path)))
        _FONT_SOURCE = font_source
    examples_by_id = {pack["pack_id"]: pack for pack in examples["packs"]}
    generated = [_build_pack(by_id[p], examples_by_id[p]) for p in selected]
    # Validate everything before creating the destination.
    root.mkdir(parents=True, exist_ok=False)
    index = {"schema_version": 1, "protocol_id": "dynamic-dev-v1", "data_class": "ai_authored_seen_development",
             "independent_sealed_packs": 0, "blueprint_sha256": sha256((HERE / "case_matrix.json").read_bytes()).hexdigest(),
             "examples_sha256": sha256((HERE / "development_examples.json").read_bytes()).hexdigest(), "packs": []}
    for manifest, gold, originals in generated:
        directory = root / "inputs" / manifest.pack_id
        directory.mkdir(parents=True)
        for name, content in originals.items():
            (directory / name).write_bytes(content)
        entry = {"pack_id": manifest.pack_id, "primary_cells": len(gold.cells)}
        for kind, value in [("config", manifest), ("gold", gold)]:
            parent = root / kind
            parent.mkdir(exist_ok=True)
            data = _json_bytes(value.model_dump(mode="json"))
            (parent / f"{manifest.pack_id}.json").write_bytes(data)
            entry["manifest_sha256" if kind == "config" else "gold_sha256"] = sha256(data).hexdigest()
        index["packs"].append(entry)
    (root / "index.json").write_bytes(_json_bytes(index))
    return index


def load_pack(root: Path, pack_id: str) -> tuple[Manifest, Gold, dict[str, bytes]]:
    root = Path(root).absolute()
    if root.is_symlink() or root.is_junction():
        raise ValueError("unsafe_asset_path")
    if not IDENTIFIER.fullmatch(pack_id):
        raise ValueError("pack_selection")
    index = _read_json(_safe_path(root, "index.json"))
    rows = [p for p in index["packs"] if p["pack_id"] == pack_id]
    if len(rows) != 1:
        raise ValueError("pack_selection")
    row = rows[0]
    config_path, gold_path = (_safe_path(root, kind, f"{pack_id}.json") for kind in ["config", "gold"])
    config, gold_value = _read_json(config_path), _read_json(gold_path)
    if (sha256(config_path.read_bytes()).hexdigest() != row["manifest_sha256"]
            or sha256(gold_path.read_bytes()).hexdigest() != row["gold_sha256"]):
        raise ValueError("config_or_gold_integrity")
    manifest, gold = Manifest.model_validate(config), Gold.model_validate(gold_value)
    if manifest.pack_id != pack_id or len(gold.cells) != row["primary_cells"]:
        raise ValueError("config_or_gold_integrity")
    validate_pair(manifest, gold)
    directory = _safe_path(root, "inputs", pack_id)
    if not directory.is_dir() or {p.name for p in directory.iterdir()} != {d.filename for d in manifest.documents}:
        raise ValueError("input_inventory")
    originals = {}
    for document in manifest.documents:
        path = _safe_path(root, "inputs", pack_id, document.filename)
        if not path.is_file() or path.stat().st_size != document.bytes:
            raise ValueError("input_integrity")
        content = path.read_bytes()
        if sha256(content).hexdigest() != document.sha256:
            raise ValueError("input_integrity")
        originals[document.filename] = content
    return manifest, gold, originals


def validate_corpus(root: Path) -> dict:
    index = _read_json(Path(root) / "index.json")
    errors, cells, documents = [], 0, 0
    for entry in index["packs"]:
        try:
            manifest, gold, originals = load_pack(root, entry["pack_id"])
            texts, original_errors = read_originals(manifest, originals)
            errors.extend({"pack_id": manifest.pack_id, **error} for error in original_errors)
            anchors = [a for cell in gold.cells for group in cell.evidence_groups for a in group.alternatives]
            anchors.extend(binding.anchor for binding in manifest.bindings)
            for anchor in anchors:
                original = texts.get((anchor.document, anchor.page, anchor.sheet, anchor.cell), "")
                if _normal(anchor.text) not in _normal(original):
                    errors.append({"pack_id": manifest.pack_id, "document": anchor.document,
                                   "code": "anchor_text_missing", "text": anchor.text})
            cells += len(gold.cells)
            documents += len(manifest.documents)
        except (ValueError, TypeError, KeyError, OSError) as error:
            errors.append({"pack_id": entry["pack_id"], "code": str(error)[:200]})
    return {"packs": len(index["packs"]), "primary_cells": cells, "documents": documents,
            "data_class": "ai_authored_seen_development", "independent_sealed_packs": 0, "errors": errors}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--font", type=Path, default=Path("C:/Windows/Fonts/simhei.ttf"))
    parser.add_argument("--pack", action="append")
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    if not args.validate_only:
        build_corpus(args.root, font_path=args.font, pack_ids=args.pack)
    result = validate_corpus(args.root)
    print(json.dumps(result, ensure_ascii=False))
    return 1 if result["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
