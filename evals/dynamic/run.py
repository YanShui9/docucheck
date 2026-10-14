"""Preview frozen dynamic development inputs, or explicitly send to loopback HTTP.

No credentials, .env, database or service startup is used during preview. The
target receives only original business files and the predetermined checklist
and bindings. Gold remains on the local scoring side.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import time
import unicodedata
from datetime import UTC, datetime
from hashlib import sha256
from os import fsync
from os import replace as atomic_replace
from pathlib import Path
from urllib.parse import urlsplit
from uuid import uuid4
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import httpx
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

TERMINAL = {"completed", "partial", "failed", "cancelled", "interrupted", "waiting_user"}
BUSINESS = {"满足", "不满足", "资料冲突", "证据不足", "待人工确认"}
QUESTION = "请完整审查所有供应商和全部已确认指标，并列出结果与来源"
MAX_JSON_BYTES = 8 * 1024 * 1024
MAX_EXPORT_BYTES = 20 * 1024 * 1024
REPO = Path(__file__).resolve().parents[2]


class EvaluationError(ValueError):
    """A bounded error code; transport exception text never enters output."""


def validate_api_url(value: str) -> str:
    try:
        url = urlsplit(value)
        if (url.scheme not in {"http", "https"} or url.hostname not in {"127.0.0.1", "::1"}
                or url.username or url.password or url.query or url.fragment
                or url.path not in {"", "/"}
                or not 1 <= (url.port if url.port is not None else 80) <= 65535):
            raise ValueError
    except ValueError:
        raise ValueError("evaluation_api_must_be_literal_loopback") from None
    return value.rstrip("/")


def _identifier(value) -> str:
    if not isinstance(value, str) or not re.fullmatch(r"[a-f0-9]{32}|[a-f0-9]{64}", value):
        raise EvaluationError("api_identifier_invalid")
    return value


def _encoded(value) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode("utf-8")


def _write_new(path: Path, value) -> None:
    with path.open("xb") as stream:
        stream.write(value if isinstance(value, bytes) else _encoded(value))
        stream.flush()
        fsync(stream.fileno())


def _unresolved_mutation(value: dict) -> bool:
    return any(row.get("method") in {"POST", "PUT"} and (
        row.get("state") == "sending" or row.get("state") == "received"
        and 200 <= row.get("http_status", 0) < 300 and not row.get("receipt_validated")
    ) for row in value.get("raw", []))


class AttemptJournal:
    """Own one reserved in-flight record; final reports remain append-only.

    Sync a new sibling before atomic replacement so process termination cannot
    truncate the last known request/receipt. This is not a power-loss guarantee
    or permission to resume/retry an unknown request.
    """

    def __init__(self, path: Path, initial: dict):
        self.path = path
        _write_new(path, self._snapshot(initial, state="reserved"))

    @staticmethod
    def _snapshot(value: dict, *, state: str) -> dict:
        result = {key: data for key, data in value.items() if key != "xlsx_bytes"}
        result["outcome_unknown"] = bool(result.get("outcome_unknown") or _unresolved_mutation(result))
        result["journal_state"] = state
        result["needs_reconciliation"] = result["outcome_unknown"] or state == "in_flight"
        result["recovery_policy"] = "inspect_receipts_and_server_state_never_automatically_resend"
        return result

    def update(self, value: dict, *, finished: bool = False) -> None:
        temporary = self.path.with_name(f".{self.path.name}.{uuid4().hex}.tmp")
        _write_new(temporary, self._snapshot(value, state="finished" if finished else "in_flight"))
        # Both names belong to this invocation's fresh output directory; the
        # immutable final attempt file is never passed to this replacement.
        atomic_replace(temporary, self.path)


def _unstarted(task: dict, error_code: str | None = None) -> dict:
    return {"pack_id": task["pack_id"], "attempt_number": 1, "stage": "not_started",
            "error_code": error_code, "outcome_unknown": False, "automatic_retries": False,
            "raw": [], "usage": {"state": "unknown", "actual": None, "cost": None},
            "observation": {"pack_id": task["pack_id"], "run_status": "not_started", "error_code": error_code}}


def _source_identity() -> dict:
    paths = [p for folder in ("backend", "evals") for p in (REPO / folder).rglob("*.py")
             if not p.is_symlink() and not any(part.startswith(".env") for part in p.parts)]
    paths += [REPO / name for name in ("pyproject.toml", "backend/requirements.lock.txt",
                                      "backend/requirements-dev.lock.txt", "backend/environment.yml")
              if (REPO / name).is_file()]
    rows = [{"path": p.relative_to(REPO).as_posix(), "sha256": sha256(p.read_bytes()).hexdigest()}
            for p in sorted(paths)]
    return {"kind": "python_sources_and_locked_dependencies", "files": rows,
            "sha256": sha256(_encoded(rows)).hexdigest()}


def prepare_protocol(root: Path, *, packs: list[str] | None = None) -> dict:
    from backend.app.agent.policy import MODE_LIMITS

    from .corpus import load_pack, validate_corpus

    validation = validate_corpus(root)
    if validation.get("errors"):
        raise ValueError("evaluation_corpus_invalid")
    index_path = root / "index.json"
    index = json.loads(index_path.read_bytes())
    available = {row["pack_id"]: row for row in index["packs"]}
    selected = packs or list(available)
    if not selected or len(selected) != len(set(selected)) or set(selected) - set(available):
        raise ValueError("evaluation_selection_invalid")
    tasks = []
    for pack_id in selected:
        manifest, _gold, _inputs = load_pack(root, pack_id)
        business = manifest.model_dump(mode="json")
        config = (root / "config" / f"{pack_id}.json").read_bytes()
        gold = (root / "gold" / f"{pack_id}.json").read_bytes()
        tasks.append({"task_id": f"{pack_id}-01", "pack_id": pack_id, "attempt_number": 1,
                      "manifest": business, "manifest_sha256": sha256(config).hexdigest(),
                      "gold_sha256": sha256(gold).hexdigest(),
                      "primary_cells": available[pack_id]["primary_cells"]})
    limits = MODE_LIMITS["thorough"]
    return {"schema_version": 1, "protocol_id": index["protocol_id"],
            "created_at": datetime.now(UTC).isoformat(), "track": "formal_review",
            "data_class": index["data_class"], "independent_sealed_packs": 0,
            "formal_independent_acceptance": False, "index_sha256": sha256(index_path.read_bytes()).hexdigest(),
            "source_identity": _source_identity(), "tasks": tasks,
            "planned_tasks": len(tasks), "planned_files": sum(len(t["manifest"]["documents"]) for t in tasks),
            "planned_positions": sum(t["primary_cells"] for t in tasks),
            "mode": "thorough", "maximum_agent_requests": len(tasks) * limits["max_calls"],
            "maximum_input_tokens": len(tasks) * limits["max_input"],
            "maximum_output_tokens": len(tasks) * limits["max_output"],
            "automatic_retries": False, "new_project": True, "new_case_per_attempt": True,
            "candidate_generation_measured": False,
            "environment": {key: "unknown" for key in (
                "api_source", "worker_source", "migration", "model", "parser", "embedding", "actual_bill")}}


def _input_bytes(pack_dir: Path, manifest: dict) -> dict[str, bytes]:
    directory = pack_dir.resolve(strict=True)
    if pack_dir.is_symlink():
        raise EvaluationError("evaluation_input_changed")
    names = {row["filename"] for row in manifest["documents"]}
    if len(names) != len(manifest["documents"]) or {p.name for p in directory.iterdir()} != names:
        raise EvaluationError("evaluation_input_changed")
    result = {}
    for doc in manifest["documents"]:
        name = doc["filename"]
        path = directory / name
        if (Path(name).name != name or path.is_symlink() or not path.is_file()
                or path.resolve().parent != directory or path.stat().st_size != doc["bytes"]
                or path.stat().st_size > MAX_EXPORT_BYTES):
            raise EvaluationError("evaluation_input_changed")
        content = path.read_bytes()
        if sha256(content).hexdigest() != doc["sha256"]:
            raise EvaluationError("evaluation_input_changed")
        result[name] = content
    return result


def _anchor_context(anchor: dict, evidence: list[dict]) -> dict:
    def compact(value):
        return re.sub(r"\s+", " ", value).strip()

    def normalized(value):
        return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value))

    matches = [row for row in evidence
               if row.get("sheet") == anchor.get("sheet") and row.get("cell") == anchor.get("cell")
               and (None if row.get("page_idx") is None else row["page_idx"] + 1) == anchor.get("page")
               and normalized(anchor["text"]) in normalized(row["source_text"])
               and (anchor.get("bbox") is None or row.get("bbox") == anchor["bbox"])]
    if len(matches) != 1:
        raise EvaluationError("binding_anchor_unresolved")
    block = matches[0]
    # Parsing may normalize punctuation or wrap inside a Chinese phrase. Use
    # the actual single source block whenever strict whitespace matching fails.
    quote = anchor["text"] if compact(anchor["text"]) in compact(block["source_text"]) else block["source_text"]
    return {"evidence_refs": [_identifier(block["id"])], "quote": quote}


def _selector(scope: dict) -> dict:
    if scope.get("sheet") is not None:
        min_col, min_row, max_col, max_row = range_boundaries(scope["cell_range"])
        return {"kind": "xlsx", "sheet": scope["sheet"], "start_row": min_row,
                "end_row": max_row, "start_column": min_col, "end_column": max_col}
    if scope.get("page_start") is not None:
        return {"kind": "pdf", "start_page": scope["page_start"], "end_page": scope["page_end"]}
    return {"kind": "whole"}


def _in_scope(row: dict, scope: dict) -> bool:
    if scope.get("sheet") is not None:
        if row.get("sheet") != scope["sheet"] or not row.get("cell"):
            return False
        left, top, right, bottom = range_boundaries(scope["cell_range"])
        c1, r1, c2, r2 = range_boundaries(row["cell"])
        return left <= c1 <= c2 <= right and top <= r1 <= r2 <= bottom
    if scope.get("page_start") is not None:
        return row.get("page_idx") is not None and scope["page_start"] <= row["page_idx"] + 1 <= scope["page_end"]
    return True


def _cells(pack_id: str, checks: list[dict], *, needs_review: bool) -> list[dict]:
    return [{"pack_id": pack_id, "supplier_handle": row["supplier"], "item_id": row["item_id"],
             "criterion_id": row["criterion_id"], "processing_status": row["processing_status"],
             "business_status": row.get("business_status"), "is_formal": row["is_formal"],
             "needs_review": needs_review,
             "evidence_refs": list(dict.fromkeys([*row["evidence_refs"], *row["reason_evidence_refs"]])),
             "surface": "full"} for row in checks]


def _visible_cells(manifest: dict, answer: str | None) -> tuple[list[dict], list[str]]:
    items = {row["item_id"]: row["name"] for row in manifest["checklist_definition"]["items"]}
    identities = {}
    for supplier in manifest["supplier_handles"]:
        for criterion in manifest["checklist_definition"]["criteria"]:
            if criterion.get("enabled", True):
                label = f"{supplier} · {items[criterion['item_id']]} · {criterion['name']}"
                identities.setdefault(label, []).append((supplier, criterion["item_id"], criterion["criterion_id"]))
    result, errors = [], []
    for line in (answer or "").splitlines():
        if " · " not in line or "：" not in line:
            continue
        label, status = line.split("：", 1)
        status = status.removesuffix("。").strip()
        matches = identities.get(label, [])
        if len(matches) != 1 or status not in BUSINESS | {"处理未完成", "待人工确认（资料处理未完成）"}:
            errors.append("visible_answer_identity_or_status_unresolved")
            continue
        supplier, item, criterion = matches[0]
        incomplete = status in {"处理未完成", "待人工确认（资料处理未完成）"}
        result.append({"pack_id": manifest["pack_id"], "supplier_handle": supplier,
                       "item_id": item, "criterion_id": criterion,
                       "processing_status": "partial" if incomplete else "complete",
                       "business_status": "待人工确认" if status.startswith("待人工确认") else (None if incomplete else status),
                       "is_formal": None, "needs_review": None, "evidence_refs": [], "surface": "status_only"})
    return result, errors


def _xlsx_cells(pack_id: str, content: bytes, *, needs_review: bool) -> list[dict]:
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        sheet = workbook["核验结果"]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise EvaluationError("xlsx_export_shape_invalid")
        headers = {value: index for index, value in enumerate(rows[0])}
        required = {"供应商", "品目 ID", "指标 ID", "处理状态", "业务结果", "正式", "来源 ID", "原因来源 ID"}
        if not required <= set(headers):
            raise EvaluationError("xlsx_export_shape_invalid")
        result = []
        for row in rows[1:]:
            if not any(value is not None for value in row):
                continue
            def cell(name, values=row):
                return values[headers[name]]
            refs = []
            for name in ("来源 ID", "原因来源 ID"):
                source_ids = cell(name)
                if source_ids is not None and not isinstance(source_ids, str):
                    raise EvaluationError("xlsx_export_shape_invalid")
                refs.extend(ref for ref in (source_ids or "").splitlines() if ref)
            if cell("正式") not in {"是", "否"}:
                raise EvaluationError("xlsx_export_shape_invalid")
            result.append({"pack_id": pack_id, "supplier_handle": cell("供应商"), "item_id": cell("品目 ID"),
                           "criterion_id": cell("指标 ID"), "processing_status": cell("处理状态"),
                           "business_status": cell("业务结果"), "is_formal": cell("正式") == "是",
                           "needs_review": needs_review, "evidence_refs": list(dict.fromkeys(refs)), "surface": "full"})
        return result
    except (BadZipFile, InvalidFileException, ParseError, DefusedXmlException):
        raise EvaluationError("xlsx_export_shape_invalid") from None
    finally:
        if workbook is not None:
            workbook.close()


def run_pack(client, pack_dir: Path, *, manifest, project_id: str, deadline_seconds: float = 240,
             clock=time.monotonic, sleep=time.sleep, checkpoint=None) -> dict:
    """Only original inputs plus business configuration cross this target boundary."""
    from .contracts import Manifest

    business = (manifest.model_dump(mode="json") if hasattr(manifest, "model_dump") else manifest)
    pack_id = business.get("pack_id", "invalid-manifest") if isinstance(business, dict) else "invalid-manifest"
    observation = {"pack_id": pack_id, "run_status": "not_started", "api_cells": [], "visible_cells": [],
                   "export_cells": [], "xlsx_export_cells": [], "sources": [], "coverage": [],
                   "coverage_unknown": True, "surface_errors": []}
    attempt = {"pack_id": pack_id, "attempt_number": 1, "stage": "input_validation", "error_code": None,
               "outcome_unknown": False, "automatic_retries": False, "raw": [], "observation": observation,
               "usage": {"state": "unknown", "actual": None, "cost": None}, "xlsx_bytes": None}
    started, documents, evidence = clock(), {}, {}
    deadline = started + deadline_seconds

    def save():
        if checkpoint is not None:
            checkpoint(attempt)

    def acknowledge_receipt():
        # A successful HTTP status is insufficient: the caller must validate
        # the returned resource identity/version before the mutation is known.
        attempt["raw"][-1]["receipt_validated"] = True
        save()

    def request(method, path, *, expected=200, binary=False, **kwargs):
        remaining = deadline - clock()
        if remaining <= 0:
            raise EvaluationError(f"{attempt['stage']}_timeout")
        record = {"method": method, "path": path, "stage": attempt["stage"], "state": "sending"}
        for key in ("json", "data", "params"):
            if key in kwargs:
                record[key] = kwargs[key]
        if "headers" in kwargs:
            record["idempotency_key"] = kwargs["headers"].get("Idempotency-Key")
        attempt["raw"].append(record)
        save()
        try:
            response = client.request(method, path, timeout=min(15, remaining), **kwargs)
        except httpx.HTTPError:
            record["error_code"] = "http_transport_error"
            record["state"] = "failed"
            if method in {"POST", "PUT"}:
                attempt["outcome_unknown"] = True
            save()
            raise EvaluationError("http_transport_error") from None
        record["http_status"] = response.status_code
        record["state"] = "received"
        if response.status_code != expected:
            if method in {"POST", "PUT"} and response.status_code >= 500:
                attempt["outcome_unknown"] = True
            try:
                error = response.json().get("error", {}).get("code")
                if isinstance(error, str) and re.fullmatch(r"[a-z][a-z0-9_]{0,99}", error):
                    record["api_error_code"] = error
            except (ValueError, AttributeError):
                pass
            save()
            raise EvaluationError(f"http_status_{response.status_code}")
        if binary:
            if len(response.content) > MAX_EXPORT_BYTES:
                raise EvaluationError("api_response_limit")
            record.update(bytes=len(response.content), sha256=sha256(response.content).hexdigest())
            save()
            return response.content
        if len(response.content) > MAX_JSON_BYTES:
            raise EvaluationError("api_response_limit")
        try:
            value = response.json()
        except ValueError:
            raise EvaluationError("api_json_invalid") from None
        if not isinstance(value, dict):
            raise EvaluationError("api_json_invalid")
        record["response"] = value
        save()
        return value

    def edit(method, path, body, *, expected=200):
        return request(method, path, expected=expected, json=body, headers={"Idempotency-Key": uuid4().hex})

    def bound(value, *, run=False):
        if (value["case_id"] != attempt["case_id"] or value["snapshot_id"] != attempt["snapshot_id"]
                or value["revision"] != attempt["revision"]):
            raise EvaluationError("run_receipt_mismatch" if run else "snapshot_document_binding")

    try:
        try:
            business = Manifest.model_validate(business).model_dump(mode="json")
        except ValueError:
            raise EvaluationError("evaluation_manifest_invalid") from None
        if not 0 < deadline_seconds <= 600:
            raise EvaluationError("evaluation_deadline_invalid")
        originals = _input_bytes(pack_dir, business)
        project_id = _identifier(project_id)
        attempt["stage"] = "create_case"
        case = request("POST", "/api/v1/cases", expected=201, json={"project_id": project_id})
        attempt["case_id"] = _identifier(case["id"])
        if case["revision"] != 0:
            raise EvaluationError("evaluation_requires_empty_case")
        acknowledge_receipt()
        revision = 0
        attempt["stage"] = "upload"
        for doc in business["documents"]:
            name = doc["filename"]
            skus = {b["sku"] for b in business["bindings"] if b["document"] == name and b.get("sku")}
            data = {"role": doc["role"], "expected_revision": str(revision)}
            if doc.get("supplier_handle") is not None:
                data.update(supplier_handle=doc["supplier_handle"], supplier_name=doc["supplier_handle"])
                if len(skus) == 1:
                    data["model_name"] = next(iter(skus))
            mime = ("application/pdf" if name.endswith(".pdf") else
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            receipt = request("POST", f"/api/v1/cases/{attempt['case_id']}/documents", expected=201,
                              data=data, files={"file": (name, originals[name], mime)},
                              headers={"Idempotency-Key": uuid4().hex})
            if (receipt.get("case_id") != attempt["case_id"] or receipt["revision"] != revision + 1
                    or receipt["source_sha256"] != doc["sha256"]):
                raise EvaluationError("upload_receipt_mismatch")
            documents[name] = {**doc, "id": _identifier(receipt["document_id"]), "model_name": data.get("model_name")}
            revision = receipt["revision"]
            attempt["snapshot_id"] = _identifier(receipt["snapshot_id"])
            acknowledge_receipt()
        attempt["revision"] = revision
        prefix = f"/api/v1/cases/{attempt['case_id']}/snapshots/{attempt['snapshot_id']}"
        attempt["stage"] = "parse"
        while True:
            listed = request("GET", f"{prefix}/documents")
            bound(listed)
            rows = listed["documents"]
            expected = {(name, doc["sha256"], doc["role"], doc.get("supplier_handle"), doc.get("model_name"))
                        for name, doc in documents.items()}
            actual = {(r["original_name"], r["source_sha256"], r["role"], r.get("supplier_handle"), r.get("model_name")) for r in rows}
            if expected != actual or len(rows) != len(expected):
                raise EvaluationError("snapshot_document_binding")
            # Uploading creates a new snapshot and clones retained documents.
            # Runtime evidence and bindings must use the final snapshot's IDs.
            for row in rows:
                documents[row["original_name"]]["id"] = _identifier(row["id"])
            attempt["document_ids"] = {name: document["id"] for name, document in documents.items()}
            states = {row["parse_status"] for row in rows}
            if "failed" in states:
                raise EvaluationError("parser_failed")
            if states <= {"complete", "partial"}:
                break
            sleep(min(0.5, max(0, deadline - clock())))
        attempt["stage"] = "evidence"
        for name, document in documents.items():
            after, blocks = -1, []
            while True:
                page = request("GET", f"{prefix}/documents/{document['id']}/evidence",
                               params={"after_ordinal": after, "limit": 100})
                if (page["case_id"] != attempt["case_id"] or page["snapshot_id"] != attempt["snapshot_id"]
                        or page["document_id"] != document["id"] or not page.get("parser_version")):
                    raise EvaluationError("snapshot_document_binding")
                document["parser_version"] = page["parser_version"]
                batch = page["evidence"]
                if not batch:
                    break
                if any(row["document_id"] != document["id"] or row["ordinal"] <= after
                       or row.get("parser_version") != document["parser_version"] for row in batch):
                    raise EvaluationError("snapshot_document_binding")
                if page["next_ordinal"] != batch[-1]["ordinal"]:
                    raise EvaluationError("evidence_pagination_invalid")
                blocks.extend(batch)
                after = batch[-1]["ordinal"]
            if len({row["id"] for row in blocks}) != len(blocks):
                raise EvaluationError("evidence_pagination_invalid")
            evidence[name] = blocks
        definition = business["checklist_definition"]
        buyer_refs = {}
        for criterion in definition["criteria"]:
            if not criterion.get("enabled", True):
                continue
            anchor = business["buyer_anchors"][criterion["criterion_id"]]
            context = _anchor_context(anchor, evidence[anchor["document"]])
            criterion["source"] = {**criterion["source"], "context": context}
            for ref in context["evidence_refs"]:
                if ref in buyer_refs and buyer_refs[ref]["item_id"] != anchor["item_id"]:
                    raise EvaluationError("buyer_source_identity_ambiguous")
                buyer_refs[ref] = anchor
        attempt["stage"] = "checklist"
        draft = edit("PUT", f"{prefix}/checklist", {**definition, "expected_revision": revision,
                     "expected_checklist_revision": 0, "reason": "冻结合成开发评估的预定业务清单"})
        bound(draft)
        acknowledge_receipt()
        confirmed = edit("POST", f"{prefix}/checklist/confirm", {
            "criterion_ids": [c["criterion_id"] for c in definition["criteria"] if c.get("enabled", True)],
            "confirmed": True, "expected_revision": revision,
            "expected_checklist_revision": draft["checklist_revision"], "reason": "明确确认预定评估清单"})
        bound(confirmed)
        if not confirmed.get("execution_available"):
            raise EvaluationError("checklist_confirmation_mismatch")
        version = confirmed["confirmed_version"]
        attempt["checklist_version_id"] = _identifier(version["id"])
        # Compare normalized definitions because API materializes default fields.
        from backend.app.criteria.models import ChecklistDefinition
        expected_definition = ChecklistDefinition.model_validate(definition).model_dump(mode="json")
        if ChecklistDefinition.model_validate(version["definition"]).model_dump(mode="json") != expected_definition:
            raise EvaluationError("checklist_confirmation_mismatch")
        acknowledge_receipt()
        attempt["stage"] = "bindings"
        entries = []
        for binding in business["bindings"]:
            document = documents[binding["document"]]
            if document["role"] == "requirement":
                continue
            selector = _selector(binding["scope"])
            entries.append({"document_id": document["id"], "source_sha256": document["sha256"],
                            "parser_version": document["parser_version"], "item_id": binding["item_id"],
                            "selector": selector, "anchor": _anchor_context(binding["anchor"], evidence[binding["document"]]),
                            "identifier_class": "sku" if binding.get("sku") else "category",
                            "supplier_sku": binding.get("sku"),
                            "sku_anchor": (_anchor_context(binding["sku_anchor"], evidence[binding["document"]])
                                           if binding.get("sku_anchor") else None),
                            "single_item_confirmed": selector["kind"] == "whole"})
        bindings = edit("PUT", f"{prefix}/item-bindings", {"entries": entries,
                        "checklist_version_id": attempt["checklist_version_id"], "expected_revision": revision,
                        "expected_checklist_revision": confirmed["checklist_revision"], "expected_binding_revision": 0,
                        "confirmed": True, "reason": "明确确认合成原件的预定品目范围与型号"})
        if (bindings["case_id"] != attempt["case_id"] or bindings["snapshot_id"] != attempt["snapshot_id"]
                or bindings["state"] != "confirmed"):
            raise EvaluationError("binding_confirmation_mismatch")
        attempt["item_binding_version_id"] = _identifier(bindings["current_version"]["id"])
        attempt["item_binding_digest"] = bindings["current_version"]["digest"]
        acknowledge_receipt()
        attempt["stage"] = "submit"
        conversation = request("POST", f"/api/v1/cases/{attempt['case_id']}/conversations", expected=201)
        attempt["conversation_id"] = _identifier(conversation["id"])
        acknowledge_receipt()
        receipt = edit("POST", f"/api/v1/conversations/{attempt['conversation_id']}/messages", {
            "content": QUESTION, "mode": "thorough", "operation": "review", "expected_revision": revision}, expected=202)
        attempt["run_id"] = _identifier(receipt["run_id"])
        attempt["attempt_id"] = _identifier(receipt["attempt_id"])
        attempt["message_id"] = _identifier(receipt["message_id"])
        acknowledge_receipt()
        attempt["stage"] = "run"
        run_path = f"/api/v1/agent-runs/{attempt['run_id']}"
        while True:
            run = request("GET", run_path)
            bound(run, run=True)
            if (run["id"] != attempt["run_id"] or run["active_attempt_id"] != attempt["attempt_id"]
                    or run["conversation_id"] != attempt["conversation_id"]
                    or run.get("checklist_version_id") != attempt["checklist_version_id"]
                    or run.get("item_binding_version_id") != attempt["item_binding_version_id"]
                    or run.get("item_binding_digest") != attempt["item_binding_digest"]):
                raise EvaluationError("run_receipt_mismatch")
            attempt["run"] = run
            observation["run_status"] = run["status"]
            attempt["outcome_unknown"] = bool(run.get("outcome_unknown"))
            usage = run.get("usage")
            unknown_calls = usage.get("unknown_usage_calls") if usage else None
            attempt["usage"] = {"state": ("partial" if unknown_calls else "actual") if usage is not None else "unknown",
                                "actual": usage, "unknown_calls": unknown_calls, "cost": None}
            if attempt["outcome_unknown"]:
                observation["run_status"] = "outcome_unknown"
                raise EvaluationError("server_outcome_unknown")
            if run["status"] in TERMINAL:
                break
            sleep(min(0.5, max(0, deadline - clock())))
        needs_review = run.get("correction_review_status") == "needs_review"
        visible, errors = _visible_cells(business, run.get("answer_text"))
        observation["visible_cells"], observation["surface_errors"] = visible, errors
        attempt["stage"] = "artifacts"
        for kind, suffix in (("checks", "/checks"), ("progress", "/progress"),
                             ("export", "/export?format=json"), ("xlsx", "/export?format=xlsx")):
            try:
                payload = request("GET", run_path + suffix, binary=kind == "xlsx")
                if kind in {"checks", "export"}:
                    checks = payload["result"] if kind == "export" else payload
                    bound(checks, run=True)
                    if (checks["run_id"] != attempt["run_id"]
                            or checks.get("checklist_version_id") != attempt["checklist_version_id"]
                            or checks.get("item_binding_version_id") != attempt["item_binding_version_id"]
                            or checks.get("item_binding_digest") != attempt["item_binding_digest"]):
                        raise EvaluationError("run_receipt_mismatch")
                    observation["export_cells" if kind == "export" else "api_cells"] = _cells(pack_id, checks["checks"], needs_review=needs_review)
                elif kind == "progress":
                    if (payload["run_id"] != attempt["run_id"]
                            or payload.get("item_binding_digest") != attempt["item_binding_digest"]):
                        raise EvaluationError("run_receipt_mismatch")
                    item_coverage = payload.get("item_coverage", {})
                    for binding in business["bindings"]:
                        supplier = documents[binding["document"]].get("supplier_handle")
                        done = item_coverage.get(supplier, {}).get(binding["item_id"])
                        if type(done) is bool:
                            observation["coverage"].append({"document": binding["document"], "item_id": binding["item_id"],
                                "scope": binding["scope"], "processing_status": "complete" if done else "partial", "source": run_path + suffix})
                    observation["coverage_unknown"] = not bool(observation["coverage"])
                else:
                    attempt["xlsx_bytes"] = payload
                    observation["xlsx_export_cells"] = _xlsx_cells(pack_id, payload, needs_review=needs_review)
            except (EvaluationError, KeyError, TypeError, ValueError, OSError):
                observation["surface_errors"].append(f"{kind}_artifact_unavailable_or_invalid")
        if observation["xlsx_export_cells"] != observation["export_cells"]:
            observation["surface_errors"].append("json_xlsx_export_mismatch")
        if observation["api_cells"] != observation["export_cells"]:
            observation["surface_errors"].append("api_json_export_mismatch")
        observed_sources = {}

        def source(block, document, item_id):
            value = {"evidence_id": block["id"], "sha256": document["sha256"],
                     "document": document["filename"], "role": document["role"],
                     "supplier_handle": document.get("supplier_handle"), "item_id": item_id,
                     "text": block["source_text"], "page": None if block.get("page_idx") is None else block["page_idx"] + 1,
                     "sheet": block.get("sheet"), "cell": block.get("cell"), "bbox": block.get("bbox")}
            old = observed_sources.get(block["id"])
            if old is not None and old != value:
                raise EvaluationError("observed_source_identity_ambiguous")
            observed_sources[block["id"]] = value

        for ref, anchor in buyer_refs.items():
            block = next(row for row in evidence[anchor["document"]] if row["id"] == ref)
            source(block, documents[anchor["document"]], anchor["item_id"])
        for binding in business["bindings"]:
            document = documents[binding["document"]]
            for block in evidence[binding["document"]]:
                if _in_scope(block, binding["scope"]):
                    source(block, document, binding["item_id"])
        observation["sources"] = list(observed_sources.values())
        attempt["error_code"] = run.get("error_code")
        attempt["stage"] = "finished"
    except (EvaluationError, KeyError, TypeError, ValueError, OSError) as exc:
        attempt["error_code"] = str(exc) if isinstance(exc, EvaluationError) else "invalid_input_or_api_shape"
    except KeyboardInterrupt:
        attempt["error_code"] = "client_interrupted"
    finally:
        still_active = bool(attempt.get("run_id") and observation["run_status"] not in TERMINAL)
        attempt["outcome_unknown"] = bool(attempt["outcome_unknown"] or _unresolved_mutation(attempt) or still_active)
        observation["error_code"] = attempt["error_code"]
        attempt["elapsed_ms"] = round((clock() - started) * 1000, 2)
        save()
    return attempt


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path("data/dynamic-dev-v1"))
    parser.add_argument("--pack", action="append")
    parser.add_argument("--send", action="store_true")
    parser.add_argument("--api-url", default="http://127.0.0.1:8000")
    parser.add_argument("--deadline-seconds", type=int, default=240)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        if not 1 <= args.deadline_seconds <= 600:
            raise ValueError("evaluation_deadline_invalid")
        protocol = prepare_protocol(args.root, packs=args.pack)
        output = args.output or Path("data/dynamic-runs") / (datetime.now(UTC).strftime("%Y%m%d-%H%M%S-") + uuid4().hex[:8])
        output.mkdir(parents=True, exist_ok=False)
        _write_new(output / "protocol.json", protocol)
        if not args.send:
            print(json.dumps({"mode": "preview", "path": str(output),
                              **{k: protocol[k] for k in ("planned_tasks", "planned_files", "planned_positions", "maximum_agent_requests")}}, ensure_ascii=False))
            return 0
        url = validate_api_url(args.api_url)
        token = os.environ.get("DOCUCHECK_EVAL_TOKEN")
        if not token:
            raise ValueError("evaluation_oidc_token_missing")
        observations, stopped = [], None
        journals = {task["task_id"]: AttemptJournal(output / f"{task['task_id']}.inflight.json", _unstarted(task))
                    for task in protocol["tasks"]}
        project_attempt = {"stage": "create_project", "outcome_unknown": False, "raw": [],
                           "automatic_retries": False}
        project_journal = AttemptJournal(output / "project.inflight.json", project_attempt)
        with httpx.Client(base_url=url, headers={"Authorization": f"Bearer {token}"},
                          timeout=15, follow_redirects=False, trust_env=False) as client:
            try:
                project_body = {"name": "Synthetic dynamic development evaluation " + uuid4().hex[:8]}
                project_key = uuid4().hex
                project_record = {"method": "POST", "path": "/api/v1/projects", "state": "sending",
                                  "json": project_body, "idempotency_key": project_key}
                project_attempt["raw"].append(project_record)
                project_journal.update(project_attempt)
                response = client.post("/api/v1/projects", json=project_body,
                                       headers={"Idempotency-Key": project_key})
                project_record.update(state="received", http_status=response.status_code)
                if response.status_code != 201:
                    raise EvaluationError(f"http_status_{response.status_code}")
                project = response.json()
                project_record["response"] = project
                project_journal.update(project_attempt)
                project_id = _identifier(project["id"])
                project_attempt["project_id"] = project_id
                project_record["receipt_validated"] = True
                _write_new(output / "project.json", project)
            except (httpx.HTTPError, EvaluationError, ValueError, KeyError, TypeError):
                project_id, stopped = None, "project_creation_failed_or_unknown"
                project_attempt.update(outcome_unknown=True, error_code=stopped)
            project_journal.update(project_attempt, finished=True)
            for task in protocol["tasks"]:
                if not stopped:
                    try:
                        if sha256((args.root / "index.json").read_bytes()).hexdigest() != protocol["index_sha256"]:
                            stopped = "evaluation_frozen_index_changed"
                    except OSError:
                        stopped = "evaluation_frozen_index_changed"
                if not stopped:
                    try:
                        config = args.root / "config" / f"{task['pack_id']}.json"
                        gold = args.root / "gold" / f"{task['pack_id']}.json"
                        if (sha256(config.read_bytes()).hexdigest() != task["manifest_sha256"]
                                or sha256(gold.read_bytes()).hexdigest() != task["gold_sha256"]):
                            stopped = "evaluation_frozen_configuration_changed"
                    except OSError:
                        stopped = "evaluation_frozen_configuration_changed"
                if stopped:
                    attempt = _unstarted(task, stopped)
                else:
                    attempt = run_pack(client, args.root / "inputs" / task["pack_id"],
                                       manifest=task["manifest"], project_id=project_id, deadline_seconds=args.deadline_seconds,
                                       checkpoint=journals[task["task_id"]].update)
                    if attempt["error_code"] == "client_interrupted":
                        stopped = "client_interrupted"
                xlsx = attempt.pop("xlsx_bytes", None)
                if xlsx is not None:
                    name = f"{task['task_id']}.xlsx"
                    _write_new(output / name, xlsx)
                    attempt["xlsx_export"] = {"filename": name, "sha256": sha256(xlsx).hexdigest(), "bytes": len(xlsx)}
                _write_new(output / f"{task['task_id']}.json", attempt)
                journals[task["task_id"]].update(attempt, finished=True)
                observations.append(attempt["observation"])
                print(json.dumps({"task_id": task["task_id"], "stage": attempt["stage"], "error_code": attempt["error_code"]}, ensure_ascii=False), flush=True)
        _write_new(output / "observations.json", observations)
        return int(any(row.get("run_status") != "completed" or row.get("error_code") or row.get("surface_errors")
                       for row in observations))
    except (ValueError, OSError, KeyError, TypeError):
        print(json.dumps({"error": "evaluation_setup_or_version_invalid", "formal_independent_acceptance": False}))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
